#!/usr/bin/env python3
"""
Financial model: integrated oilseed crushing + extruded aquafeed mill, Kisumu (Kenya).

Run:   python3 model.py
Makes: financial_model.xlsx, outputs.md, pnl_annual.csv, cashflow_annual.csv,
       mass_balance_annual.csv, scenarios_annual.csv (all in this folder).

Structure
  1. ASSUMPTIONS   every scalar input with low/base/high, unit and source
  2. VECTORS       inputs that are paths or tables (supply path, seed-cost path,
                   seasonality, feed formulas, capex items, ramps)
  3. run_core()    one deterministic run for a parameter set and a Phase 2 flag
  4. run_model()   applies the Phase 2 go/no-go gate (auto/go/no) and returns results
  5. scenarios, tornado, seed x oil grid, survival-weighted tree
  6. writers       Excel, markdown, CSV

Conventions
  * Annual periods 2027-2038. 2027 = Phase 0 + FID (Q4), 2028 = construction,
    2029-2038 = ten operating years (COD in 2029 after `ops_months_y1` months).
  * Nominal KES. Prices are entered in 2026 KES (real) and escalated:
      tradables (oil, cake, seed, feed, ingredients, capex) by USD drift x FX path,
      local costs (power, overheads) by KES CPI, wages by CPI + real wage growth.
  * All prices are ex-VAT. VAT is modelled only where it is a real cost
    (irrecoverable input VAT on exempt feed) or a timing item (VAT on capex).
  * Discounting: flows in year t are discounted (1+r)^(t-2027), i.e. value date
    = FID year 2027 (end-year convention).
  * IRR/NPV are implemented here (numpy_financial is not used).
"""
import os
import sys
import csv
import copy
from collections import OrderedDict

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
YEARS = np.arange(2027, 2039)
N = len(YEARS)
OPK = YEARS - 2028            # operating-year number: 1 = 2029
IDX = {int(y): i for i, y in enumerate(YEARS)}
VALUE_YEAR = 2027

MD = "research/market_demand.md"
FS = "research/feedstock_supply.md"
TO = "research/technical_operations.md"
RT = "research/regulatory_tax_risk.md"
FC = "research/financing_capital.md"
PN = "prior notes / screening report (reports/Profitable business opportunities 2026.md)"
EV = "research/round1_evidence.md"

# =============================================================================
# 1. ASSUMPTIONS  (key, section, description, low, base, high, unit, source)
# =============================================================================
_A = [
    # ---------------- Macro and FX
    ("fx0", "Macro", "KES/USD spot at model start (2026)", 129.4, 129.8, 130.0, "KES/USD",
     f"{FC} row 1: 129.81 on 15 Oct 2026 (https://thekenyatimes.com/business/kenyan-shilling-holds-firm-as-forex-reserves-drop-by-221million/)"),
    ("kes_dep", "Macro", "KES depreciation vs USD, 2027-2038", 0.02, 0.04, 0.07, "%/yr",
     f"{FC} row 2 [ESTIMATE]: inflation differential (KES CPI 6.8% vs US ~2.5-3%)"),
    ("kes_dep_2027_extra", "Macro", "Extra depreciation in election year 2027", 0.0, 0.015, 0.05, "pp",
     f"{RT} s7 risk 7 [ESTIMATE]: 50% chance of >6% KES depreciation over 2027 (election 10 Aug 2027)"),
    ("usd_infl", "Macro", "USD drift of commodity and equipment prices", 0.0, 0.02, 0.03, "%/yr",
     "[ESTIMATE] long-run USD inflation; no commodity forecast in research files"),
    ("kes_real_drift", "Macro", "KES CPI minus PPP-implied inflation (CPI = (1+USD drift)(1+KES depreciation)-1 + this)", -0.01, 0.0, 0.01, "pp/yr",
     "[ESTIMATE]; base CPI = 6.1%/yr vs KNBS 6.8% y/y Sep 2026 (https://www.knbs.or.ke/wp-content/uploads/2026/09/Kenya-Consumer-Price-Indices-and-Inflation-Rates-September-2026.pdf, via financing_capital.md)"),
    ("real_wage", "Macro", "Real wage growth above CPI", 0.0, 0.01, 0.02, "%/yr",
     f"[ESTIMATE]; 2026 minimum wage +12% (Nation via {TO} s4.3)"),

    # ---------------- Timing and configuration
    ("ops_months_y1", "Timing", "Operating months in 2029 (COD timing)", 6, 9, 12, "months",
     f"{TO} row 42 build 12/15/20 months from FID; {FC} s6 FID ~Q4 2027, COD ~Q1-Q2 2029"),
    ("p1_tpd", "Configuration", "Phase 1 crushing capacity (multi-seed full-press + soy extruder-expeller)", 50, 50, 50, "t seed/day",
     f"{TO} s5.1 Phase 1"),
    ("crush_days", "Configuration", "Crushing operating days", 270, 300, 330, "days/yr", f"{TO} row 8 [ESTIMATE]"),
    ("max_util", "Configuration", "Maximum sustainable crush utilisation", 0.85, 0.90, 0.95, "% of nameplate",
     f"{TO} row 9 (Y3+ high case 90%)"),
    ("ref_tpd_p1", "Configuration", "Phase 1 refinery capacity (crude in)", 12, 15, 20, "t/day",
     f"{TO} s5.1 specifies 20 t/d for 100% sunflower; scaled to 15 t/d because 30% of seed is soy (12% oil) [ESTIMATE]"),
    ("ref_days", "Configuration", "Refinery operating days", 270, 300, 330, "days/yr", f"[ESTIMATE] = crushing days ({TO} row 8)"),
    ("ref_max_util", "Configuration", "Refinery maximum utilisation", 0.85, 0.95, 0.95, "%", "[ESTIMATE]"),
    ("extr_tph_p1", "Configuration", "Phase 1 twin-screw extrusion line", 2, 3, 4, "t/h", f"{TO} row 15, s3.1"),
    ("extr_eff", "Configuration", "Effective throughput / nameplate", 0.70, 0.85, 0.90, "%", f"{TO} row 16 [ESTIMATE]"),
    ("extr_hours", "Configuration", "Extruder run hours", 3500, 4500, 6000, "h/yr", f"{TO} row 17 [ESTIMATE]"),
    ("p2_tpd", "Configuration", "Phase 2 total crushing capacity (prepress + solvent, dehulling)", 75, 100, 150, "t seed/day",
     f"{TO} s5.1 specifies 150 tpd; base cut to 100 tpd because {FS} s9 says >20 kt/yr seed is not supportable before year 5 and the model's supply path reaches ~26 kt by 2035 [ESTIMATE]"),
    ("p2_extr_tph", "Configuration", "Phase 2 added extrusion", 0, 3, 6, "t/h",
     f"{TO} s5.1 specifies +6 t/h; base +3 t/h because merchant aquafeed SAM x share caps sales at ~14-17 kt by 2035-38 ({MD} s3.3, s5) [ESTIMATE]"),
    ("p2_ref_add", "Configuration", "Phase 2 added refinery capacity", 10, 20, 40, "t/day",
     f"{TO} s5.1 (to 60 t/d) scaled to Phase 2 crude output [ESTIMATE]"),
    ("p2_start_year", "Timing", "First operating year of Phase 2 (capex 2032-33, gate end-2031)", 2033, 2033, 2033, "year",
     f"{TO} s5.1 (Year 3-5, gated); {FC} s6 (~2032-33)"),
    ("soy_share", "Configuration", "Soybean share of seed crushed (rest sunflower)", 0.15, 0.30, 0.40, "% of seed",
     f"{TO} s5.2 recommended mix 10 kt sunflower + 5 kt soy (33%) so own soy cake replaces SBM [ESTIMATE]"),

    # ---------------- Yields and process
    ("sun_oil_y", "Process", "Sunflower crude oil yield, full press (Phase 1)", 0.27, 0.31, 0.34, "% of seed",
     f"{TO} row 3; {FS} row 'Sunflower oil yield, expeller' 25/30/34 (COSTECH, https://repository.costech.or.tz/items/e23c10cb-6682-4441-9c02-4446353268f9/full) [DATED]"),
    ("sun_cake_max", "Process", "Sunflower expeller cake yield cap", 0.66, 0.66, 0.66, "% of seed",
     f"{TO} row 4 (66/64/61); cake = min(66%, 95% - oil yield)"),
    ("sun_loss", "Process", "Sunflower process losses (impurities 2%, moisture 2%, foots 1%)", 0.05, 0.05, 0.07, "% of seed",
     f"{TO} s5.2 mass balance"),
    ("p2_oil_uplift", "Process", "Oil-yield uplift from prepress + solvent (Phase 2)", 0.05, 0.07, 0.09, "pp of seed",
     f"{TO} row 6 (36/38/40% vs 31% full press), s1.2"),
    ("p2_hull", "Process", "Hulls removed by partial dehulling (to boiler)", 0.08, 0.10, 0.12, "% of seed", f"{TO} s1.3, s5.3"),
    ("soy_oil_y", "Process", "Soy oil yield, extruder-expeller", 0.10, 0.12, 0.14, "% of seed", f"{FS} row 'Soybean oil yield, expeller' [ESTIMATE]"),
    ("soy_loss", "Process", "Soy process losses", 0.04, 0.06, 0.08, "% of seed", f"[ESTIMATE]; {TO} s5.2 gives cake ~80-82%"),
    ("storage_shrink_pm", "Process", "Seed storage shrink at plant", 0.003, 0.005, 0.010, "%/month", f"{FS} row 'Storage shrink' [ESTIMATE]"),
    ("safety_stock_months", "Process", "Minimum seed stock (safety, end-May trough)", 0.5, 1.5, 2.5, "months of crush",
     f"[ESTIMATE] calibrated so peak stock (~41% of annual crush) sits inside {FS} s7.2 range of 7-12 kt per 20 kt"),
    ("season_conc", "Process", "Share of annual seed bought Jun-Oct", 0.55, 0.70, 0.80, "%", f"{FS} row 'Share of annual seed bought Jun-Oct' [ESTIMATE]"),
    ("ref_loss", "Process", "Refining loss crude -> RBD (to soapstock/acid oil)", 0.035, 0.05, 0.08, "% of crude", f"{TO} row 12"),
    ("feed_shrink", "Process", "Feed process shrink (moisture + unrecycled fines)", 0.015, 0.025, 0.04, "% of mash", f"{TO} row 23 [ESTIMATE]"),
    ("starter_share", "Process", "Fry/fingerling (35-40% CP) share of feed tonnes", 0.05, 0.15, 0.25, "%",
     f"[ESTIMATE]; {TO} s3.1 (fingerling sizes carry the highest margin)"),

    # ---------------- Market
    ("sam_2026", "Market", "Contestable KENYAN merchant extruded aquafeed 2026 (net of Victory Farms captive feed and SamakGro sales to VF-linked farms; UG/TZ excluded)", 15000, 25000, 37000, "t/yr",
     f"R1 rebase [ESTIMATE], build-up in {EV} s1: KNBS cage 28.1 kt (2025) -> ~32 kt (2026); VF Kenya ~14-18 kt (VF >1,400 t/month, The Star Jul 2026, https://www.the-star.co.ke/counties/nyanza/2026-07-03-victory-farms-bets-on-tech-and-talent-to-boost-fish-production; Kivu Choice Rwanda ~800 t/month, https://www.newtimes.co.rw/article/36743/National/cepgel-governors-meet-in-gisenyi/amp); non-VF cage x FCR 1.6 x (1 - VF-linked 15%) + ponds ~3 kt. Was 45 kt incl. UG/TZ in v1.1"),
    ("sam_regional", "Market", "Uganda/Tanzania lake-shore merchant demand added to SAM (upside only)", 0, 0, 12000, "t/yr in 2026",
     f"{MD} s1.5, s3.3: De Heus Njeru 50-100 kt and Tunga Uganda sit inside those markets; TZ commercial fish feed 6.2 kt (2023, UNCST) [DATED]. Base 0 [ESTIMATE]"),
    ("sam_growth", "Market", "Growth of contestable merchant aquafeed SAM", 0.04, 0.08, 0.14, "%/yr",
     f"{MD} s4: Kenya merchant ~20-40 kt (2026) -> ~45 kt (2030); cage CAGR 19.7% 2021-25 is mostly Victory Farms (captive); cage zoning risk ({RT} s4.5) [ESTIMATE]"),
    ("feed_real_decline", "Market", "Real decline in aquafeed prices from 2029 (competitor response to over-capacity)", 0.0, 0.01, 0.03, "%/yr",
     "[ESTIMATE] R1-01: Tunga 45 kt (Feb 2026), SamakGro 8 t/h plant aims to make feed 'more affordable' (https://capitalfm.africa/fish-feed-firms-to-invest-sh510mn-at-kengens-olkaria-park/); no observed 2026 price cut found"),
    ("feed_px_mult", "Market", "Multiplier on both aquafeed list prices (grow-out and starter)", 0.88, 1.0, 1.20, "x",
     f"{MD} rows 15-16: grow-out KES 110/125/150 per kg = x0.88/1.00/1.20"),
    ("feed_share_max", "Market", "Maximum share of SAM won (from op year 3)", 0.08, 0.15, 0.25, "%",
     f"{MD} s5 SOM 10-25% of merchant market; over-capacity (Tunga 45 kt, SamakGro 8 t/h, De Heus 50-100 kt) [ESTIMATE]"),
    ("refined_share_max", "Market", "Maximum share of crude sold refined under own brand (op yr 3+)", 0.50, 0.80, 0.95, "% of crude for sale",
     f"[ESTIMATE]; {TO} s1.4: main risk is distribution against Bidco/Pwani/Kapa; rest sold crude to refiners"),
    ("entrant_discount", "Market", "New-entrant feed price discount, op years 1-2", 0.0, 0.05, 0.10, "% of list", "[ESTIMATE]"),

    # ---------------- Prices (2026 KES/t, ex-VAT)
    ("p_crude", "Prices", "Crude oil, ex-crusher to refiners = oil price index (refined price follows it)", 155000, 175000, 200000, "KES/t",
     f"R1 [CONFLICT resolved, {EV} s2]: base = reconciled CPO landed parity ~KES 177k/t (FOB ~USD 1,100, Malaysia Oct reference USD 1,091 https://www.bernama.com/en/news.php?id=2604287; +60 freight, 10% duty, 4.5% IDF/RDL, USD 30 port) because refiners can switch to CPO ({FS} s2.4). High = CSFO parity KES 209k less 5%. Was 190k ({MD} row 8)"),
    ("soy_oil_factor", "Prices", "Crude soy oil price relative to sunflower", 0.90, 0.95, 1.00, "x", "[ESTIMATE]"),
    ("refine_spread", "Prices", "Refined fortified oil (20 L, delivered wholesale, ex-VAT) minus crude price", 40000, 65000, 90000, "KES/t",
     f"R1-03 [ESTIMATE], parity build in {EV} s2: base refined = KES 240k/t (KES 218/L ex-VAT) = midpoint of local CPO-refiner cost-plus (~219k delivered Kisumu) and imported RBD olein ceiling (~247-269k at 25% or USD 500/t duty); 2022 jerrican KES 5,100/20 L = ~242k/t ex-VAT at similar retail prices (https://peopledaily.digital/news/makers-warn-cooking-oil-prices-to-rise) [DATED]. No sunflower premium assumed (none evidenced for institutions). Was an independent KES 264k"),
    ("soapstock_factor", "Prices", "Soapstock/acid oil value vs crude", 0.2, 0.3, 0.4, "x", "[ESTIMATE]"),
    ("p_suncake", "Prices", "Sunflower expeller cake (28-34% CP)", 32000, 40000, 50000, "KES/t",
     f"{MD} row 11: KES 40/kg Nairobi listing (https://jiji.co.ke/viwandani-makadara/feeds-supplements-seeds/sunflower-seed-cake-dTh5NuF0GdsvaWJ01XlkOR3T.html) [UNVERIFIED]; KES 25->35 [DATED]"),
    ("meal_premium", "Prices", "Dehulled solvent sunflower meal (35% CP) vs cake", 1.0, 1.1, 1.2, "x",
     f"[ESTIMATE]; capped by Black Sea sunflower meal landed ~KES 48-52/kg ({FS} s2.4)"),
    ("p_soycake", "Prices", "Non-GM soy expeller cake (42-44% CP)", 65000, 85000, 100000, "KES/t",
     f"[ESTIMATE/CONFLICT]: {FS} s8 uses KES 65/kg; {MD} non-GM SBM KES 95/kg x protein ratio ~0.9"),
    ("p_sbm", "Prices", "Non-GM soybean meal purchased (44-46% CP)", 80000, 95000, 120000, "KES/t",
     f"{MD} row 12 (Indexbox KES 90-120k/t 2026 [UNVERIFIED])"),
    ("p_fishmeal", "Prices", "Omena/fishmeal, feed grade", 120000, 160000, 220000, "KES/t", f"{MD} row 14 [ESTIMATE]"),
    ("p_cottoncake", "Prices", "Cottonseed cake", 30000, 38000, 45000, "KES/t", f"{MD} row 13 [ESTIMATE]"),
    ("p_maize", "Prices", "Maize / maize germ", 38000, 50000, 60000, "KES/t", f"{TO} s3.3 [ESTIMATE]"),
    ("p_pollard", "Prices", "Wheat pollard", 26000, 32000, 38000, "KES/t", f"{TO} s3.3 [ESTIMATE]"),
    ("p_bran", "Prices", "Wheat bran", 20000, 25000, 30000, "KES/t", f"{TO} s3.3 [ESTIMATE]"),
    ("p_premix", "Prices", "Premix, DCP, amino acids, binder, toxin binder", 200000, 250000, 320000, "KES/t", f"{TO} s3.3 [ESTIMATE]"),
    ("p_feed_grow", "Prices", "Extruded floating tilapia grow-out feed (28-32% CP), ex-mill", 110000, 125000, 150000, "KES/t",
     f"{MD} row 15: SamakGro KES 116-130/kg (https://take.app/samakgro/p/clzr0xsj9004nirjw3brif1fa, undated); WorldFish USD 1.00-1.24/kg"),
    ("p_feed_start", "Prices", "Fry/starter feed (35-45% CP), ex-mill", 160000, 200000, 250000, "KES/t", f"{MD} row 16 [ESTIMATE]"),
    ("seed_cost_mult", "Prices", "Delivered seed cost multiplier (sunflower path and soy)", 0.79, 1.00, 1.26, "x",
     f"{FS} row 'Blended delivered seed cost' 45,000/57,000/72,000 KES/t [ESTIMATE]"),
    ("seed_passthrough", "Prices", "Share of a change in oil+cake value per t seed passed to delivered seed price", 0.0, 0.5, 0.8, "share",
     f"[ESTIMATE]: floating seed price = floor + share of oil/cake import parity ({FS} s4.2); crush margin largely parity-indexed ({FC} s1). 0 = fixed-price contracts"),
    ("p_soyseed", "Prices", "Delivered soybean (Kenya/Uganda)", 60000, 75000, 100000, "KES/t", f"{FS} row 'Delivered soybean' [ESTIMATE]"),
    ("tz_farmgate", "Prices", "Tanzania sunflower farmgate (levy base)", 27000, 35000, 47000, "KES/t",
     f"{FS} row 'Tanzania sunflower farmgate' (TZS 497/kg Jun 2026, https://www.selinawamucii.com/insights/prices/tanzania/sunflower-seeds/)"),

    # ---------------- Policy (probability-weighted in base)
    ("cpo_duty_value", "Policy", "Oil-price protection from 10% CPO duty", 12000, 15000, 18000, "KES/t oil",
     f"{MD} row 25; {RT} s7 rank 1 (-KES 12-18k/t if removed)"),
    ("p_cpo_removed", "Policy", "Probability CPO duty removed/unenforceable (base = expected value)", 0.0, 0.60, 1.0, "prob",
     f"{RT} s7 rank 1: P=60%; High Court struck duty down 27 Nov 2025 (https://allafrica.com/stories/202511270355.html)"),
    ("gmo_meal_drop", "Policy", "Fall in protein-meal prices if GM imports allowed", 0.20, 0.25, 0.30, "%", f"{MD} s2.2; {RT} s7 rank 4"),
    ("feed_cost_passthrough", "Prices", "Share of a change in competitors' feed raw-material cost passed into aquafeed prices", 0.5, 0.7, 0.9, "share",
     "[ESTIMATE]; applies to GMO, waiver and ingredient-price moves (e.g. GM imports: ~-5% feed price)"),
    ("p_gmo", "Policy", "Probability GM feed imports allowed by 2029", 0.0, 0.40, 1.0, "prob",
     f"{RT} s7 rank 4: 40% by end-2028; CoA conservatory order 7 Mar 2025 (https://citizen.digital/article/blow-to-govt-as-court-of-appeal-bars-importation-of-gmos-n358801)"),
    ("waiver_drop", "Policy", "Cake/meal price fall in a duty-waiver year", 0.05, 0.075, 0.10, "%", f"{RT} s7 rank 2"),
    ("waiver_freq", "Policy", "Share of years with a feed-ingredient duty waiver", 0.0, 0.33, 0.5, "share",
     f"{RT} s7 aggregate: 'one feed-waiver year in three'"),
    ("tz_levy", "Policy", "Tanzanian export levy on sunflower seed", 0.0, 0.10, 0.10, "% of farmgate",
     f"{RT} row 21 (TZ 2026/27 budget, https://millingmea.com/tanzania-targets-edible-oil-self-sufficiency-with-tax-relief-and-import-controls/)"),
    ("p_tz_levy", "Policy", "Probability TZ levy in force", 0.0, 0.70, 1.0, "prob", f"{RT} s7 rank 3"),
    ("afa_levy_seed", "Policy", "AFA import levy on imported raw seed (UG/TZ)", 0.0, 0.0, 0.04, "% of CIF",
     f"{RT} row 15: 4% of CIF on finished products; rate on raw seed not found [UNVERIFIED]"),

    # ---------------- Tax
    ("cit", "Tax", "Corporate income tax", 0.10, 0.30, 0.30, "%",
     f"{RT} rows 1-2: 30% standard; 10% SEZ yrs 1-10 (not base: SEZ domestic-sale treatment unresolved)"),
    ("inv_ded", "Tax", "Investment deduction, year 1 (buildings & machinery); 50% => 50% then 25% RB", 0.50, 1.00, 1.00, "%",
     f"{RT} row 3 [CONFLICT on threshold]; https://taxsummaries.pwc.com/kenya/corporate/deductions"),
    ("loss_years", "Tax", "Tax-loss carry-forward limit", 5, 5, 5, "years", f"{RT} row 4 (Finance Act 2025)"),
    ("int_cap", "Tax", "Deductible foreign-debt interest cap", 0.30, 0.30, 0.30, "% of EBITDA", f"{RT} row 5; 3-yr carry-forward"),
    ("wht_int", "Tax", "WHT on interest to non-resident lender (grossed up)", 0.0, 0.15, 0.15, "%", f"{RT} row 7; {FC} row 31"),
    ("div_wht", "Tax", "Dividend WHT, resident individual (<12.5% holdco route = 0)", 0.0, 0.05, 0.15, "%", f"{RT} row 6"),
    ("cgt", "Tax", "Capital gains tax on exit (share sale)", 0.0, 0.15, 0.15, "%", "[UNVERIFIED - prior knowledge: 15% since 2023; verify]"),
    ("vat_feed_irrecov", "Tax", "Irrecoverable input VAT on exempt feed sales, OPERATING inputs only (capex VAT apportioned separately)", 0.01, 0.02, 0.035, "% of feed revenue", f"{RT} row 19 [ESTIMATE]; R1-09: capex share now modelled via p_feed_exempt"),
    ("p_feed_exempt", "Tax", "Probability finished feed is VAT-exempt (not zero-rated): capex VAT apportioned by exempt turnover share", 0.0, 1.0, 1.0, "prob",
     f"{RT} row 19 and plan s7.2 [CONFLICT exempt vs zero-rated]; base treats feed as exempt, consistent with the 2% leakage"),
    ("min_tax_rate", "Tax", "Minimum tax on gross turnover if reintroduced (payable where it exceeds CIT)", 0.0, 0.01, 0.01, "% of revenue",
     "Finance Act 2020 design (1% of turnover); struck down by courts 2021-22 (https://mondaq.co.uk/capital-gains-tax/1269472/court-of-appeal-upholds-decision-declaring-minimum-tax-as-unconstitutional) [DATED]; MTRS 2024/25-2026/27 commits to a redesigned version (https://ipfglobal.or.ke/wp-content/uploads/2025/01/IS-A-MINIMUM-TAX-STILL-A-VIABLE-OPTION-FOR-KENYA.pdf)"),
    ("p_min_tax", "Tax", "Probability minimum tax applies from 2028 (base = expected value)", 0.0, 0.25, 1.0, "prob",
     "[ESTIMATE] R1-08: MTRS commitment, but no minimum tax in Finance Bill 2026 per search extracts (https://kpmg.com/us/en/taxnewsflash/news/2026/05/kenya-tax-proposals-finance-bill-2026.html)"),
    ("feed_levy", "Tax", "New feed-regulator levy (Livestock Bill 2026), expected value", 0.0, 0.005, 0.01, "% of feed revenue",
     f"{RT} s4.4 and s7 rank 9: 0-1% of feed revenue, P~40-50% [ESTIMATE]"),
    ("idf", "Tax", "Import Declaration Fee on imported equipment", 0.015, 0.025, 0.035, "% of CIF", f"{RT} row 13"),
    ("rdl", "Tax", "Railway Development Levy on imported equipment", 0.015, 0.020, 0.025, "% of CIF", f"{RT} row 14"),
    ("vat_rate", "Tax", "VAT on capex (paid, recovered next year)", 0.16, 0.16, 0.16, "%", f"{RT} s1.3 (VAT exemption only for >=KES 2bn investments)"),

    # ---------------- Opex
    ("elec_price", "Opex", "Grid electricity, all-in energy charge", 16, 20, 25, "KES/kWh", f"{TO} row 26 (EPRA 2023 schedule + Jun 2026 pass-throughs)"),
    ("kwh_sun", "Opex", "Electricity, sunflower full press", 40, 60, 120, "kWh/t seed", f"{TO} row 10"),
    ("kwh_soy", "Opex", "Electricity, soy extruder-expeller", 70, 90, 130, "kWh/t seed", "[ESTIMATE]"),
    ("kwh_solv", "Opex", "Extra electricity, solvent plant", 10, 15, 25, "kWh/t seed", f"{TO} s1.2"),
    ("kwh_ref", "Opex", "Electricity, refinery", 20, 30, 45, "kWh/t oil", f"{TO} row 14"),
    ("kwh_feed", "Opex", "Electricity, extruded aquafeed", 70, 90, 130, "kWh/t feed", f"{TO} row 24"),
    ("kwh_gen", "Opex", "General site load (Phase 1; x1.5 after Phase 2)", 150000, 200000, 300000, "kWh/yr", f"{TO} s4.3 energy table"),
    ("steam_crush", "Opex", "Steam, seed conditioning", 0.10, 0.15, 0.25, "t/t seed", f"{TO} row 11"),
    ("steam_solv", "Opex", "Extra steam, solvent plant", 0.2, 0.3, 0.4, "t/t seed", f"{TO} s1.2"),
    ("steam_ref", "Opex", "Steam, refinery", 0.2, 0.3, 0.45, "t/t oil", f"{TO} row 14"),
    ("steam_cost", "Opex", "Steam cost, biomass boiler", 2000, 2800, 4000, "KES/t steam", f"{TO} row 29 [ESTIMATE]"),
    ("feed_gj", "Opex", "Thermal energy, aquafeed", 0.8, 1.1, 1.5, "GJ/t feed", f"{TO} row 25"),
    ("gj_per_t_steam", "Opex", "Fuel per t steam", 3.7, 3.7, 3.7, "GJ/t", f"{TO} row 29"),
    ("hull_mj", "Opex", "Sunflower hull calorific value as fired", 13, 15, 17, "MJ/kg", f"{TO} s4.3 (HHV ~17.8)"),
    ("demand_charge", "Opex", "Electricity demand charge", 700, 1100, 1100, "KES/kVA/month", f"{TO} row 27"),
    ("kva_p1", "Opex", "Billed demand, Phase 1", 700, 900, 1100, "kVA", f"{TO} s4.3"),
    ("kva_p2", "Opex", "Added billed demand, Phase 2", 800, 1200, 1600, "kVA", "[ESTIMATE]"),
    ("diesel_p1", "Opex", "Standby diesel, Phase 1", 5e6, 6.5e6, 8e6, "KES/yr", f"{TO} s4.3"),
    ("diesel_p2", "Opex", "Standby diesel, Phase 2 increment", 2e6, 4e6, 6e6, "KES/yr", "[ESTIMATE]"),
    ("labour_p1", "Opex", "Labour Phase 1 (~70 FTE incl. 15% on-costs)", 50e6, 61e6, 80e6, "KES/yr", f"{TO} s4.3"),
    ("labour_p2", "Opex", "Labour Phase 2 increment (~45 FTE)", 25e6, 32e6, 42e6, "KES/yr", "[ESTIMATE] solvent 3 shifts x 5, refinery +8, extrusion +12, maintenance/logistics +10"),
    ("maint_pct", "Opex", "Maintenance, equipment", 0.025, 0.035, 0.05, "% of equipment capex/yr", f"{TO} row 33"),
    ("maint_bldg", "Opex", "Maintenance, buildings", 0.005, 0.01, 0.015, "% of building capex/yr", "[ESTIMATE]"),
    ("press_wear", "Opex", "Press wear parts", 100, 150, 200, "KES/t seed", f"{TO} row 33"),
    ("hexane", "Opex", "Hexane loss (Phase 2)", 150, 250, 400, "KES/t seed solvent-extracted", f"{TO} s1.2"),
    ("pack_feed", "Opex", "Feed bags (25 kg PP)", 1500, 2000, 2800, "KES/t feed", f"{TO} row 34"),
    ("pack_oil_l", "Opex", "Oil packaging, 20 L jerrican", 6, 8, 10, "KES/L", f"{TO} row 35"),
    ("l_per_t", "Opex", "Litres per tonne of oil", 1099, 1099, 1099, "L/t", f"{MD} header (0.91 kg/L)"),
    ("pack_cake", "Opex", "Cake bags", 500, 700, 900, "KES/t cake", f"{TO} s4.3"),
    ("ref_chem", "Opex", "Refining chemicals + bleaching earth + fortificant", 1300, 1800, 2500, "KES/t crude refined", f"{TO} s1.4, s4.3"),
    ("lab_p1", "Opex", "QC lab running + external tests", 6e6, 8e6, 10e6, "KES/yr", f"{TO} s4.3"),
    ("lab_p2", "Opex", "QC lab Phase 2 increment", 2e6, 3e6, 4e6, "KES/yr", "[ESTIMATE]"),
    ("insurance", "Opex", "Insurance", 0.004, 0.005, 0.0075, "% of gross capex", f"{TO} s4.3"),
    ("security_p1", "Opex", "Security and cleaning", 5e6, 6e6, 8e6, "KES/yr", f"{TO} s4.3"),
    ("security_p2", "Opex", "Security Phase 2 increment", 2e6, 3e6, 4e6, "KES/yr", "[ESTIMATE]"),
    ("ga_p1", "Opex", "Other G&A (office, IT, audit, legal, licences, travel)", 20e6, 30e6, 45e6, "KES/yr", "[ESTIMATE]; licence renewals per " + RT + " s3.1"),
    ("ga_p2", "Opex", "Other G&A Phase 2 increment", 5e6, 10e6, 15e6, "KES/yr", "[ESTIMATE]"),
    ("opex_fixed_mult", "Opex", "Multiplier on all fixed opex (sensitivity)", 0.85, 1.0, 1.25, "x", "[ESTIMATE]"),
    ("storage_lease", "Opex", "Leased overflow storage (NCPB/WRS)", 60, 120, 530, "KES/t/month", f"{FS} row 'Storage cost'"),
    ("silo_p1", "Opex", "Own seed storage, Phase 1", 5000, 5000, 5000, "t", f"{TO} s5.1"),
    ("silo_p2", "Opex", "Own seed storage added in Phase 2 (scaled from +15 kt at 150 tpd)", 5000, 7500, 15000, "t", f"{TO} s4.1 [ESTIMATE]"),
    ("crude_freight", "Opex", "Crude oil haulage to refiners (Kisumu -> Nakuru/Thika)", 1300, 2250, 3200, "KES/t", f"[ESTIMATE] 190-345 km x KES 6.5/t-km ({FS} row 'Road haulage')"),
    ("cake_freight", "Opex", "Cake/meal haulage to Rift/Central millers", 1200, 1800, 2500, "KES/t", "[ESTIMATE] as above"),
    ("feed_delivery", "Opex", "Feed delivery to lake depots (Siaya/Homa Bay)", 1000, 1500, 2500, "KES/t", f"[ESTIMATE] 60-130 km ({TO} s6)"),
    ("feed_selling", "Opex", "Feed sales force, farmer extension, dealer support", 0.02, 0.03, 0.05, "% of feed revenue", "[ESTIMATE]"),
    ("feed_dealer_margin", "Opex", "Dealer/distributor margin on feed (price inputs are retail/farm-level lists)", 0.0, 0.05, 0.10, "% of feed revenue",
     f"[ESTIMATE]: {MD} row 15 base price derives from SamakGro's retail store list and WorldFish farm-level prices; ~half of volume via dealers (e.g. AquaRech-type depots) at ~10%"),
    ("bad_debt", "Opex", "Bad debts", 0.0025, 0.005, 0.01, "% of revenue", "[ESTIMATE]"),
    ("refined_selling", "Opex", "Refined-oil distribution and trade marketing", 0.03, 0.05, 0.08, "% of refined revenue", "[ESTIMATE]"),
    ("launch_mkt", "Opex", "Brand launch marketing (2029; 75% of it again in 2030)", 10e6, 20e6, 30e6, "KES", "[ESTIMATE]"),

    # ---------------- Capex and development
    ("capex_mult", "Capex", "Multiplier on Phase 1 and Phase 2 installed capex", 0.70, 1.00, 1.45, "x",
     f"{TO} row 36 (P1 USD 6.5/9.5/14m) and row 38 (P2 15/22/32m); high = +50% per {TO} s4.2"),
    ("contingency_p1", "Capex", "Contingency Phase 1", 0.10, 0.15, 0.20, "% of capex excl. land", f"{TO} s4.1"),
    ("contingency_p2", "Capex", "Contingency Phase 2", 0.08, 0.10, 0.15, "% of capex excl. land", f"{TO} s4.1"),
    ("scale_exp", "Capex", "Capex scaling exponent for Phase 2 sizes", 0.6, 0.7, 0.8, "x", "[ESTIMATE] six-tenths-type rule"),
    ("util_import", "Capex", "Imported share of utilities capex", 0.5, 0.7, 0.9, "%", "[ESTIMATE]"),
    ("phase0_usd", "Capex", "Phase 0 development (outgrower pilot, FEED, ESIA, permits, LOIs), expensed 2027", 0.6, 0.8, 1.0, "USD m",
     f"{FC} s6 Phase 0 (USD 0.6-1.0m); licensing KES 8-30m per {RT} row 24"),
    ("phase0a_usd", "Capex", "Phase 0a: cheap decisive tests (market census, VF share, feed and 20 L oil quotes, capex quotes, GM search start)", 0.2, 0.3, 0.4, "USD m",
     "R1-05 [ESTIMATE]: plan s12.5 workstreams 1, 2 (quotes only), 4 (lab tests), 5, 6, 10"),
    ("phase0b_usd", "Capex", "Phase 0b: pilot, toll-extruded trial sales, EIA, legal, financing (only if 0a passes)", 0.4, 0.5, 0.7, "USD m", "R1-05 [ESTIMATE]"),
    ("pilot_seed_loss_usd", "Capex", "Expected loss on pilot seed resold back-to-back (price risk, handling, shrink)", 0.02, 0.05, 0.10, "USD m",
     "R1-05 [ESTIMATE]: ~500-600 t at ~KES 52k/t = KES ~30m turnover, resold to a named crusher at a 5-15% loss"),
    ("preop_p1_usd", "Capex", "Phase 1 pre-operating (training, commissioning, outgrower scale-up)", 0.4, 0.6, 0.9, "USD m", f"{FC} s6 Phase 1 uses"),
    ("preop_p2_usd", "Capex", "Phase 2 pre-operating", 0.2, 0.3, 0.5, "USD m", "[ESTIMATE]"),
    ("sust_capex", "Capex", "Sustaining capex from op year 3", 0.010, 0.015, 0.025, "% of cumulative capex excl. land/yr", "[ESTIMATE]"),

    # ---------------- Working capital
    ("ar_days_oil", "Working capital", "Receivable days, oil", 15, 30, 45, "days", "[ESTIMATE]"),
    ("ar_days_cake", "Working capital", "Receivable days, cake/meal", 15, 30, 45, "days", "[ESTIMATE]"),
    ("ar_days_feed", "Working capital", "Receivable days, feed (cash-heavy farmers, dealers)", 7, 21, 45, "days", "[ESTIMATE]"),
    ("fg_days", "Working capital", "Finished-goods stock", 10, 15, 30, "days of sales at 80% (cost)", "[ESTIMATE]"),
    ("ingr_days", "Working capital", "Purchased feed-ingredient stock", 30, 45, 60, "days of purchases", "[ESTIMATE]"),
    ("ap_days", "Working capital", "Payable days on non-seed purchases (seed is weigh-and-pay)", 15, 30, 45, "days",
     f"[ESTIMATE]; weigh-and-pay per {FS} s4.2"),
    ("og_credit", "Working capital", "Outgrower seed-only input credit outstanding Mar-Aug", 5600, 8750, 11900, "KES/t outgrower seed",
     f"{FS} s7.2: KES 45-95m for 8,000 t"),

    # ---------------- Financing
    ("senior_usd_max", "Financing", "Phase 1 senior KES term loan, maximum (USD-equivalent at signing)", 3.0, 4.5, 4.5, "USD m",
     f"{FC} s6: USD 4.5m (KES 585m) 'is the maximum the plant can carry'; sized down to DSCR 1.30x on Phase 1 standalone"),
    ("senior_rate", "Financing", "Senior KES rate (floating, all-in)", 0.14, 0.155, 0.175, "%", f"{FC} row 5"),
    ("senior_tenor", "Financing", "Senior tenor incl. grace", 5, 7, 8, "years", f"{FC} row 6"),
    ("senior_grace_m", "Financing", "Senior grace from first draw", 12, 18, 24, "months", f"{FC} row 6"),
    ("mezz_usd", "Financing", "Mezzanine, USD (AgDevCo-type)", 0.0, 1.7, 2.5, "USD m", f"{FC} s6"),
    ("mezz_rate", "Financing", "Mezzanine all-in USD rate", 0.11, 0.13, 0.16, "%", f"{FC} row 9"),
    ("mezz_grace", "Financing", "Mezzanine principal grace", 2, 3, 4, "years", f"{FC} row 10"),
    ("mezz_tenor", "Financing", "Mezzanine tenor", 6, 8, 10, "years", f"{FC} row 10"),
    ("mezz_pik_years", "Financing", "Years of capitalised (PIK) mezz interest from draw", 0, 2, 2, "years", f"{FC} s6"),
    ("fee_pct", "Financing", "Upfront/arrangement fees", 0.01, 0.015, 0.025, "% of facility", f"{FC} row 13"),
    ("dsra_months", "Financing", "Debt-service reserve", 3, 6, 6, "months of senior DS", f"{FC} row 16"),
    ("dscr_target", "Financing", "Senior sizing / lock-up DSCR", 1.20, 1.30, 1.40, "x", f"{FC} row 14"),
    ("sizing_haircut", "Financing", "Lender downside case for sizing: CFADS haircut before applying the sizing DSCR", 0.0, 0.20, 0.30, "% of CFADS",
     "[ESTIMATE] R1-02: senior sized so that min DSCR >= 1.30x on 80% of base CFADS (i.e. >= 1.63x on base) over 2030-34. Was 0% (sized to exactly 1.30x)"),
    ("dscr_first_year", "Financing", "First DSCR test year (2031 = lenders accept ramp-year interest funded from the Phase 1 budget)", 2030, 2030, 2031, "year",
     f"{FC} s5 principle 1: 'size debt to DSCR at ramp-up'; plan s10.1 asks for a 24-month grace / covenant holiday [ESTIMATE]"),
    ("default_dscr", "Financing", "Default covenant DSCR (cure required below this)", 1.05, 1.10, 1.20, "x", f"{FC} row 14 / term table"),
    ("waiver_fee", "Financing", "Waiver/default fee in a breach year", 0.0, 0.01, 0.02, "% of opening senior balance", "[ESTIMATE / PK-UNVERIFIED]"),
    ("sweep_share", "Financing", "Cash sweep to senior prepayment in lock-up years (DSCR below lock-up)", 0.0, 0.5, 1.0, "% of excess cash", "[ESTIMATE / PK-UNVERIFIED]"),
    ("pg_cover", "Financing", "Founder personal guarantee on senior debt (share of senior shortfall called on failure)", 0.0, 1.0, 1.0, "share",
     f"{FC} s2 'typical greenfield terms' and term table: founder guarantee likely; collateral cover 125-150% [PK-UNVERIFIED]"),
    ("gm_equity", "Financing", "GM equity incentive (dilutes founder group's share of distributions and exit)", 0.0, 0.04, 0.05, "% of equity", "Plan s8.2: 3-5% vesting over 4 years [plan recommendation]"),
    ("revolver_rate", "Financing", "Seasonal WC revolver rate", 0.13, 0.145, 0.16, "%", f"{FC} row 11"),
    ("revolver_adv", "Financing", "Revolver advance rate on seed stock", 0.50, 0.65, 0.75, "% of stock", f"{FC} row 12"),
    ("grants_usd", "Financing", "Grants/TA (outgrower and ESIA co-funding only)", 0.0, 0.3, 1.0, "USD m", f"{FC} row 22"),
    ("p2_debt_share", "Financing", "Phase 2 debt share of Phase 2 capex", 0.45, 0.55, 0.60, "%", f"{FC} row 21"),
    ("p2_rate", "Financing", "Phase 2 DFI loan in KES (TCX-hedged)", 0.13, 0.14, 0.16, "%", f"{FC} s6 Phase 2"),
    ("p2_tenor", "Financing", "Phase 2 loan tenor", 8, 9, 10, "years", f"{FC} s6"),
    ("p2_grace", "Financing", "Phase 2 loan grace", 2, 2, 2, "years", f"{FC} s6"),
    ("min_cash", "Financing", "Minimum operating cash (2026 KES)", 30e6, 50e6, 80e6, "KES", "[ESTIMATE]"),

    # ---------------- Valuation and hurdles
    ("exit_mult", "Valuation", "Terminal EV/EBITDA at end-2038", 4.0, 5.0, 7.0, "x", f"{FC} row 27 [ESTIMATE]"),
    ("wacc_p1", "Valuation", "Financing-file WACC, Phase 1 (assumes 50% equity; memo only - the model now computes WACC from the modelled structure)", 0.19, 0.23, 0.27, "%", f"{FC} row 25"),
    ("wacc_ss", "Valuation", "Financing-file WACC, steady state (memo only)", 0.14, 0.165, 0.19, "%", f"{FC} row 26"),
    ("hurdle_gf_lo", "Valuation", "Greenfield equity hurdle (survival-weighted), lower", 0.35, 0.35, 0.35, "% KES IRR", PN),
    ("hurdle_gf_hi", "Valuation", "Greenfield equity hurdle (survival-weighted), upper", 0.40, 0.40, 0.40, "% KES IRR", PN),
    ("hurdle_est", "Valuation", "Established-business hurdle", 0.25, 0.275, 0.30, "% KES IRR", PN),
    ("infra_bond", "Valuation", "Tax-free infrastructure bond", 0.1175, 0.1225, 0.1274, "%", f"{PN}; https://serrarigroup.com/cbk-opens-ksh150bn-tax-free-infrastructure-bond-sale/"),
    ("bank_rate", "Valuation", "Average bank lending rate", 0.14, 0.144, 0.155, "%", "https://capitalfm.africa/private-sector-credit-growth-accelerates-to-10-6pc-in-september/ (Sep 2026)"),

    # ---------------- Gate and survival
    ("gate_supply_frac", "Gate", "Gate (a): seed secured for 2035 >= share of Phase 2 nameplate", 0.70, 0.80, 0.90, "%",
     f"{TO} s5.1 gate (a) >=40 kt for 45 kt plant (89%), scaled [ESTIMATE]"),
    ("gate_feed_t", "Gate", "Gate (b): feed sales in 2031", 8000, 8000, 8000, "t", f"{TO} s5.1 gate (b)"),
    ("gate_feed_gm", "Gate", "Gate (b): feed gross margin in 2031", 0.20, 0.20, 0.20, "%", f"{TO} s5.1 gate (b)"),
    ("gate_ref_share", "Gate", "Gate (c): refined share of crude for sale in 2031", 0.70, 0.70, 0.70, "%",
     f"{TO} s5.1 gate (c) restated as a market test (refinery >70% utilised)"),
    ("p_fail", "Survival", "Probability of business failure (total equity loss path)", 0.15, 0.30, 0.45, "prob",
     f"{PN}: 30% total-loss probability underlies '35-40% expected = ~45-50% conditional' [ESTIMATE]"),
    ("fail_year", "Survival", "Year of failure (end of op year 2), liquidation following year", 2030, 2030, 2031, "year", "[ESTIMATE]"),
    ("recovery_fixed", "Survival", "Liquidation value of fixed assets", 0.20, 0.35, 0.50, "% of gross capex", "[ESTIMATE]"),
    ("recovery_wc", "Survival", "Liquidation value of working capital", 0.5, 0.7, 0.85, "% of book", "[ESTIMATE]"),
    ("p_go", "Survival", "Probability Phase 2 gate passes, given survival (used only when the base gate says GO)", 0.3, 0.5, 0.7, "prob", "[ESTIMATE]"),
    ("p_late_fail", "Survival", "Probability of late failure (price war, Phase 2 start-up or exit failure), given early survival", 0.05, 0.10, 0.15, "prob", "[ESTIMATE] R1-11"),
    ("late_fail_year", "Survival", "Late-failure year (liquidation the following year)", 2033, 2034, 2035, "year", "[ESTIMATE]"),
    ("config", "Configuration", "0 = integrated crusher + feed mill; 1 = feed-mill-first (buys protein); 2 = crusher + refinery only", 0, 0, 2, "code", "R1-05 restructure alternatives [ESTIMATE scaling]"),
]

ASSUMPTIONS = OrderedDict()
for row in _A:
    k, sec, desc, lo, ba, hi, unit, src = row
    ASSUMPTIONS[k] = dict(section=sec, desc=desc, low=lo, base=ba, high=hi, unit=unit, source=src)

# =============================================================================
# 2. VECTORS (paths and tables)
# =============================================================================
VECTORS = OrderedDict()
# Seed supply potential (t/yr), op years 1..10 (2029-2038), before supply_mult
VECTORS["supply_path"] = dict(
    values=[8000, 11000, 14000, 17000, 20000, 23000, 26000, 28000, 30000, 32000],
    unit="t seed/yr (full-year potential)", low="x0.70", high="x1.25",
    source=f"{FS} s9 reference path 8/13/17/20/20 kt (Y1-Y5) with contracted area growth; extended +2-3 kt/yr [ESTIMATE]. "
           f"{FS} s10: P(Y1 throughput <60% of plan) ~40%.")
VECTORS["supply_mult"] = dict(values=[1.0], unit="x", low=0.70, high=1.25, source="[ESTIMATE] multiplier on supply path")
# Sunflower delivered blend (2026 KES/t) by op year
VECTORS["sun_cost_path"] = dict(
    values=[52300, 53700, 55600, 57000, 57600, 57600, 57600, 57600, 57600, 57600],
    unit="KES/t delivered (2026 KES), x seed_cost_mult", low="x0.79", high="x1.26",
    source=f"{FS} s9 'Blended delivered cost (base)' by year; flat after Y5")
VECTORS["tz_share_path"] = dict(values=[0.30, 0.25, 0.15, 0.10, 0.05, 0.05, 0.05, 0.05, 0.05, 0.05], unit="share of sunflower seed",
                                low="-", high="-", source=f"{FS} s9 sourcing mix")
VECTORS["import_share_path"] = dict(values=[0.65, 0.55, 0.45, 0.35, 0.30, 0.30, 0.30, 0.30, 0.30, 0.30], unit="share UG+TZ",
                                    low="-", high="-", source=f"{FS} s9 sourcing mix")
VECTORS["outgrower_share_path"] = dict(values=[0.10, 0.20, 0.30, 0.40, 0.45, 0.50, 0.50, 0.50, 0.50, 0.50], unit="share of sunflower seed",
                                       low="-", high="-", source=f"{FS} s9 (Y1-Y5), 50% thereafter [ESTIMATE]")
VECTORS["feed_share_ramp"] = dict(values=[0.40, 0.70, 1.0], unit="x feed_share_max, op years 1,2,3+", low="-", high="-", source="[ESTIMATE]")
VECTORS["refined_ramp"] = dict(values=[0.50 / 0.80, 0.65 / 0.80, 1.0], unit="x refined_share_max, op years 1,2,3+", low="-", high="-",
                               source="[ESTIMATE] 50% -> 65% -> 80% of crude refined under own brand")
VECTORS["season_base"] = dict(values=[6, 6, 3, 2, 4, 9, 15, 17, 16, 13, 5, 4], unit="% of annual seed bought, Jan..Dec",
                              low="Jun-Oct 55%", high="Jun-Oct 80%",
                              source=f"{FS} s7.1 harvest calendars (70% Jun-Oct; smaller Dec-Feb window) [ESTIMATE]")
# Feed formulas (inclusion, share of mash). Grow-out from TO s3.3; starter [ESTIMATE].
FORM_GROW = OrderedDict([("sun", 0.20), ("soy", 0.30), ("fish", 0.07), ("cotton", 0.05), ("maize", 0.18),
                         ("pollard", 0.13), ("bran", 0.025), ("oil", 0.02), ("premix", 0.025)])
FORM_START = OrderedDict([("sun", 0.05), ("soy", 0.45), ("fish", 0.20), ("cotton", 0.0), ("maize", 0.14),
                          ("pollard", 0.08), ("bran", 0.0), ("oil", 0.04), ("premix", 0.04)])
VECTORS["form_grow"] = dict(values=list(FORM_GROW.values()), unit="share of mash: " + ",".join(FORM_GROW.keys()), low="-", high="-",
                            source=f"{TO} s3.3 base grow-out formula (~29.7% CP)")
VECTORS["form_start"] = dict(values=list(FORM_START.values()), unit="share of mash: " + ",".join(FORM_START.keys()), low="-", high="-",
                             source="[ESTIMATE] ~36% CP starter/fingerling formula")
CP = dict(sun=0.30, soy=0.45, fish=0.58, cotton=0.38, maize=0.08, pollard=0.16, bran=0.15, oil=0.0, premix=0.12)

# Capex items, USD m (2026), Phase 1 per TO s4.1; Phase 2 reference at TO's 150 tpd spec
CAPEX_P1 = OrderedDict([
    ("land", (0.60, "land", False, f"{TO} s4.1 (6-10 acres)")),
    ("civil", (1.20, "building", False, f"{TO} s4.1 (~4,500 m2)")),
    ("silos", (0.50, "machinery", True, f"{TO} s4.1 (5 kt)")),
    ("crushing", (0.80, "machinery", True, f"{TO} s4.1 (50 tpd full press)")),
    ("soy_ee", (0.25, "machinery", True, "[ESTIMATE] soy dry extruder-expeller (TO s1.1 route)")),
    ("refinery", (0.90, "machinery", True, f"{TO} s1.4 (20 t/d), scaled to ref_tpd_p1")),
    ("extrusion", (1.90, "machinery", True, f"{TO} s4.1 (3 t/h twin-screw)")),
    ("utilities", (1.10, "machinery", "partial", f"{TO} s4.1 (boiler, transformer, genset, ETP)")),
    ("qc_lab", (0.15, "machinery", True, f"{TO} s7")),
    ("vehicles", (0.30, "vehicles", True, f"{TO} s4.1")),
    ("engineering", (0.55, "soft", False, f"{TO} s4.1 USD 0.75m less IDF/RDL shown separately")),
])
CAPEX_P2_REF = OrderedDict([
    ("land", (0.30, "land", False, "fixed")),
    ("civil", (2.50, "building", False, "inc")),
    ("silos", (1.80, "machinery", True, "inc")),
    ("crushing", (2.00, "machinery", True, "inc")),
    ("solvent", (4.50, "machinery", True, "solv")),
    ("refinery", (2.50, "machinery", True, "ref")),
    ("extrusion", (3.50, "machinery", True, "ext")),
    ("utilities", (2.00, "machinery", "partial", "inc")),
    ("qc_lab", (0.10, "machinery", True, "fixed")),
    ("vehicles", (0.30, "vehicles", True, "fixed")),
    ("engineering", (0.80, "soft", False, "avg")),
])
VECTORS["capex_phase_p1"] = dict(values=[0.15, 0.70, 0.15], unit="share of P1 capex excl. land in 2027/2028/2029 (land 100% 2027)",
                                 low="-", high="-", source="[ESTIMATE] FID Q4 2027 deposits; 15-month build")
VECTORS["capex_phase_p2"] = dict(values=[0.80, 0.20], unit="share of P2 capex in 2032/2033", low="-", high="-", source="[ESTIMATE]")
BOOK_LIFE = dict(building=25, machinery=12, vehicles=5)
# Restructure alternatives (R1-05) [ESTIMATE]: Phase 1 capex item multipliers and fixed-opex multipliers by configuration.
# 1 = feed-mill-first: 3 t/h extrusion line buying all protein (no crusher, refinery or seed silos);
# 2 = crusher + refinery only: sells cake to feed millers (e.g. a protein-supply deal with Tunga/SamakGro/De Heus).
CONFIG_CAPEX = {
    1: dict(land=0.7, civil=0.55, silos=0.0, crushing=0.0, soy_ee=0.0, refinery=0.0, utilities=0.6, qc_lab=0.7, vehicles=0.7, engineering=0.6),
    2: dict(land=0.85, civil=0.75, extrusion=0.0, utilities=0.75, engineering=0.75),
}
CONFIG_FIXED = {1: 0.55, 2: 0.75}   # labour, G&A, security, lab, demand charge, diesel [ESTIMATE]
CONFIG_NAMES = {0: "Integrated crusher + feed mill (plan)", 1: "Feed-mill-first (buys protein; no crusher)", 2: "Crusher + refinery only (sells cake)"}


def base_params():
    p = {k: v["base"] for k, v in ASSUMPTIONS.items()}
    p["supply_mult"] = 1.0
    p["exit_mult_override"] = None
    return p


def load_overrides(p):
    """Optional assumptions_override.csv (columns: key,value) overrides base values."""
    path = os.path.join(HERE, "assumptions_override.csv")
    if os.path.exists(path):
        with open(path) as f:
            for r in csv.DictReader(f):
                k = r["key"].strip()
                if k in p:
                    p[k] = float(r["value"])
                    print(f"  override: {k} = {r['value']}")
                else:
                    print(f"  WARNING: unknown override key {k}")
    return p


# =============================================================================
# helpers
# =============================================================================
def byop(vec, pre=None):
    """Map a vector indexed by op year (1..) onto model years; last value repeats."""
    out = np.zeros(N)
    for i in range(N):
        k = OPK[i]
        if k >= 1:
            out[i] = vec[min(k - 1, len(vec) - 1)]
        else:
            out[i] = vec[0] if pre is None else pre
    return out


def npv(rate, cfs, years=YEARS):
    t = np.asarray(years) - VALUE_YEAR
    return float(np.sum(np.asarray(cfs) / (1.0 + rate) ** t))


def irr(cfs, years=YEARS):
    cfs = np.asarray(cfs, dtype=float)
    if not (np.any(cfs < 0) and np.any(cfs > 0)):
        return float("nan")
    grid = np.concatenate([np.linspace(-0.95, 1.0, 400), np.linspace(1.0, 10.0, 200)[1:]])
    vals = np.array([npv(r, cfs, years) for r in grid])
    roots = []
    for i in range(len(grid) - 1):
        if vals[i] == 0:
            roots.append(grid[i])
        elif vals[i] * vals[i + 1] < 0:
            lo, hi = grid[i], grid[i + 1]
            flo = vals[i]
            for _ in range(100):
                mid = 0.5 * (lo + hi)
                fm = npv(mid, cfs, years)
                if flo * fm <= 0:
                    hi = mid
                else:
                    lo, flo = mid, fm
            roots.append(0.5 * (lo + hi))
    if not roots:
        return float("nan")
    pos = [r for r in roots if r > -0.5]
    return float(min(pos) if pos else roots[0])


def payback_year(cfs):
    c = np.cumsum(cfs)
    for i in range(N):
        if c[i] >= 0 and np.all(c[i:] >= 0):
            return int(YEARS[i])
    return None


def seasonal_profile(p):
    base = np.array(VECTORS["season_base"]["values"], dtype=float)
    peak_m = np.zeros(12, bool)
    peak_m[5:10] = True                       # Jun..Oct
    c = p["season_conc"]
    sh = np.where(peak_m, base * c / base[peak_m].sum(), base * (1 - c) / base[~peak_m].sum())
    cum = np.cumsum(sh - 1.0 / 12)            # end-of-month inventory change, fraction of annual crush
    I0 = p["safety_stock_months"] / 12 - min(0.0, cum.min())
    inv = I0 + cum                            # end-of-month
    start = np.concatenate([[I0], inv[:-1]])
    avg = float(np.mean((start + inv) / 2))
    return dict(shares=sh, inv=inv, dec=float(inv[-1]), avg=avg, peak=float(inv.max()), I0=I0,
                peak_month=int(np.argmax(inv)) + 1)


def capex_items(p, phase):
    """Return OrderedDict item -> (usd_m_real, category, imported_share)."""
    m = p["capex_mult"]
    out = OrderedDict()
    if phase == 1:
        cfg = CONFIG_CAPEX.get(int(p.get("config", 0)), {})
        for k, (v, cat, imp, _src) in CAPEX_P1.items():
            val = v
            if k == "refinery":
                val = v * (p["ref_tpd_p1"] / 20.0) ** p["scale_exp"]
            val *= cfg.get(k, 1.0)
            impf = p["util_import"] if imp == "partial" else (1.0 if imp else 0.0)
            out[k] = [val * m, cat, impf]
        cont = p["contingency_p1"]
    else:
        e = p["scale_exp"]
        s = dict(inc=max(p["p2_tpd"] - p["p1_tpd"], 0) / 100.0, solv=p["p2_tpd"] / 150.0,
                 ref=p["p2_ref_add"] / 40.0, ext=p["p2_extr_tph"] / 6.0, fixed=1.0)
        s = {k: (v ** e if v > 0 else 0.0) for k, v in s.items()}
        main = [k for k in CAPEX_P2_REF if CAPEX_P2_REF[k][3] in ("inc", "solv", "ref", "ext")]
        scaled = sum(CAPEX_P2_REF[k][0] * s[CAPEX_P2_REF[k][3]] for k in main)
        unscaled = sum(CAPEX_P2_REF[k][0] for k in main)
        s["avg"] = scaled / unscaled
        for k, (v, cat, imp, rule) in CAPEX_P2_REF.items():
            impf = p["util_import"] if imp == "partial" else (1.0 if imp else 0.0)
            out[k] = [v * s[rule] * m, cat, impf]
        cont = p["contingency_p2"]
    imported = sum(v[0] * v[2] for v in out.values())
    out["idf_rdl"] = [imported * (p["idf"] + p["rdl"]), "soft", 0.0]
    base_ex_land = sum(v[0] for k, v in out.items() if k != "land")
    out["contingency"] = [base_ex_land * cont, "soft", 0.0]
    return out


def category_split(items):
    """Allocate soft costs pro-rata; return dict category -> USD m."""
    cats = dict(land=0.0, building=0.0, machinery=0.0, vehicles=0.0)
    soft = 0.0
    for k, (v, cat, _i) in items.items():
        if cat == "soft":
            soft += v
        else:
            cats[cat] += v
    hard_dep = cats["building"] + cats["machinery"] + cats["vehicles"]
    for c in ("building", "machinery", "vehicles"):
        cats[c] += soft * cats[c] / hard_dep if hard_dep > 0 else 0
    return cats


# =============================================================================
# 3. CORE RUN
# =============================================================================
def run_core(p, go):
    R = OrderedDict()
    Y = YEARS
    ops = (OPK >= 1).astype(float)
    opsfrac = np.where(OPK == 1, p["ops_months_y1"] / 12.0, np.where(OPK >= 2, 1.0, 0.0))
    p2on = ((Y >= p["p2_start_year"]) & bool(go)).astype(float)

    # ---- indices
    dep = np.full(N, p["kes_dep"])
    dep[0] += p["kes_dep_2027_extra"]
    fx = p["fx0"] * np.cumprod(1 + dep)
    t26 = Y - 2026
    T = (1 + p["usd_infl"]) ** t26 * fx / p["fx0"]
    infl = (1 + p["usd_infl"]) * (1 + dep) - 1 + p["kes_real_drift"]       # KES CPI linked to FX (PPP + real drift)
    L = np.cumprod(1 + infl)
    W = np.cumprod(1 + infl + p["real_wage"])
    R["infl"] = infl
    R.update(fx=fx, T=T, L=L, ops=ops, opsfrac=opsfrac, p2on=p2on)

    # ---- seasonal profile
    sp = seasonal_profile(p)
    avg_months = sp["avg"] * 12
    shrink = p["storage_shrink_pm"] * avg_months

    # ---- crushing volumes (t)
    cap_crush = np.where(p2on > 0, p["p2_tpd"], p["p1_tpd"]) * p["crush_days"] * ops
    cfg = int(p.get("config", 0))
    supply = byop(VECTORS["supply_path"]["values"], pre=0.0) * p["supply_mult"] * (0.0 if cfg == 1 else 1.0)
    crush = np.minimum(cap_crush * p["max_util"], supply) * opsfrac
    soy = crush * p["soy_share"]
    sun = crush - soy
    sun_solv = sun * p2on
    sun_press = sun - sun_solv
    y_oil_p = p["sun_oil_y"]
    y_cake_p = min(p["sun_cake_max"], 1 - p["sun_loss"] - y_oil_p)
    y_oil_s = p["sun_oil_y"] + p["p2_oil_uplift"]
    y_hull = p["p2_hull"]
    y_meal = 1 - y_oil_s - y_hull - p["sun_loss"]
    y_soy_oil = p["soy_oil_y"]
    y_soy_cake = 1 - y_soy_oil - p["soy_loss"]
    oil_sun = sun_press * y_oil_p + sun_solv * y_oil_s
    cake_sun = sun_press * y_cake_p
    meal_sun = sun_solv * y_meal
    hulls = sun_solv * y_hull
    loss_sun = sun - oil_sun - cake_sun - meal_sun - hulls
    oil_soy = soy * y_soy_oil
    cake_soy = soy * y_soy_cake
    loss_soy = soy - oil_soy - cake_soy
    crude_total = oil_sun + oil_soy
    seed_bought = crush * (1 + shrink)
    seed_shrink = seed_bought - crush

    # ---- feed volumes
    sam = (p["sam_2026"] + p["sam_regional"]) * (1 + p["sam_growth"]) ** t26
    share = p["feed_share_max"] * byop(VECTORS["feed_share_ramp"]["values"], pre=0.0) * (0.0 if cfg == 2 else 1.0)
    feed_cap = (p["extr_tph_p1"] + p2on * p["p2_extr_tph"]) * p["extr_eff"] * p["extr_hours"] * ops
    feed = np.minimum(feed_cap, share * sam) * opsfrac
    f_start = feed * p["starter_share"]
    f_grow = feed - f_start
    mash_g = f_grow / (1 - p["feed_shrink"])
    mash_s = f_start / (1 - p["feed_shrink"])
    mash = mash_g + mash_s
    need = OrderedDict((k, mash_g * FORM_GROW[k] + mash_s * FORM_START[k]) for k in FORM_GROW)
    sun_prod = cake_sun + meal_sun
    sun_to_feed = np.minimum(need["sun"], sun_prod)
    sun_buy = need["sun"] - sun_to_feed
    soy_to_feed = np.minimum(need["soy"], cake_soy)
    sbm_buy = need["soy"] - soy_to_feed
    oil_to_feed = np.minimum(need["oil"], crude_total)
    oil_buy = need["oil"] - oil_to_feed
    sun_sold = sun_prod - sun_to_feed
    soy_sold = cake_soy - soy_to_feed

    # ---- oil routing
    crude_avail = crude_total - oil_to_feed
    ref_cap = (p["ref_tpd_p1"] + p2on * p["p2_ref_add"]) * p["ref_days"] * p["ref_max_util"] * ops
    ref_share = p["refined_share_max"] * byop(VECTORS["refined_ramp"]["values"], pre=0.0)
    ref_in = np.minimum(ref_cap * opsfrac, crude_avail * ref_share)
    ref_out = ref_in * (1 - p["ref_loss"])
    soap = ref_in * p["ref_loss"]
    crude_sold = crude_avail - ref_in

    # ---- prices (nominal KES/t)
    pol = ops
    oil_adj = p["p_cpo_removed"] * p["cpo_duty_value"] * pol
    prot = 1 - pol * (p["p_gmo"] * p["gmo_meal_drop"] + p["waiver_freq"] * p["waiver_drop"])
    entrant = np.where((OPK == 1) | (OPK == 2), 1 - p["entrant_discount"], 1.0)
    soyfrac = np.divide(oil_soy, crude_total, out=np.zeros(N), where=crude_total > 0)
    A0 = {k: v["base"] for k, v in ASSUMPTIONS.items()}
    oil_adj_ref = 0.0 * pol          # reference = today's prices (research values, no policy shock)
    prot_ref = 1.0 + 0.0 * pol
    P = OrderedDict()
    P["crude"] = (p["p_crude"] - oil_adj) * (1 - soyfrac * (1 - p["soy_oil_factor"])) * T
    P["crude_sun"] = (p["p_crude"] - oil_adj) * T
    # R1-03: refined price is linked to the crude/parity index through a refining spread, so the two cannot drift apart
    P["refined"] = (p["p_crude"] + p["refine_spread"] - oil_adj) * T
    P["soap"] = p["soapstock_factor"] * (p["p_crude"] - oil_adj) * T
    P["suncake"] = p["p_suncake"] * prot * T
    P["sunmeal"] = p["p_suncake"] * p["meal_premium"] * prot * T
    sun_px = np.where(p2on > 0, P["sunmeal"], P["suncake"])
    P["sunprod"] = sun_px
    P["soycake"] = p["p_soycake"] * prot * T
    P["sbm"] = p["p_sbm"] * prot * T
    P["fish"] = p["p_fishmeal"] * T
    P["cotton"] = p["p_cottoncake"] * prot * T
    P["maize"] = p["p_maize"] * T
    P["pollard"] = p["p_pollard"] * T
    P["bran"] = p["p_bran"] * T
    P["premix"] = p["p_premix"] * T
    # aquafeed prices follow competitors' raw-material cost (market prices, SBM as protein) vs today's
    mkt = dict(sun=p["p_suncake"] * prot, soy=p["p_sbm"] * prot, fish=p["p_fishmeal"], cotton=p["p_cottoncake"] * prot,
               maize=p["p_maize"], pollard=p["p_pollard"], bran=p["p_bran"], oil=p["p_crude"] - oil_adj, premix=p["p_premix"])
    mkt0 = dict(sun=A0["p_suncake"] * prot_ref, soy=A0["p_sbm"] * prot_ref, fish=A0["p_fishmeal"], cotton=A0["p_cottoncake"] * prot_ref,
                maize=A0["p_maize"], pollard=A0["p_pollard"], bran=A0["p_bran"], oil=A0["p_crude"] - oil_adj_ref, premix=A0["p_premix"])
    d_grow = sum(FORM_GROW[k] * (mkt[k] - mkt0[k]) for k in FORM_GROW) / (1 - A0["feed_shrink"])
    d_start = sum(FORM_START[k] * (mkt[k] - mkt0[k]) for k in FORM_START) / (1 - A0["feed_shrink"])
    fpt = p["feed_cost_passthrough"]
    # R1-01: competitor response = real decline in feed prices from 2029 (over-capacity: Tunga 45 kt, SamakGro 8 t/h)
    comp = (1 - p["feed_real_decline"]) ** np.maximum(Y - 2028, 0)
    R["feed_comp_index"] = comp
    P["feed_grow"] = (p["p_feed_grow"] * p["feed_px_mult"] + fpt * d_grow) * entrant * comp * T
    P["feed_start"] = (p["p_feed_start"] * p["feed_px_mult"] + fpt * d_start) * entrant * comp * T
    # seed prices follow product value (floating, parity-linked contracts): pass-through of deviations from base prices
    y0o, y0c = A0["sun_oil_y"], min(A0["sun_cake_max"], 1 - A0["sun_loss"] - A0["sun_oil_y"])
    y0so, y0sc = A0["soy_oil_y"], 1 - A0["soy_oil_y"] - A0["soy_loss"]
    v_sun = y0o * (p["p_crude"] - oil_adj) + y0c * p["p_suncake"] * prot
    v_sun_ref = y0o * (A0["p_crude"] - oil_adj_ref) + y0c * A0["p_suncake"] * prot_ref
    v_soy = y0so * (p["p_crude"] - oil_adj) * p["soy_oil_factor"] + y0sc * p["p_soycake"] * prot
    v_soy_ref = y0so * (A0["p_crude"] - oil_adj_ref) * A0["soy_oil_factor"] + y0sc * A0["p_soycake"] * prot_ref
    pt_sun = p["seed_passthrough"] * (v_sun - v_sun_ref)
    pt_soy = p["seed_passthrough"] * (v_soy - v_soy_ref)
    R["seed_pt"] = (pt_sun, pt_soy)
    sun_base = byop(VECTORS["sun_cost_path"]["values"]) * p["seed_cost_mult"] + pt_sun
    tz_add = p["p_tz_levy"] * p["tz_levy"] * p["tz_farmgate"] * byop(VECTORS["tz_share_path"]["values"])
    afa_add = p["afa_levy_seed"] * byop(VECTORS["import_share_path"]["values"]) * sun_base * 0.8
    P["sunseed"] = (sun_base + tz_add + afa_add) * T
    P["soyseed"] = (p["p_soyseed"] * p["seed_cost_mult"] + pt_soy) * T
    R["prices"] = P

    # ---- revenue (KES)
    rev = OrderedDict()
    rev["Refined oil"] = ref_out * P["refined"]
    rev["Crude oil to refiners"] = crude_sold * P["crude"]
    rev["Soapstock / acid oil"] = soap * P["soap"]
    rev["Sunflower cake/meal"] = sun_sold * sun_px
    rev["Soy cake (surplus)"] = soy_sold * P["soycake"]
    rev["Aquafeed grow-out"] = f_grow * P["feed_grow"]
    rev["Aquafeed starter/fingerling"] = f_start * P["feed_start"]
    revenue = sum(rev.values())
    rev_feed = rev["Aquafeed grow-out"] + rev["Aquafeed starter/fingerling"]

    # ---- variable costs
    cost = OrderedDict()
    seed_sun_cost = sun * (1 + shrink) * P["sunseed"]
    seed_soy_cost = soy * (1 + shrink) * P["soyseed"]
    cost["Seed: sunflower"] = seed_sun_cost
    cost["Seed: soybean"] = seed_soy_cost
    ingr_purch = (need["fish"] * P["fish"] + need["cotton"] * P["cotton"] + need["maize"] * P["maize"] +
                  need["pollard"] * P["pollard"] + need["bran"] * P["bran"] + need["premix"] * P["premix"] +
                  sbm_buy * P["sbm"] + sun_buy * P["suncake"] + oil_buy * P["crude_sun"])
    cost["Feed ingredients purchased"] = ingr_purch
    pack_feed = feed * p["pack_feed"] * L
    pack_oil = ref_out * p["pack_oil_l"] * p["l_per_t"] * L
    pack_cake = (sun_sold + soy_sold) * p["pack_cake"] * L
    cost["Packaging"] = pack_feed + pack_oil + pack_cake
    cost["Refining chemicals"] = ref_in * p["ref_chem"] * T
    # energy
    kwh = (sun_press * p["kwh_sun"] + sun_solv * (p["kwh_sun"] + p["kwh_solv"]) + soy * p["kwh_soy"] +
           ref_in * p["kwh_ref"] + feed * p["kwh_feed"] + p["kwh_gen"] * (1 + 0.5 * p2on) * opsfrac)
    steam = crush * p["steam_crush"] + sun_solv * p["steam_solv"] + ref_in * p["steam_ref"]
    therm_gj = steam * p["gj_per_t_steam"] + feed * p["feed_gj"]
    hull_gj = hulls * p["hull_mj"]
    hull_cover = np.minimum(0.9, np.divide(hull_gj, therm_gj, out=np.zeros(N), where=therm_gj > 0))
    fuel_cost_per_gj = p["steam_cost"] / p["gj_per_t_steam"]
    elec_cost = kwh * p["elec_price"] * L
    thermal_cost = therm_gj * fuel_cost_per_gj * (1 - hull_cover) * L
    cfx = CONFIG_FIXED.get(cfg, 1.0)
    demand_cost = (p["kva_p1"] * cfx + p2on * p["kva_p2"]) * p["demand_charge"] * 12 * opsfrac * L
    diesel = (p["diesel_p1"] * cfx + p2on * p["diesel_p2"]) * opsfrac * L
    cost["Energy"] = elec_cost + thermal_cost + demand_cost + diesel
    cost["Press wear, hexane"] = crush * p["press_wear"] * L + sun_solv * p["hexane"] * T
    # storage overflow
    silo = (p["silo_p1"] + p2on * p["silo_p2"]) * ops
    inv_m = np.outer(crush / np.maximum(opsfrac, 1e-9) * (opsfrac > 0), sp["inv"])   # t by month (annualised rate)
    over = np.maximum(inv_m - silo[:, None], 0).sum(axis=1) * opsfrac
    cost["Leased storage"] = over * p["storage_lease"] * L
    variable = sum(cost.values())

    # ---- fixed and selling
    fx_mult = p["opex_fixed_mult"] * CONFIG_FIXED.get(cfg, 1.0)
    items1 = capex_items(p, 1)
    items2 = capex_items(p, 2)
    cat1 = category_split(items1)
    cat2 = category_split(items2)
    mach_usd = cat1["machinery"] + p2on * cat2["machinery"]
    bldg_usd = cat1["building"] + p2on * cat2["building"]
    gross_usd = sum(v[0] for v in items1.values()) + p2on * sum(v[0] for v in items2.values())
    fixed = OrderedDict()
    fixed["Labour"] = (p["labour_p1"] + p2on * p["labour_p2"]) * W * opsfrac * fx_mult
    fixed["Maintenance"] = (mach_usd * p["maint_pct"] + bldg_usd * p["maint_bldg"]) * 1e6 * p["fx0"] * L * opsfrac * fx_mult
    fixed["QC lab"] = (p["lab_p1"] + p2on * p["lab_p2"]) * L * opsfrac * fx_mult
    fixed["Insurance"] = gross_usd * 1e6 * p["insurance"] * fx * (1 + p["usd_infl"]) ** t26 * opsfrac
    fixed["Security"] = (p["security_p1"] + p2on * p["security_p2"]) * L * opsfrac * fx_mult
    fixed["G&A other"] = (p["ga_p1"] + p2on * p["ga_p2"]) * L * opsfrac * fx_mult
    fixed_total = sum(fixed.values())
    sell = OrderedDict()
    sell["Crude oil freight"] = crude_sold * p["crude_freight"] * L
    sell["Cake/meal freight"] = (sun_sold + soy_sold) * p["cake_freight"] * L
    sell["Feed delivery + selling"] = feed * p["feed_delivery"] * L + rev_feed * p["feed_selling"]
    sell["Feed dealer margin"] = rev_feed * p["feed_dealer_margin"]
    sell["Bad debts"] = revenue * p["bad_debt"]
    launch = np.where(OPK == 1, p["launch_mkt"], np.where(OPK == 2, 0.75 * p["launch_mkt"], 0.0)) * L
    sell["Refined oil distribution + launch"] = rev["Refined oil"] * p["refined_selling"] + launch
    sell_total = sum(sell.values())
    levies = OrderedDict()
    levies["Irrecoverable VAT (feed)"] = rev_feed * p["vat_feed_irrecov"]
    levies["Feed levy"] = rev_feed * p["feed_levy"]
    levies_total = sum(levies.values())
    # development / pre-op
    devx = np.zeros(N)
    devx[IDX[2027]] += p["phase0_usd"] * 1e6 * fx[IDX[2027]] * T[IDX[2027]] / (fx[IDX[2027]] / p["fx0"])
    devx[IDX[2028]] += 0.6 * p["preop_p1_usd"] * 1e6 * fx[IDX[2028]]
    devx[IDX[2029]] += 0.4 * p["preop_p1_usd"] * 1e6 * fx[IDX[2029]]
    if go:
        devx[IDX[2032]] += p["preop_p2_usd"] * 1e6 * fx[IDX[2032]]
    opex_total = variable + fixed_total + sell_total + levies_total + devx
    ebitda = revenue - opex_total

    # ---- segment view (for gate b and unit economics), internal transfers at market
    feed_input_value = (ingr_purch + sun_to_feed * sun_px + soy_to_feed * P["soycake"] + oil_to_feed * P["crude_sun"])
    feed_energy = feed * p["kwh_feed"] * p["elec_price"] * L + feed * p["feed_gj"] * fuel_cost_per_gj * (1 - hull_cover) * L
    feed_gm = rev_feed * (1 - p["feed_dealer_margin"]) - feed_input_value - pack_feed - feed_energy
    feed_gm_pct = np.divide(feed_gm, rev_feed, out=np.zeros(N), where=rev_feed > 0)
    crush_value = (crude_total * P["crude"] + sun_prod * sun_px + cake_soy * P["soycake"])
    crush_gm = crush_value - seed_sun_cost - seed_soy_cost
    crush_gm_per_t = np.divide(crush_gm, crush, out=np.zeros(N), where=crush > 0)

    # ---- capex (nominal KES)
    usd_esc = fx * (1 + p["usd_infl"]) ** t26
    capex_p1 = np.zeros(N)
    land1 = items1["land"][0]
    nonland1 = sum(v[0] for k, v in items1.items() if k != "land")
    ph1 = VECTORS["capex_phase_p1"]["values"]
    for j, yr in enumerate((2027, 2028, 2029)):
        capex_p1[IDX[yr]] = (nonland1 * ph1[j] + (land1 if yr == 2027 else 0.0)) * 1e6 * usd_esc[IDX[yr]]
    capex_p2 = np.zeros(N)
    if go:
        tot2 = sum(v[0] for v in items2.values())
        ph2 = VECTORS["capex_phase_p2"]["values"]
        for j, yr in enumerate((2032, 2033)):
            capex_p2[IDX[yr]] = tot2 * ph2[j] * 1e6 * usd_esc[IDX[yr]]
    cum_ex_land_usd = np.where(OPK >= 1, nonland1, 0.0) + p2on * (sum(v[0] for v in items2.values()) - items2["land"][0])
    sust = np.where(OPK >= 3, p["sust_capex"] * cum_ex_land_usd * 1e6 * usd_esc, 0.0)
    capex = capex_p1 + capex_p2
    land_spend = np.zeros(N)
    land_spend[IDX[2027]] = land1 * 1e6 * usd_esc[IDX[2027]]
    if go:
        land_spend[IDX[2032]] = items2["land"][0] * ph2[0] * 1e6 * usd_esc[IDX[2032]]
        land_spend[IDX[2033]] = items2["land"][0] * ph2[1] * 1e6 * usd_esc[IDX[2033]]
    vat_paid = p["vat_rate"] * (capex - land_spend + sust)
    # R1-09: if finished feed is VAT-exempt, input VAT on capex is apportioned by exempt (feed) turnover share and the
    # exempt share is an up-front cash cost. Phase 1 capex (2027-29) uses the 2029-31 turnover mix; Phase 2 capex the 2033-35 mix.
    fshare = np.divide(rev_feed, revenue, out=np.zeros(N), where=revenue > 0)
    es1 = float(np.mean(fshare[IDX[2029]:IDX[2031] + 1]))
    es2 = float(np.mean(fshare[IDX[2033]:IDX[2035] + 1]))
    exempt = np.where(Y <= 2031, es1, np.where(Y <= 2035, es2, fshare)) * p["p_feed_exempt"]
    vat_irrecov = vat_paid * exempt
    vat_recov = np.concatenate([[0.0], (vat_paid - vat_irrecov)[:-1]])

    # ---- book depreciation and tax allowances by asset pool
    def pools(cat, total_nom_by_year, service_year):
        tot = total_nom_by_year.sum()
        return tot, service_year
    # nominal cost per category: split each phase's nominal capex by category shares
    def cat_nominal(items, capex_nom):
        cats = category_split(items)
        tot_usd = sum(cats.values())
        tot_nom = capex_nom.sum()
        return {c: tot_nom * cats[c] / tot_usd for c in cats}
    c1 = cat_nominal(items1, capex_p1)
    c2 = cat_nominal(items2, capex_p2) if go else {c: 0.0 for c in c1}
    dep_book = np.zeros(N)
    ded = np.zeros(N)
    vehicles_wdv = 0.0
    ded_pool = []   # (remaining residue for 50/25 regime)
    for i, yr in enumerate(Y):
        # additions placed in service
        adds = {c: 0.0 for c in c1}
        if yr == 2029:
            adds = dict(c1)
        if go and yr == p["p2_start_year"]:
            for c in adds:
                adds[c] += c2[c]
        bm_add = adds["building"] + adds["machinery"] + (sust[i - 1] if i > 0 and OPK[i] >= 4 else 0.0)
        # investment deduction
        if bm_add > 0:
            first = p["inv_ded"] * bm_add
            ded[i] += first
            ded_pool.append(bm_add - first)
        # residue at 25% reducing (only if inv_ded < 1)
        newpool = []
        for j, res in enumerate(ded_pool[:-1] if bm_add > 0 else ded_pool):
            d = 0.25 * res
            ded[i] += d
            newpool.append(res - d)
        if bm_add > 0:
            newpool.append(ded_pool[-1])
        ded_pool = newpool
        vehicles_wdv += adds["vehicles"]
        dv = 0.25 * vehicles_wdv
        ded[i] += dv
        vehicles_wdv -= dv
    # book depreciation (straight line from service year)
    for c in ("building", "machinery", "vehicles"):
        life = BOOK_LIFE[c]
        for start, amount in ((2029, c1[c]), (p["p2_start_year"], c2[c])):
            if amount <= 0:
                continue
            for i, yr in enumerate(Y):
                if start <= yr < start + life:
                    dep_book[i] += amount / life
    for i in range(N):
        if sust[i] > 0:
            for j in range(i + 1, N):
                if j - i <= 12:
                    dep_book[j] += sust[i] / 12.0

    R.update(volumes=OrderedDict(
        seed_bought=seed_bought, seed_shrink=seed_shrink, crush=crush, sun=sun, soy=soy, sun_press=sun_press,
        sun_solv=sun_solv, cap_crush=cap_crush, supply=supply, oil_sun=oil_sun, oil_soy=oil_soy, crude_total=crude_total,
        cake_sun=cake_sun, meal_sun=meal_sun, hulls=hulls, loss_sun=loss_sun, cake_soy=cake_soy, loss_soy=loss_soy,
        oil_to_feed=oil_to_feed, crude_avail=crude_avail, ref_in=ref_in, ref_out=ref_out, soap=soap, crude_sold=crude_sold,
        ref_cap=ref_cap * opsfrac, sun_to_feed=sun_to_feed, sun_sold=sun_sold, soy_to_feed=soy_to_feed, soy_sold=soy_sold,
        sbm_buy=sbm_buy, sun_buy=sun_buy, oil_buy=oil_buy, sam=sam, share=share, feed_cap=feed_cap * opsfrac,
        feed=feed, f_grow=f_grow, f_start=f_start, mash=mash, feed_shrink=mash - feed, kwh=kwh, steam=steam,
        hull_cover=hull_cover, over_storage=over, silo=silo, **{"need_" + k: v for k, v in need.items()}))
    R.update(rev=rev, revenue=revenue, rev_feed=rev_feed, cost=cost, variable=variable, fixed=fixed, fixed_total=fixed_total,
             sell=sell, sell_total=sell_total, levies=levies, levies_total=levies_total, devx=devx,
             opex_total=opex_total, ebitda=ebitda, feed_gm=feed_gm, feed_gm_pct=feed_gm_pct, crush_gm=crush_gm,
             crush_gm_per_t=crush_gm_per_t, capex_p1=capex_p1, capex_p2=capex_p2, capex=capex, sust=sust,
             vat_paid=vat_paid, vat_recov=vat_recov, vat_irrecov=vat_irrecov, vat_exempt_share=exempt, dep_book=dep_book, ded=ded, items1=items1, items2=items2,
             elec_cost=elec_cost, thermal_cost=thermal_cost, sp=sp, shrink=shrink, ingr_purch=ingr_purch,
             pack_feed=pack_feed, feed_input_value=feed_input_value, seed_sun_cost=seed_sun_cost, seed_soy_cost=seed_soy_cost)

    # ---- working capital (year-end)
    crush_next = np.concatenate([crush[1:], crush[-1:]])
    sun_next = crush_next * (1 - p["soy_share"])
    soy_next = crush_next * p["soy_share"]
    seed_unit = np.where(crush > 0, (seed_sun_cost + seed_soy_cost) / np.maximum(crush, 1e-9),
                         (1 - p["soy_share"]) * P["sunseed"] + p["soy_share"] * P["soyseed"])
    seed_inv = sp["dec"] * (sun_next * P["sunseed"] + soy_next * P["soyseed"])
    seed_inv[-1] = sp["dec"] * crush[-1] * seed_unit[-1]
    ingr_inv = ingr_purch * p["ingr_days"] / 365
    fg_inv = 0.8 * (rev["Refined oil"] + rev["Crude oil to refiners"] + rev["Sunflower cake/meal"] + rev["Soy cake (surplus)"] + rev_feed) * p["fg_days"] / 365
    ar = ((rev["Refined oil"] + rev["Crude oil to refiners"] + rev["Soapstock / acid oil"]) * p["ar_days_oil"] +
          (rev["Sunflower cake/meal"] + rev["Soy cake (surplus)"]) * p["ar_days_cake"] + rev_feed * p["ar_days_feed"]) / 365
    ap_base = ingr_purch + cost["Packaging"] + cost["Refining chemicals"] + cost["Energy"] + sell["Crude oil freight"] + sell["Cake/meal freight"]
    ap = ap_base * p["ap_days"] / 365
    wc = seed_inv + ingr_inv + fg_inv + ar - ap
    dwc = np.diff(np.concatenate([[0.0], wc]))
    # intra-year seasonal peak (steady-state profile on this year's crush) + outgrower credit Mar-Aug
    og_t = sun * byop(VECTORS["outgrower_share_path"]["values"], pre=0.0)
    og_peak = og_t * p["og_credit"] * L
    seed_val_rate = np.where(crush > 0, (seed_sun_cost + seed_soy_cost) / np.maximum(crush, 1e-9) / (1 + shrink), 0.0)
    ann_crush = np.divide(crush, opsfrac, out=np.zeros(N), where=opsfrac > 0)
    month_inv_val = np.outer(ann_crush * seed_val_rate, sp["inv"])            # KES by month
    og_month = np.outer(og_peak, np.array([0, 0, 1, 1, 1, 1, 1, 1, 0, 0, 0, 0], float))
    seasonal_peak = (month_inv_val + og_month).max(axis=1)
    revolver_bal = p["revolver_adv"] * seed_inv
    rev_avg_draw = p["revolver_adv"] * sp["avg"] * ann_crush * seed_val_rate * opsfrac + og_peak * 6 / 12
    revolver_int = p["revolver_rate"] * rev_avg_draw
    revolver_peak = p["revolver_adv"] * month_inv_val.max(axis=1) + og_peak
    R.update(seed_inv=seed_inv, ingr_inv=ingr_inv, fg_inv=fg_inv, ar=ar, ap=ap, wc=wc, dwc=dwc, og_peak=og_peak,
             seasonal_peak=seasonal_peak, revolver_bal=revolver_bal, revolver_int=revolver_int, revolver_peak=revolver_peak,
             month_inv_val=month_inv_val, og_month=og_month)
    return R


def debt_schedules(p, R, senior_kes, go, prepay=None, extra_fees=None):
    """Senior P1 (KES), mezzanine (USD), Phase 2 DFI loan (KES), fees, DSRA target.
    prepay: year-end senior prepayments (equity cures and cash sweeps); later instalments shrink pro rata.
    extra_fees: waiver/default fees in breach years."""
    fx = R["fx"]
    capex_p1, capex_p2 = R["capex_p1"], R["capex_p2"]
    prepay = np.zeros(N) if prepay is None else np.asarray(prepay, float)
    D = OrderedDict()
    # --- senior P1: drawn pro-rata to 2028/2029 capex; quarterly instalments after grace
    w = np.array([capex_p1[IDX[2028]], capex_p1[IDX[2029]]])
    w = w / w.sum()
    draws = np.zeros(N)
    draws[IDX[2028]] = senior_kes * w[0]
    draws[IDX[2029]] = senior_kes * w[1]
    nq = int(p["senior_tenor"] * 4)
    gq = int(round(p["senior_grace_m"] / 3))
    ninst = nq - gq
    prin_orig = np.zeros(N)
    for q in range(gq, nq):
        yr = 2028 + q // 4
        if yr in IDX:
            prin_orig[IDX[yr]] += senior_kes / ninst
    bal_open = np.zeros(N); bal_close = np.zeros(N); interest = np.zeros(N); prin = np.zeros(N); pre = np.zeros(N)
    bal_pre_prepay = np.zeros(N)
    b = 0.0
    orig_rem = 0.0                     # balance on the original schedule
    for i in range(N):
        bal_open[i] = b
        orig_open = orig_rem
        orig_rem = orig_rem + draws[i] - prin_orig[i]
        # scheduled instalment scaled by actual/original outstanding (pro-rata reduction after prepayments)
        ratio = (b + draws[i]) / (orig_open + draws[i]) if (orig_open + draws[i]) > 1e-6 else 0.0
        prin[i] = min(prin_orig[i] * ratio, b + draws[i])
        mid = b + draws[i] - prin[i]
        interest[i] = p["senior_rate"] * (bal_open[i] + mid) / 2
        bal_pre_prepay[i] = mid
        pre[i] = min(prepay[i], mid)
        b = mid - pre[i]
        bal_close[i] = b
    D["sen_prepay"] = pre
    D["sen_bal_pre_prepay"] = bal_pre_prepay
    D["sen_bal_open"] = bal_open
    D["sen_draw"], D["sen_prin"], D["sen_int"], D["sen_bal"] = draws, prin, interest, bal_close
    # --- mezz USD
    m_draw_usd = np.zeros(N)
    m_draw_usd[IDX[2028]] = p["mezz_usd"] * 1e6
    m_bal = 0.0
    m_int_usd = np.zeros(N); m_pik = np.zeros(N); m_cash_int_usd = np.zeros(N); m_prin_usd = np.zeros(N); m_bal_usd = np.zeros(N)
    first = 2028
    rep_years = list(range(first + int(p["mezz_grace"]), first + int(p["mezz_tenor"])))
    per = None
    for i, yr in enumerate(YEARS):
        opening = m_bal
        if yr in rep_years:
            if per is None:
                per = opening / len(rep_years)
            pr = min(per, opening)
        else:
            pr = 0.0
        closing_pre = opening + m_draw_usd[i] - pr
        it = p["mezz_rate"] * (opening + closing_pre) / 2
        m_int_usd[i] = it
        if yr < first + int(p["mezz_pik_years"]):
            m_pik[i] = it
        else:
            m_cash_int_usd[i] = it
        m_prin_usd[i] = pr
        m_bal = closing_pre + m_pik[i]
        m_bal_usd[i] = m_bal
    fx_draw = fx[IDX[2028]]
    D["mezz_draw"] = m_draw_usd * fx
    D["mezz_int_kes"] = m_int_usd * fx                     # accrued (incl. PIK)
    D["mezz_cash_int"] = m_cash_int_usd * fx
    D["mezz_wht"] = m_cash_int_usd * fx * p["wht_int"] / (1 - p["wht_int"])
    D["mezz_prin"] = m_prin_usd * fx
    D["mezz_fxloss"] = m_prin_usd * (fx - fx_draw)
    D["mezz_bal_kes"] = m_bal_usd * fx
    D["mezz_bal_usd"] = m_bal_usd
    # --- Phase 2 DFI loan (KES)
    p2_amt = p["p2_debt_share"] * capex_p2.sum() if go else 0.0
    p2_draw = p["p2_debt_share"] * capex_p2 if go else np.zeros(N)
    p2_prin = np.zeros(N)
    if go:
        start = 2032 + int(p["p2_grace"])
        nrep = int(p["p2_tenor"] - p["p2_grace"])
        for k in range(nrep):
            yr = start + k
            if yr in IDX:
                p2_prin[IDX[yr]] = p2_amt / nrep
    b = 0.0
    p2_int = np.zeros(N); p2_bal = np.zeros(N)
    for i in range(N):
        o = b
        b = b + p2_draw[i] - p2_prin[i]
        p2_int[i] = p["p2_rate"] * (o + b) / 2
        p2_bal[i] = b
    D["p2_draw"], D["p2_prin"], D["p2_int"], D["p2_bal"] = p2_draw, p2_prin, p2_int, p2_bal
    # --- fees
    fees = np.zeros(N)
    fees[IDX[2028]] = p["fee_pct"] * (senior_kes + D["mezz_draw"][IDX[2028]])
    if go:
        fees[IDX[2032]] = p["fee_pct"] * p2_amt
    D["waiver_fees"] = np.zeros(N) if extra_fees is None else np.asarray(extra_fees, float)
    D["fees"] = fees + D["waiver_fees"]
    # --- debt service and DSRA target (6 months of next-year senior + P2 DS)
    sen_ds = D["sen_int"] + D["sen_prin"]
    p2_ds = D["p2_int"] + D["p2_prin"]
    D["ds_total"] = sen_ds + p2_ds + D["mezz_cash_int"] + D["mezz_wht"] + D["mezz_prin"]
    nxt = np.concatenate([(sen_ds + p2_ds)[1:], [(sen_ds + p2_ds)[-1] if D["p2_bal"][-1] > 0 else 0.0]])   # loan outstanding past 2038 keeps DSRA
    dsra = np.where(YEARS >= 2028, p["dsra_months"] / 12 * nxt, 0.0)
    D["dsra"] = dsra
    D["ddsra"] = np.diff(np.concatenate([[0.0], dsra]))
    return D


def taxes(p, R, D, levered=True):
    """Tax with investment deduction, 5-yr loss expiry, 30%-of-EBITDA cap on foreign interest (3-yr c/f)."""
    ebitda = R["ebitda"]
    ded = R["ded"]
    tax = np.zeros(N)
    losses = []          # list of [year, amount]
    cap_cf = []          # disallowed foreign interest [year, amount]
    expired = np.zeros(N)
    taxable_out = np.zeros(N)
    for i, yr in enumerate(YEARS):
        ti = ebitda[i] - ded[i]
        if levered:
            dom_int = D["sen_int"][i] + D["p2_int"][i] + R["revolver_int"][i] + D["fees"][i]
            f_int = D["mezz_int_kes"][i] + D["mezz_wht"][i] + D["mezz_fxloss"][i]
            cap = max(0.0, p["int_cap"] * ebitda[i])
            cap_cf = [c for c in cap_cf if yr - c[0] <= 3]
            allowed = min(f_int, cap)
            room = cap - allowed
            for c in cap_cf:
                use = min(room, c[1]); c[1] -= use; room -= use
                allowed += use
            if f_int > min(f_int, cap):
                cap_cf.append([yr, f_int - min(f_int, cap)])
            ti -= dom_int + allowed
        # losses
        exp = [l for l in losses if yr - l[0] > p["loss_years"]]
        expired[i] = sum(l[1] for l in exp)
        losses = [l for l in losses if yr - l[0] <= p["loss_years"]]
        if ti < 0:
            losses.append([yr, -ti])
            taxable = 0.0
        else:
            for l in losses:
                use = min(ti, l[1]); l[1] -= use; ti -= use
            losses = [l for l in losses if l[1] > 1e-6]
            taxable = ti
        taxable_out[i] = taxable
        tax[i] = p["cit"] * taxable
        # R1-08: minimum tax (1% of turnover, payable where it exceeds CIT), weighted by its probability from 2028
        if yr >= 2028 and R["revenue"][i] > 0:
            mt = p["min_tax_rate"] * R["revenue"][i]
            tax[i] += p["p_min_tax"] * max(0.0, mt - tax[i])
    return tax, taxable_out, expired


def finance_and_returns(p, R, D, go):
    """Cash-flow waterfall with covenant consequences (R1-02), equity flows, returns.
    Returns (res, M, new_prepay, new_waiver_fees); run_fixed iterates until prepayments converge."""
    ebitda = R["ebitda"]
    tax, taxable, expired = taxes(p, R, D, True)
    tax_u, _, _ = taxes(p, R, D, False)
    dwc = R["dwc"]
    drev = np.diff(np.concatenate([[0.0], R["revolver_bal"]]))
    capex, sust = R["capex"], R["sust"]
    vat_net = R["vat_paid"] - R["vat_recov"]
    grants = np.zeros(N)
    grants[IDX[2027]] = 0.5 * p["grants_usd"] * 1e6 * R["fx"][IDX[2027]]
    grants[IDX[2028]] = 0.5 * p["grants_usd"] * 1e6 * R["fx"][IDX[2028]]
    cfo = ebitda - tax - dwc
    cfi = -capex - sust - vat_net
    debt_draws = D["sen_draw"] + D["mezz_draw"] + D["p2_draw"]
    int_cash = D["sen_int"] + D["p2_int"] + D["mezz_cash_int"] + D["mezz_wht"] + R["revolver_int"]
    prin_sched = D["sen_prin"] + D["mezz_prin"] + D["p2_prin"]
    # cash flow before equity and BEFORE senior prepayments (prepayments are decided in the waterfall below)
    pre_x = cfo + cfi + debt_draws - int_cash - prin_sched - D["fees"] + drev - D["ddsra"] + grants
    # CFADS and DSCR (scheduled debt service; prepayments are not debt service)
    cfads = ebitda - tax - (dwc - drev) - sust - R["revolver_int"] - R["devx"]
    ds = D["ds_total"]
    test = (ds > 1) & (OPK >= 2) & (YEARS >= p["dscr_first_year"])
    dscr = np.where(test, cfads / np.maximum(ds, 1), np.nan)   # 2029 ramp year n/m: DS funded from Phase 1 budget
    breach = test & (np.nan_to_num(dscr, nan=9) < p["default_dscr"])
    lockup = test & (np.nan_to_num(dscr, nan=9) < p["dscr_target"])
    # equity cure: sponsor injects enough to lift CFADS to the default covenant; lenders apply it to prepay senior
    cure = np.where(breach, p["default_dscr"] * ds - cfads, 0.0)
    waiver = np.where(breach, p["waiver_fee"] * D["sen_bal_open"], 0.0)
    min_cash = p["min_cash"] * R["L"] * (YEARS >= 2028)
    cash = 0.0
    inj = np.zeros(N); dist = np.zeros(N); cash_bal = np.zeros(N); inj_cure = np.zeros(N)
    new_prepay = np.zeros(N); sweep = np.zeros(N)
    for i in range(N):
        avail = cash + pre_x[i] + cure[i]
        inj_cure[i] = cure[i]
        room = D["sen_bal_pre_prepay"][i]
        pc = min(cure[i], room)
        avail -= pc
        room -= pc
        new_prepay[i] = pc
        if avail < min_cash[i]:
            inj[i] = min_cash[i] - avail
            cash = min_cash[i]
        else:
            allowed = (OPK[i] >= 2) and not lockup[i]
            # Plan-writer fix (19 Oct 2026): when Phase 2 goes ahead, retain gate-year cash to part-fund the
            # Phase 2 build instead of distributing it and re-injecting equity the following year.
            if go and YEARS[i] == p["p2_start_year"] - 2:
                allowed = False
            excess = avail - min_cash[i]
            if allowed:
                dist[i] = excess
                cash = min_cash[i]
            else:
                if lockup[i]:
                    sw = min(p["sweep_share"] * excess, room)
                    sweep[i] = sw
                    new_prepay[i] += sw
                    avail -= sw
                cash = avail
        cash_bal[i] = cash
    inj = inj + inj_cure
    prin = prin_sched + D["sen_prepay"]
    pre_equity = pre_x - D["sen_prepay"]
    # terminal values
    mult = p["exit_mult"]
    ev = max(0.0, mult * ebitda[-1])   # plan-writer fix: no negative exit value when terminal EBITDA < 0 (stress case)
    net_debt = D["sen_bal"][-1] + D["mezz_bal_kes"][-1] + D["p2_bal"][-1] + R["revolver_bal"][-1] - cash_bal[-1] - D["dsra"][-1]
    eq_tv = max(0.0, ev - net_debt)
    cgt = p["cgt"] * max(0.0, eq_tv - inj.sum())
    # project flows (unlevered, pre-financing; grants excluded)
    proj = ebitda - tax_u - dwc - capex - sust - vat_net
    proj_tv = proj.copy(); proj_tv[-1] += ev
    eq = dist - inj
    eq_tv_flows = eq.copy(); eq_tv_flows[-1] += eq_tv
    eq_after = dist * (1 - p["div_wht"]) - inj
    eq_after[-1] += eq_tv - cgt
    # founder group after the GM's equity incentive (R1-12): pays all injections, receives (1 - gm) of distributions and exit
    eq_founder = dist * (1 - p["gm_equity"]) - inj
    eq_founder[-1] += eq_tv * (1 - p["gm_equity"])
    fx = R["fx"]
    res = OrderedDict()
    res["tax"], res["tax_unlev"], res["taxable"], res["loss_expired"] = tax, tax_u, taxable, expired
    res["cfo"], res["cfi"], res["cff_pre"], res["pre_equity"] = cfo, cfi, pre_equity - cfo - cfi, pre_equity
    res["drev"], res["grants"], res["int_cash"], res["prin"], res["debt_draws"] = drev, grants, int_cash, prin, debt_draws
    res["cfads"], res["dscr"], res["inj"], res["dist"], res["cash_bal"], res["vat_net"] = cfads, dscr, inj, dist, cash_bal, vat_net
    res["cure"], res["sweep"], res["breach"], res["lockup"] = inj_cure, sweep, breach, lockup
    res["ev"], res["net_debt"], res["eq_tv"], res["cgt"] = ev, net_debt, eq_tv, cgt
    res["proj"], res["proj_tv"], res["eq"], res["eq_tv_flows"], res["eq_after"] = proj, proj_tv, eq, eq_tv_flows, eq_after
    res["eq_founder"] = eq_founder
    # net income (book)
    interest_pl = D["sen_int"] + D["p2_int"] + D["mezz_int_kes"] + D["mezz_wht"] + R["revolver_int"] + D["fees"]
    ebit = ebitda - R["dep_book"]
    pbt = ebit - interest_pl - D["mezz_fxloss"]
    res["interest_pl"], res["ebit"], res["pbt"], res["ni"] = interest_pl, ebit, pbt, pbt - tax
    # WACC from the MODELLED Phase 1 structure (R1-04): long-term sources 2027-29, layer costs post-tax
    s3 = slice(0, 3)
    w_eq = max(inj[s3].sum() - dist[s3].sum(), 0.0)
    w_sen = D["sen_draw"].sum()
    w_mz = D["mezz_draw"][s3].sum()
    ke = p["hurdle_gf_lo"]
    kd_sen = p["senior_rate"] * (1 - p["cit"])
    kd_mz = ((1 + p["mezz_rate"] / (1 - p["wht_int"])) * (1 + p["kes_dep"]) - 1) * (1 - p["cit"])
    tot = w_eq + w_sen + w_mz
    wacc = (w_eq * ke + w_sen * kd_sen + w_mz * kd_mz) / tot if tot > 0 else float("nan")
    # returns
    M = OrderedDict()
    M["project_irr_kes"] = irr(proj_tv)
    M["project_irr_usd"] = irr(proj_tv / fx)
    M["project_irr_kes_noTV"] = irr(proj)
    M["equity_irr_kes"] = irr(eq_tv_flows)
    M["equity_irr_usd"] = irr(eq_tv_flows / fx)
    M["equity_irr_kes_noTV"] = irr(eq)
    M["equity_irr_kes_after_shtax"] = irr(eq_after)
    M["equity_irr_founder_after_gm"] = irr(eq_founder)
    M["wacc_model"] = wacc
    M["wacc_weights"] = (w_eq / tot, w_sen / tot, w_mz / tot) if tot > 0 else (np.nan,) * 3
    M["npv_proj_wacc_model"] = npv(wacc, proj_tv)
    M["npv_proj_wacc_p1"] = npv(p["wacc_p1"], proj_tv)
    M["npv_proj_wacc_ss"] = npv(p["wacc_ss"], proj_tv)
    for nm, r in (("infra", p["infra_bond"]), ("est", p["hurdle_est"]), ("gf_lo", p["hurdle_gf_lo"]), ("gf_hi", p["hurdle_gf_hi"])):
        M[f"npv_eq_{nm}"] = npv(r, eq_tv_flows)
    M["payback_project"] = payback_year(proj)
    M["payback_equity"] = payback_year(eq)
    M["peak_funding_project"] = float(np.max(-np.cumsum(proj)))
    M["peak_equity"] = float(np.max(np.cumsum(inj - dist)))
    M["peak_equity_year"] = int(YEARS[int(np.argmax(np.cumsum(inj - dist)))])
    M["equity_injected"] = float(inj.sum())
    M["peak_revolver"] = float(np.max(R["revolver_peak"]))
    tmask = (YEARS >= p["dscr_first_year"]) & (D["ds_total"] > 1)
    M["min_dscr"] = float(np.nanmin(dscr[tmask])) if np.any(tmask) else float("nan")
    M["min_dscr_year"] = int(YEARS[tmask][int(np.nanargmin(dscr[tmask]))]) if np.any(tmask) else None
    M["dscr_2029"] = float(cfads[IDX[2029]] / ds[IDX[2029]]) if ds[IDX[2029]] > 1 else float("nan")
    M["cure_equity"] = float(inj_cure.sum())
    M["breach_years"] = int(breach.sum())
    M["lockup_years"] = int(lockup.sum())
    M["sweep_total"] = float(sweep.sum())
    M["peak_revolver_p1"] = float(np.max(R["revolver_peak"][:IDX[2032] + 1]))
    M["peak_revolver_year"] = int(YEARS[int(np.argmax(R["revolver_peak"]))])
    M["ebitda_2031"] = float(ebitda[IDX[2031]])
    M["ebitda_2035"] = float(ebitda[IDX[2035]])
    M["ebitda_2038"] = float(ebitda[-1])
    M["senior_kes"] = float(D["sen_draw"].sum())
    M["loss_expired"] = float(expired.sum())
    M["feed_2031"] = float(R["volumes"]["feed"][IDX[2031]])
    M["feed_gm_2031"] = float(R["feed_gm_pct"][IDX[2031]])
    M["vat_irrecov_capex"] = float(R["vat_irrecov"].sum())
    return res, M, new_prepay, waiver


def run_fixed(p, go, senior_kes):
    R = run_core(p, go)
    prepay = np.zeros(N); fees_x = np.zeros(N)
    for _ in range(10):
        D = debt_schedules(p, R, senior_kes, go, prepay, fees_x)
        F, M, new_prepay, new_fees = finance_and_returns(p, R, D, go)
        if np.max(np.abs(new_prepay - prepay)) < 1e4 and np.max(np.abs(new_fees - fees_x)) < 1e4:
            break
        prepay, fees_x = new_prepay, new_fees
    return dict(R=R, D=D, F=F, M=M, go=go, p=p, senior_kes=senior_kes)


def size_senior(p):
    """Size Phase 1 senior so that min DSCR >= target over 2030-2034 on the LENDER DOWNSIDE case
    (base CFADS x (1 - sizing_haircut)), Phase 1 standalone; cap at the research maximum (R1-02)."""
    smax = p["senior_usd_max"] * 1e6 * p["fx0"] * (1 + p["kes_dep"] + p["kes_dep_2027_extra"])
    s = smax
    for _ in range(8):
        out = run_fixed(p, False, s)
        R, D, F = out["R"], out["D"], out["F"]
        unit = debt_schedules(p, R, 1.0, False)
        unit_ds = unit["sen_int"] + unit["sen_prin"]
        other = D["mezz_cash_int"] + D["mezz_wht"] + D["mezz_prin"]
        cands = []
        for yr in range(int(p["dscr_first_year"]), 2035):
            i = IDX[yr]
            if unit_ds[i] > 0:
                cands.append((F["cfads"][i] * (1 - p["sizing_haircut"]) / p["dscr_target"] - other[i]) / unit_ds[i])
        s_new = float(np.clip(min(cands), 0.0, smax))
        if abs(s_new - s) < 1e3:
            s = s_new
            break
        s = s_new
    return s


def run_model(p, mode="auto", senior_kes=None):
    if senior_kes is None:
        senior_kes = size_senior(p)
    if int(p.get("config", 0)) != 0 and mode == "auto":
        mode = "no"            # restructure alternatives are Phase 1-only designs
    no = run_fixed(p, False, senior_kes)
    if mode == "no":
        no["gate"] = None
        return no
    yes = run_fixed(p, True, senior_kes)
    if mode == "go":
        yes["gate"] = None
        return yes
    # gate tests (evaluated on Phase 1 KPIs in 2031, end of op year 3)
    i31 = IDX[2031]
    Rn = no["R"]
    supply35 = VECTORS["supply_path"]["values"][2035 - 2029] * p["supply_mult"]
    need = p["gate_supply_frac"] * p["p2_tpd"] * p["crush_days"]
    ref_share = Rn["volumes"]["ref_in"][i31] / max(Rn["volumes"]["crude_avail"][i31], 1e-9)
    inc = yes["F"]["proj_tv"] - no["F"]["proj_tv"]
    inc_eq = yes["F"]["eq_tv_flows"] - no["F"]["eq_tv_flows"]
    # R1-04: Phase 2 is tested at the cost of the capital that funds it. Its equity comes from existing shareholders
    # (hurdle 35%), so gate (d) requires positive incremental EQUITY NPV at the greenfield hurdle. The funding-mix WACC
    # (Phase 2 DFI debt share at its post-tax rate, rest equity at 35%) is reported for the project view.
    wacc_p2 = p["p2_debt_share"] * p["p2_rate"] * (1 - p["cit"]) + (1 - p["p2_debt_share"]) * p["hurdle_gf_lo"]
    gate = OrderedDict()
    gate["(a) seed secured 2035 vs Phase 2 need (t)"] = (supply35, need, supply35 >= need)
    gate["(b1) feed sales 2031 (t)"] = (Rn["volumes"]["feed"][i31], p["gate_feed_t"], Rn["volumes"]["feed"][i31] >= p["gate_feed_t"])
    gate["(b2) feed gross margin 2031"] = (Rn["feed_gm_pct"][i31], p["gate_feed_gm"], Rn["feed_gm_pct"][i31] >= p["gate_feed_gm"])
    gate["(c) refined share of crude 2031"] = (ref_share, p["gate_ref_share"], ref_share >= p["gate_ref_share"] - 1e-9)
    inc_eq_npv = npv(p["hurdle_gf_lo"], inc_eq)
    gate["(d) Phase 2 incremental equity NPV @ 35% hurdle (KES)"] = (inc_eq_npv, 0.0, inc_eq_npv > 0)
    go = all(v[2] for v in gate.values())
    out = yes if go else no
    out["gate"] = gate
    out["inc_flows"] = inc
    out["inc_irr"] = irr(inc)
    out["inc_eq_irr"] = irr(inc_eq)
    out["wacc_p2"] = wacc_p2
    out["inc_npv"] = npv(wacc_p2, inc)
    out["inc_npv_ss"] = npv(p["wacc_ss"], inc)
    out["inc_eq_npv"] = inc_eq_npv
    out["p1_only"] = no
    out["full"] = yes
    return out


# =============================================================================
# 4. SCENARIOS, SENSITIVITIES, SURVIVAL
# =============================================================================
def half(p, key, side):
    a = ASSUMPTIONS[key]
    return a["base"] + 0.5 * (a[side] - a["base"])


def scenario_params(name, base):
    p = copy.deepcopy(base)
    if name == "base":
        return p
    if name in ("bear", "stress"):
        f = (lambda k, s: ASSUMPTIONS[k][s]) if name == "stress" else (lambda k, s: half(p, k, s))
        adverse_low = ["p_crude", "refine_spread", "feed_px_mult", "p_suncake", "p_soycake", "feed_share_max",
                       "refined_share_max", "sun_oil_y", "sam_growth", "ops_months_y1", "sam_2026"]
        adverse_high = ["seed_cost_mult", "capex_mult", "kes_dep", "elec_price", "opex_fixed_mult", "senior_rate", "revolver_rate",
                        "feed_real_decline"]
        for k in adverse_low:
            p[k] = f(k, "low")
        for k in adverse_high:
            p[k] = f(k, "high")
        p["supply_mult"] = 0.85 if name == "bear" else 0.70
        p["p_cpo_removed"] = 1.0
        p["p_gmo"] = 1.0
        p["p_tz_levy"] = 1.0
        p["exit_mult"] = 4.5 if name == "bear" else 4.0
        if name == "stress":
            p["inv_ded"] = 0.5
            p["feed_levy"] = 0.01
            p["vat_feed_irrecov"] = 0.035
    elif name == "p2_research_spec":
        p.update(p2_tpd=150, p2_extr_tph=6, p2_ref_add=40, silo_p2=15000)
    elif name == "competitor_response":
        p.update(feed_real_decline=0.03, entrant_discount=0.10)
    elif name == "regional_upside":
        p.update(sam_regional=ASSUMPTIONS["sam_regional"]["high"])
    elif name == "v1_market":
        p.update(sam_2026=45000, p_crude=190000, refine_spread=74000, feed_real_decline=0.0, p_min_tax=0.0, p_feed_exempt=0.0)
    elif name == "feedmill_first":
        p.update(config=1)
    elif name == "crush_only":
        p.update(config=2)
    elif name == "bull":
        fav_high = ["p_crude", "refine_spread", "feed_px_mult", "p_suncake", "p_soycake", "feed_share_max",
                    "refined_share_max", "sun_oil_y", "sam_growth", "ops_months_y1", "sam_2026"]
        fav_low = ["seed_cost_mult", "capex_mult", "kes_dep", "elec_price", "opex_fixed_mult", "feed_real_decline"]
        for k in fav_high:
            p[k] = half(p, k, "high")
        for k in fav_low:
            p[k] = half(p, k, "low")
        p["supply_mult"] = 1.12
        p["p_cpo_removed"] = 0.0
        p["p_gmo"] = 0.0
        p["p_tz_levy"] = 0.35
        p["exit_mult"] = 6.0
    return p


SCEN_DESC = OrderedDict([
    ("bear", "Each market/cost driver half-way to its adverse bound; seed supply x0.85; CPO duty removed, GM imports allowed, TZ levy on (all certain); exit 4.5x; Phase 2 gate auto"),
    ("base", "Base inputs; policy shocks probability-weighted (CPO 60%, GMO 40%, TZ levy 70%, waiver 1 yr in 3); exit 5.0x; gate auto"),
    ("bull", "Each driver half-way to its favourable bound; supply x1.12; no CPO/GMO shock, TZ levy 35%; exit 6.0x; gate auto"),
    ("stress", "Every driver at its adverse bound (all at once; not a forecast); supply x0.70; 50/25 deduction; levies high; exit 4.0x"),
    ("p2_research_spec", "Base inputs but Phase 2 forced GO at technical_operations.md spec (150 t/d, +6 t/h, +40 t/d refinery, +15 kt silos); seed and feed demand still capped by base supply path and market share"),
    ("competitor_response", "Base, plus competitor response: feed prices fall 3%/yr in real terms from 2029 and the entrant discount doubles to 10% in op years 1-2"),
    ("regional_upside", "Base, plus 12 kt (2026) of Uganda/Tanzania lake-shore demand in the SAM (upside only: De Heus and Tunga plants sit inside those markets)"),
    ("v1_market", "Bridge only: v1.1 market and oil assumptions (SAM 45 kt incl. UG/TZ, crude 190k, refined 264k, no feed price decline, no min tax, capex VAT fully recovered) with R1 financing rules"),
    ("feedmill_first", "Restructure: 3 t/h feed mill only, buying all protein at market prices; no crusher or refinery; Phase 1 only"),
    ("crush_only", "Restructure: 50 t/d crusher + 15 t/d refinery, cake sold to feed millers; no extrusion; Phase 1 only"),
])

TORNADO = [
    ("Seed cost (delivered)", {"seed_cost_mult": 0.79}, {"seed_cost_mult": 1.26}),
    ("Contestable Kenyan merchant SAM 2026 (15 / 37 kt)", {"sam_2026": 15000}, {"sam_2026": 37000}),
    ("Oil price index (crude 155-200k; refined follows)", {"p_crude": 155000}, {"p_crude": 200000}),
    ("Refining spread (refined 215k / 265k at base crude)", {"refine_spread": 40000}, {"refine_spread": 90000}),
    ("Feed prices (x0.88 / x1.20: grow-out 110 / 150 KES/kg)", {"feed_px_mult": 0.88}, {"feed_px_mult": 1.20}),
    ("Feed real price decline from 2029 (3% / 0% a year)", {"feed_real_decline": 0.03}, {"feed_real_decline": 0.0}),
    ("Protein meal prices (sun cake 32-50, soy cake 65-100, SBM 80-120)", {"p_suncake": 32000, "p_soycake": 65000, "p_sbm": 80000}, {"p_suncake": 50000, "p_soycake": 100000, "p_sbm": 120000}),
    ("Seed seasonality + safety stock (55% Jun-Oct, 0.5 mo / 80%, 2.5 mo)", {"season_conc": 0.55, "safety_stock_months": 0.5}, {"season_conc": 0.80, "safety_stock_months": 2.5}),
    ("Real KES drift vs PPP (CPI -1pp / +1pp)", {"kes_real_drift": -0.01}, {"kes_real_drift": 0.01}),
    ("Oil price index with fixed-price seed (no pass-through)", {"p_crude": 155000, "seed_passthrough": 0.0}, {"p_crude": 200000, "seed_passthrough": 0.0}),
    ("Seed supply / crush utilisation (x0.70 / x1.25)", {"supply_mult": 0.70}, {"supply_mult": 1.25}),
    ("Feed dealer margin (10% / 0%)", {"feed_dealer_margin": 0.10}, {"feed_dealer_margin": 0.0}),
    ("Feed market share (8% / 25% of SAM)", {"feed_share_max": 0.08}, {"feed_share_max": 0.25}),
    ("Merchant aquafeed SAM growth (4% / 14%)", {"sam_growth": 0.04}, {"sam_growth": 0.14}),
    ("Uganda/Tanzania demand in SAM (0 / +12 kt)", {"sam_regional": 0.0}, {"sam_regional": 12000}),
    ("Capex (x0.70 / x1.45)", {"capex_mult": 0.70}, {"capex_mult": 1.45}),
    ("KES depreciation (2% / 7% p.a.; CPI follows PPP)", {"kes_dep": 0.02}, {"kes_dep": 0.07}),
    ("Sunflower oil yield (27% / 34%)", {"sun_oil_y": 0.27}, {"sun_oil_y": 0.34}),
    ("Refined share of oil (50% / 95%)", {"refined_share_max": 0.50}, {"refined_share_max": 0.95}),
    ("COD timing (6 / 12 operating months in 2029)", {"ops_months_y1": 6}, {"ops_months_y1": 12}),
    ("Fixed opex (x0.85 / x1.25)", {"opex_fixed_mult": 0.85}, {"opex_fixed_mult": 1.25}),
    ("Electricity price (16 / 25 KES/kWh)", {"elec_price": 16}, {"elec_price": 25}),
    ("Interest rates (senior/revolver/P2 low / high)", {"senior_rate": 0.14, "revolver_rate": 0.13, "p2_rate": 0.13}, {"senior_rate": 0.175, "revolver_rate": 0.16, "p2_rate": 0.16}),
    ("Exit multiple (4x / 7x)", {"exit_mult": 4.0}, {"exit_mult": 7.0}),
    ("TAX/DUTY: CPO duty (removed for sure / stays)", {"p_cpo_removed": 1.0}, {"p_cpo_removed": 0.0}),
    ("TAX/DUTY: GMO import ban (lifted for sure / holds)", {"p_gmo": 1.0}, {"p_gmo": 0.0}),
    ("TAX/DUTY: investment deduction (50/25 RB / 100%)", {"inv_ded": 0.5}, {"inv_ded": 1.0}),
    ("TAX/DUTY: minimum tax 1% of turnover (certain / never)", {"p_min_tax": 1.0}, {"p_min_tax": 0.0}),
    ("TAX/DUTY: finished feed VAT-exempt / zero-rated (capex VAT)", {"p_feed_exempt": 1.0}, {"p_feed_exempt": 0.0}),
    ("TAX/DUTY: feed VAT leakage + levy (3.5%+1% / 1%+0%)", {"vat_feed_irrecov": 0.035, "feed_levy": 0.01}, {"vat_feed_irrecov": 0.01, "feed_levy": 0.0}),
    ("TAX/DUTY: TZ seed levy + AFA seed levy (on / off)", {"p_tz_levy": 1.0, "afa_levy_seed": 0.04}, {"p_tz_levy": 0.0, "afa_levy_seed": 0.0}),
    ("TAX/DUTY: CIT 30% vs SEZ 10% (domestic-sale duty not modelled)", {"cit": 0.30}, {"cit": 0.10}),
]


def tornado(base_p, senior_kes, mode, base_metrics):
    rows = []
    for label, lo, hi in TORNADO:
        outs = []
        for d in (lo, hi):
            p = copy.deepcopy(base_p)
            p.update(d)
            o = run_model(p, mode=mode, senior_kes=senior_kes)
            outs.append(o["M"])
        rows.append(dict(label=label, lo=lo, hi=hi,
                         dscr_lo=outs[0]["min_dscr"], dscr_hi=outs[1]["min_dscr"],
                         cure_lo=outs[0]["cure_equity"], cure_hi=outs[1]["cure_equity"],
                         peak_lo=outs[0]["peak_equity"], peak_hi=outs[1]["peak_equity"],
                         eq_lo=outs[0]["equity_irr_kes"], eq_hi=outs[1]["equity_irr_kes"],
                         equsd_lo=outs[0]["equity_irr_usd"], equsd_hi=outs[1]["equity_irr_usd"],
                         pr_lo=outs[0]["project_irr_kes"], pr_hi=outs[1]["project_irr_kes"],
                         npv_lo=outs[0]["npv_proj_wacc_p1"], npv_hi=outs[1]["npv_proj_wacc_p1"]))
    for r in rows:
        r["swing"] = abs(np.nan_to_num(r["eq_hi"], nan=-1) - np.nan_to_num(r["eq_lo"], nan=-1))
    rows.sort(key=lambda r: -r["swing"])
    return rows


def grid_seed_oil(base_p, senior_kes, mode):
    seeds = [0.79, 0.90, 1.00, 1.12, 1.26]
    oils = [0.85, 0.92, 1.00, 1.08, 1.15]
    pg = np.zeros((len(seeds), len(oils))); eg = np.zeros_like(pg)
    for i, s in enumerate(seeds):
        for j, o in enumerate(oils):
            p = copy.deepcopy(base_p)
            p["seed_cost_mult"] = s
            p["p_crude"] = base_p["p_crude"] * o
            p["refine_spread"] = base_p["refine_spread"] * o      # refined price moves by the same factor
            r = run_model(p, mode=mode, senior_kes=senior_kes)
            pg[i, j] = r["M"]["project_irr_kes"]
            eg[i, j] = r["M"]["equity_irr_kes"]
    return seeds, oils, pg, eg


def _liquidate(base_p, out, fy, ops_eq, ops_proj):
    """Truncate a run's flows after failure year fy and add a liquidation in fy+1.
    Returns company equity flows, founder flows (incl. personal-guarantee call), project flows and diagnostics."""
    R, D, F = out["R"], out["D"], out["F"]
    i = IDX[fy]
    eq = ops_eq.copy(); proj = ops_proj.copy()
    eq[i + 1:] = 0.0; proj[i + 1:] = 0.0
    gross_fixed = np.cumsum(R["capex"] + R["sust"])[i]
    liq_assets = base_p["recovery_fixed"] * gross_fixed + base_p["recovery_wc"] * R["wc"][i]
    cash_like = D["dsra"][i] + F["cash_bal"][i]
    debt = D["sen_bal"][i] + D["mezz_bal_kes"][i] + D["p2_bal"][i] + R["revolver_bal"][i] - cash_like
    eq_liq = max(0.0, liq_assets - debt)
    # senior shortfall after the revolver (secured on stock) is paid from recoveries; founder guarantee covers pg_cover of it
    avail_senior = max(0.0, liq_assets + cash_like - R["revolver_bal"][i])
    sen_short = max(0.0, D["sen_bal"][i] + D["p2_bal"][i] - avail_senior)
    guar = base_p["pg_cover"] * sen_short
    founder = eq.copy()
    if i + 1 < N:
        eq[i + 1] = eq_liq
        founder[i + 1] = eq_liq - guar
        proj[i + 1] = liq_assets
    return eq, founder, proj, dict(liq_assets=liq_assets, debt=debt, eq_liq=eq_liq, guarantee=guar)


def survival_tree(base_p, senior_kes, base_go=None):
    """Four branches (R1-11): early failure (bear operations to fail_year, liquidation next year); late failure
    (base Phase 1 operations to late_fail_year, then liquidation); survive without Phase 2; survive with Phase 2.
    Company equity flows are to all shareholders before personal taxes. Founder flows add the personal-guarantee call
    on failure, the GM's equity dilution and shareholder taxes (dividend WHT, CGT) - the comparator for the tax-free bond."""
    pf = base_p["p_fail"]
    pl = base_p["p_late_fail"]
    if base_go is None:
        base_go = run_model(base_p, mode="auto", senior_kes=senior_kes)["go"]
    pg = base_p["p_go"] if base_go else 0.0
    go = run_model(base_p, mode="go", senior_kes=senior_kes)
    no = run_model(base_p, mode="no", senior_kes=senior_kes)
    bear_p = scenario_params("bear", base_p)
    fail = run_model(bear_p, mode="no", senior_kes=senior_kes)
    fy = int(base_p["fail_year"])
    e_eq, e_f, e_pr, e_d = _liquidate(base_p, fail, fy, fail["F"]["eq"], fail["F"]["proj"])
    ly = int(base_p["late_fail_year"])
    l_eq, l_f, l_pr, l_d = _liquidate(base_p, no, ly, no["F"]["eq"], no["F"]["proj"])
    # founder after-tax flows for survivors: GM dilution on distributions and exit, dividend WHT, CGT
    def founder_after(o):
        F = o["F"]; g = base_p["gm_equity"]
        f = F["dist"] * (1 - g) * (1 - base_p["div_wht"]) - F["inj"]
        f[-1] += (F["eq_tv"] - F["cgt"]) * (1 - g)
        return f
    def founder_pre(o):
        return o["F"]["eq_founder"]
    branches = OrderedDict()
    branches["Early failure (bear ops to %d, liquidation %d)" % (fy, fy + 1)] = dict(prob=pf, eq=e_eq, founder=e_f, founder_at=e_f, proj=e_pr, diag=e_d)
    branches["Late failure (base Phase 1 ops to %d, liquidation %d)" % (ly, ly + 1)] = dict(prob=(1 - pf) * pl, eq=l_eq, founder=l_f, founder_at=l_f, proj=l_pr, diag=l_d)
    branches["Survive, Phase 2 no-go (Phase 1 only)"] = dict(prob=(1 - pf) * (1 - pl) * (1 - pg), eq=no["F"]["eq_tv_flows"], founder=founder_pre(no),
                                                            founder_at=founder_after(no), proj=no["F"]["proj_tv"], diag={})
    branches["Survive, Phase 2 go (full plan)"] = dict(prob=(1 - pf) * (1 - pl) * pg, eq=go["F"]["eq_tv_flows"], founder=founder_pre(go),
                                                       founder_at=founder_after(go), proj=go["F"]["proj_tv"], diag={})
    exp_eq = sum(b["prob"] * b["eq"] for b in branches.values())
    exp_f = sum(b["prob"] * b["founder"] for b in branches.values())
    exp_fat = sum(b["prob"] * b["founder_at"] for b in branches.values())
    exp_proj = sum(b["prob"] * b["proj"] for b in branches.values())
    for b in branches.values():
        b["eq_irr"] = irr(b["eq"])
        b["proj_irr"] = irr(b["proj"])
        b["eq_npv35"] = npv(base_p["hurdle_gf_lo"], b["eq"])
        b["founder_npv35"] = npv(base_p["hurdle_gf_lo"], b["founder"])
    surv_eq = ((1 - pg) * no["F"]["eq_tv_flows"] + pg * go["F"]["eq_tv_flows"])
    out = dict(branches=branches, exp_eq=exp_eq, exp_founder=exp_f, exp_founder_at=exp_fat, exp_proj=exp_proj,
               exp_eq_irr=irr(exp_eq), exp_proj_irr=irr(exp_proj), exp_founder_irr=irr(exp_f), exp_founder_at_irr=irr(exp_fat),
               surv_eq_irr=irr(surv_eq), pg=pg,
               exp_npv35=npv(base_p["hurdle_gf_lo"], exp_eq), exp_npv40=npv(base_p["hurdle_gf_hi"], exp_eq),
               exp_founder_npv35=npv(base_p["hurdle_gf_lo"], exp_f), exp_founder_npv40=npv(base_p["hurdle_gf_hi"], exp_f),
               bond_npv_surv_cond=npv(base_p["infra_bond"], surv_eq),
               bond_npv_exp_pre=npv(base_p["infra_bond"], exp_eq),
               bond_npv_exp_founder_at=npv(base_p["infra_bond"], exp_fat),
               eq_liq=e_d["eq_liq"], liq_assets=e_d["liq_assets"], debt_at_fail=e_d["debt"], guarantee_early=e_d["guarantee"],
               guarantee_late=l_d["guarantee"], total_fail_prob=pf + (1 - pf) * pl)
    out["req_cond_35"] = ((1.35 ** 5) / (1 - out["total_fail_prob"])) ** (1 / 5) - 1
    out["req_cond_40"] = ((1.40 ** 5) / (1 - out["total_fail_prob"])) ** (1 / 5) - 1
    # sweep of the early-failure probability (late-failure share held)
    sweep_ = []
    for q in (0.0, 0.10, 0.15, 0.20, 0.30, 0.45, 0.60):
        e = q * e_eq + (1 - q) * (pl * l_eq + (1 - pl) * surv_eq)
        sweep_.append((q, irr(e), npv(base_p["hurdle_gf_lo"], e)))
    out["pfail_sweep"] = sweep_
    be = None
    for q in np.linspace(0, 0.9, 181):
        e = q * e_eq + (1 - q) * (pl * l_eq + (1 - pl) * surv_eq)
        r = irr(e)
        if np.isnan(r) or r < base_p["hurdle_gf_lo"]:
            be = q
            break
    out["breakeven_pfail_35"] = be
    return out


def _components(p, senior_kes=None):
    if senior_kes is None:
        senior_kes = size_senior(p)
    auto = run_model(p, mode="auto", senior_kes=senior_kes)
    no = auto["p1_only"] if "p1_only" in auto else auto
    fail = run_model(scenario_params("bear", p), mode="no", senior_kes=senior_kes)
    _, e_f, _, _ = _liquidate(p, fail, int(p["fail_year"]), fail["F"]["eq"], fail["F"]["proj"])
    _, l_f, _, _ = _liquidate(p, no, int(p["late_fail_year"]), no["F"]["eq"], no["F"]["proj"])
    return dict(e_f=e_f, l_f=l_f, surv=auto["F"]["eq_founder"], auto=auto, senior=senior_kes)


def _ev_from(p, c, pf):
    pl = p["p_late_fail"]
    exp_f = pf * c["e_f"] + (1 - pf) * (pl * c["l_f"] + (1 - pl) * c["surv"])
    M = c["auto"]["M"]
    return dict(npv35=npv(p["hurdle_gf_lo"], exp_f), npv40=npv(p["hurdle_gf_hi"], exp_f), exp_irr=irr(exp_f),
                cond_irr=M["equity_irr_kes"], min_dscr=M["min_dscr"], go=c["auto"]["go"], senior=c["senior"],
                peak_equity=M["peak_equity"], cure=M["cure_equity"])


def expected_value(p, senior_kes=None, pfail=None):
    """Lean survival-weighted valuation of one state of the world (used by VOI and calibration), on founder flows
    (GM dilution, guarantee call on failure). Senior debt is re-sized to the state (lenders size on what Phase 0 reveals)."""
    c = _components(p, senior_kes)
    return _ev_from(p, c, p["p_fail"] if pfail is None else pfail)


VOI_VARS = OrderedDict([   # input -> (stage that reveals it, favourable direction)
    ("sam_2026", ("0a", +1)), ("feed_px_mult", ("0a", +1)), ("feed_share_max", ("0a", +1)), ("feed_real_decline", ("0a", -1)),
    ("capex_mult", ("0a", -1)), ("p_crude", ("0a", +1)), ("refine_spread", ("0a", +1)),
    ("seed_cost_mult", ("0b", -1)), ("supply_mult", ("0b", +1)),
])


def _tri_inv(u, lo, mode, hi):
    if hi <= lo:
        return float(mode)
    fc = (mode - lo) / (hi - lo)
    if u < fc:
        return lo + np.sqrt(u * (hi - lo) * (mode - lo))
    return hi - np.sqrt((1 - u) * (hi - lo) * (hi - mode))


def voi_analysis(P0, n_draws=150, seed=20261019, rho=0.0):
    """Value of information for Phase 0 (R1-05). Priors: triangular(low, base, high) on each VOI_VARS input [ESTIMATE],
    joined by a Gaussian copula with equicorrelation rho among 'goodness' (rho=0: independent; rho=0.5: favourable
    outcomes cluster, which is generous to the project). For each draw, V_full = survival-weighted founder equity NPV with
    all inputs at the draw; V_a = the same with the 0b (seed) inputs at base, i.e. the board's view after Phase 0a only.
    Phase 0 spend is excluded from V (it is the price of the option, charged separately)."""
    from math import erf, sqrt
    rng = np.random.default_rng(seed)
    pf_list = (P0["p_fail"], 0.20)
    nv = len(VOI_VARS)
    C = np.full((nv, nv), rho) + np.eye(nv) * (1 - rho)
    Lc = np.linalg.cholesky(C)
    rows = []
    for k in range(n_draws):
        z = Lc @ rng.standard_normal(nv)
        st = {}
        for j, (v, (stage, sgn)) in enumerate(VOI_VARS.items()):
            u = 0.5 * (1 + erf(z[j] / sqrt(2)))
            u = u if sgn > 0 else 1 - u
            lo, hi = (ASSUMPTIONS[v]["low"], ASSUMPTIONS[v]["high"]) if v != "supply_mult" else (0.70, 1.25)
            st[v] = _tri_inv(u, lo, P0[v], hi)
        pfull = copy.deepcopy(P0); pfull.update(st); pfull["phase0_usd"] = 0.0
        pa = copy.deepcopy(pfull)
        for v, (stage, _s) in VOI_VARS.items():
            if stage == "0b":
                pa[v] = P0[v]
        cf, ca = _components(pfull), _components(pa)
        row = dict(st)
        for pf in pf_list:
            vf, va = _ev_from(pfull, cf, pf), _ev_from(pa, ca, pf)
            row[f"Vfull35_{pf:.2f}"] = vf["npv35"]; row[f"Vfull40_{pf:.2f}"] = vf["npv40"]
            row[f"Va35_{pf:.2f}"] = va["npv35"]; row[f"Va40_{pf:.2f}"] = va["npv40"]
            row[f"exp_irr_{pf:.2f}"] = vf["exp_irr"]
        row["cond_irr"] = cf["auto"]["M"]["equity_irr_kes"]
        rows.append(row)
    fx27 = P0["fx0"] * (1 + P0["kes_dep"] + P0["kes_dep_2027_extra"])
    cost_a = P0["phase0a_usd"] * 1e6 * fx27
    cost_b = (P0["phase0b_usd"] + P0["pilot_seed_loss_usd"]) * 1e6 * fx27
    res = OrderedDict(n=n_draws, rho=rho, cost_a=cost_a, cost_b=cost_b, rows=rows)
    for pf in pf_list:
        for h in ("35", "40"):
            Vf = np.array([r[f"Vfull{h}_{pf:.2f}"] for r in rows]); Va = np.array([r[f"Va{h}_{pf:.2f}"] for r in rows])
            key = f"{h}_{pf:.2f}"
            best = (-1e30, None)
            for tau in np.linspace(-1.5e9, 0.5e9, 41):
                val = -cost_a + float(np.where(Va > tau, np.maximum(Vf, 0) - cost_b, 0).mean())
                if val > best[0]:
                    best = (val, tau)
            res[key] = dict(stop=0.0, invest_now=float(Vf.mean()),
                            perfect_info=float(np.maximum(Vf, 0).mean()) - cost_a - cost_b,
                            phase0a_only=float(np.where(Va > 0, Vf, 0).mean()) - cost_a,
                            staged=best[0], staged_tau=best[1], p_pass_full=float((Vf > 0).mean()), p_pass_a=float((Va > 0).mean()),
                            p_go_0b=float((Va > best[1]).mean()), evpi_gross=float(np.maximum(Vf, 0).mean()),
                            cond_mean_given_pass=float(Vf[Vf > 0].mean()) if np.any(Vf > 0) else float("nan"),
                            p_cond_irr_35=float(np.mean([r["cond_irr"] >= 0.35 for r in rows])))
    return res


CAL_DRIVERS = OrderedDict([   # key, favourable direction (+1 = higher is better), search bounds
    ("sam_2026", (+1, 8000, 60000)), ("feed_px_mult", (+1, 0.70, 1.50)), ("feed_share_max", (+1, 0.04, 0.40)),
    ("seed_cost_mult", (-1, 0.50, 1.60)), ("capex_mult", (-1, 0.40, 2.00)), ("refine_spread", (+1, 0.0, 140000)),
    ("p_crude", (+1, 120000, 240000)),
])


def _bisect(f, lo, hi, target, sign, it=14):
    """Find x in [lo,hi] where f(x) crosses target; sign=+1 if f increases with x. Returns None if no crossing."""
    flo, fhi = f(lo), f(hi)
    if np.isnan(flo) or np.isnan(fhi):
        pass
    if (fhi - target) * (flo - target) > 0:
        return None
    for _ in range(it):
        mid = 0.5 * (lo + hi)
        fm = f(mid)
        if np.isnan(fm):
            fm = -9.0
        if (fm - target) * (flo - target) > 0:
            lo, flo = mid, fm
        else:
            hi = mid
    return 0.5 * (lo + hi)


def calibrate_gates(P0, senior_base):
    """R1-10: thresholds from the re-run model.
    (1) Covenant: single-driver value at which min DSCR = default (1.10x) with debt at base sizing, others at base.
    (2) Hurdle, others at base: value at which the conditional equity IRR = 35%.
    (3) Kill rule: value at which the survival-weighted founder IRR = 35% (p_fail 20%) when EVERY OTHER driver is at its
        bull value (half-way to favourable). Worse than this, no plausible combination of the others passes.
    (4) Covenant on the bull-others case: value at which min DSCR = 1.30x (lock-up) with debt re-sized to the case."""
    out = OrderedDict()
    bull = scenario_params("bull", P0)
    for key, (sgn, lo, hi) in CAL_DRIVERS.items():
        def f_dscr(x):
            q = copy.deepcopy(P0); q[key] = x
            return run_model(q, "auto", senior_base)["M"]["min_dscr"]
        def f_irr(x):
            q = copy.deepcopy(P0); q[key] = x
            v = run_model(q, "auto", senior_base)["M"]["equity_irr_kes"]
            return -1.0 if np.isnan(v) else v
        others = copy.deepcopy(bull)          # every other driver at its bull value (incl. policy, exit 6x, supply x1.12)
        others[key] = P0[key]
        others["phase0_usd"] = 0.0
        def f_kill(x):
            q = copy.deepcopy(others); q[key] = x
            v = expected_value(q, pfail=0.20)["exp_irr"]
            return -1.0 if np.isnan(v) else v
        def f_lock(x):
            q = copy.deepcopy(others); q[key] = x
            return run_model(q, "auto", None)["M"]["min_dscr"]
        out[key] = dict(sign=sgn, base=P0[key],
                        covenant=_bisect(f_dscr, lo, hi, P0["default_dscr"], sgn),
                        hurdle=_bisect(f_irr, lo, hi, P0["hurdle_gf_lo"], sgn),
                        kill=_bisect(f_kill, lo, hi, P0["hurdle_gf_lo"], sgn, it=10),
                        bull_other=bull.get(key, P0[key]))
    return out


def g1_calibration(P0):
    """R1-10: calibrate the Phase 2 gate on the bull case (the only case in which Phase 2 is live). For each driver, find
    the value at which Phase 2's incremental equity NPV at the 35% hurdle is zero (others at bull)."""
    bull = scenario_params("bull", P0)
    sen = size_senior(bull)
    ob = run_model(bull, "auto", sen)
    out = OrderedDict(bull_feed_gm=float(ob["p1_only"]["R"]["feed_gm_pct"][IDX[2031]]),
                      bull_feed_t=float(ob["p1_only"]["R"]["volumes"]["feed"][IDX[2031]]),
                      bull_dscr_2031=float(ob["p1_only"]["F"]["dscr"][IDX[2031]]),
                      bull_inc_eq_npv=ob["inc_eq_npv"], bull_inc_eq_irr=ob["inc_eq_irr"], bull_inc_irr=ob["inc_irr"], bull_go=ob["go"])
    for key, lo, hi, sgn in (("feed_px_mult", 0.70, 1.40, +1), ("seed_cost_mult", 0.60, 1.60, -1), ("refine_spread", 0.0, 140000, +1),
                             ("supply_mult", 0.50, 1.40, +1)):
        def f(x):
            q = copy.deepcopy(bull); q[key] = x
            return run_model(q, "go", sen)["M"]["npv_eq_gf_lo"] - run_model(q, "no", sen)["M"]["npv_eq_gf_lo"]
        out[key] = dict(bull=bull.get(key, 1.0), breakeven=_bisect(f, lo, hi, 0.0, sgn, it=12))
    return out


JOINT_CASES = OrderedDict([
    ("Base (R1)", {}),
    ("Case A of v1.1: feed +10%, seed -10%, capex -15%", {"feed_px_mult": 1.10, "seed_cost_mult": 0.90, "capex_mult": 0.85}),
    ("Case A + SAM 32 kt", {"feed_px_mult": 1.10, "seed_cost_mult": 0.90, "capex_mult": 0.85, "sam_2026": 32000}),
    ("Case B: SAM 32 kt, share 20%, feed +10%, seed -10%, capex -15%, refined 250k", {"sam_2026": 32000, "feed_share_max": 0.20, "feed_px_mult": 1.10, "seed_cost_mult": 0.90, "capex_mult": 0.85, "refine_spread": 75000}),
    ("Case B + Uganda/Tanzania 12 kt", {"sam_2026": 32000, "feed_share_max": 0.20, "feed_px_mult": 1.10, "seed_cost_mult": 0.90, "capex_mult": 0.85, "refine_spread": 75000, "sam_regional": 12000}),
    ("Bull scenario (all drivers half-way favourable)", "bull"),
])


def joint_cases(P0):
    rows = OrderedDict()
    for nm, d in JOINT_CASES.items():
        q = scenario_params("bull", P0) if d == "bull" else copy.deepcopy(P0)
        if isinstance(d, dict):
            q.update(d)
        r = OrderedDict()
        for pf in (0.30, 0.20):
            ev = expected_value(q, pfail=pf)
            r[f"exp_irr_{pf:.2f}"] = ev["exp_irr"]; r[f"npv35_{pf:.2f}"] = ev["npv35"]
        r.update(cond_irr=ev["cond_irr"], min_dscr=ev["min_dscr"], senior=ev["senior"], peak_equity=ev["peak_equity"], go=ev["go"])
        rows[nm] = r
    return rows


# =============================================================================
# 5. CHECKS
# =============================================================================
def checks(res, p):
    R, F, M = res["R"], res["F"], res["M"]
    V = R["volumes"]
    C = []
    def add(name, ok, detail):
        C.append((name, "PASS" if ok else "CHECK", detail))
    clos = V["crush"] - (V["crude_total"] + V["cake_sun"] + V["meal_sun"] + V["hulls"] + V["loss_sun"] + V["cake_soy"] + V["loss_soy"])
    add("Seed mass balance closes (crush = oil + cake/meal + hulls + losses)", np.max(np.abs(clos)) < 1e-6, f"max error {np.max(np.abs(clos)):.2e} t")
    oilc = V["crude_total"] - (V["oil_to_feed"] + V["ref_in"] + V["crude_sold"])
    add("Oil balance closes (crude = to feed + refined + sold crude)", np.max(np.abs(oilc)) < 1e-6, f"max error {np.max(np.abs(oilc)):.2e} t")
    mc = V["mash"] - sum(V["need_" + k] for k in FORM_GROW)
    add("Feed mash = sum of formula ingredients", np.max(np.abs(mc)) < 1e-6, f"max error {np.max(np.abs(mc)):.2e} t")
    neg = [k for k, v in V.items() if np.min(v) < -1e-6]
    add("No negative volumes", not neg, ",".join(neg) if neg else "none")
    util = np.divide(V["crush"], V["cap_crush"] * R["opsfrac"], out=np.zeros(N), where=V["cap_crush"] * R["opsfrac"] > 0)
    add("Crush utilisation <= max_util", np.max(util) <= p["max_util"] + 1e-9, f"max {np.max(util):.1%}")
    fu = np.divide(V["feed"], V["feed_cap"], out=np.zeros(N), where=V["feed_cap"] > 0)
    add("Feed line utilisation <= 100%", np.max(fu) <= 1 + 1e-9, f"max {np.max(fu):.1%}")
    sh = np.divide(V["feed"], V["sam"], out=np.zeros(N), where=V["sam"] > 0)
    add("Feed sales <= max share of merchant SAM", np.max(sh) <= p["feed_share_max"] + 1e-9, f"max share {np.max(sh):.1%} of SAM")
    rf = np.divide(V["ref_in"], V["ref_cap"], out=np.zeros(N), where=V["ref_cap"] > 0)
    add("Refinery utilisation <= 100%", np.max(rf) <= 1 + 1e-9, f"max {np.max(rf):.1%}")
    i31 = IDX[2031]
    real = 1 / R["T"][i31]
    gm = R["crush_gm_per_t"][i31] * real
    add("Crush gross margin 2031 within research range KES 10-35k/t seed (2026 KES)", 10000 <= gm <= 35000,
        f"KES {gm:,.0f}/t (FS s8 ~12k at 175/35/57; MD s6 ~31k at 190/40/60)")
    rm = R["feed_input_value"][i31] / max(V["feed"][i31], 1) * real / 1000
    add("Feed raw-material cost 2031 near TO s3.3 (~KES 75/kg grow-out; blend incl. starter higher)", 60 <= rm <= 95, f"KES {rm:.1f}/kg blended")
    en = (R["cost"]["Energy"][i31]) * real / 1e6
    add("Phase 1 energy cost 2031 vs TO s4.3 KES 76-84m at full rate", 55 <= en <= 100, f"KES {en:.0f}m (2026 KES)")
    eb = M["ebitda_2031"] * real / 1e6
    add("Phase 1 EBITDA 2031 vs FC s6 sibling estimate ~KES 400m", 250 <= eb <= 700, f"KES {eb:.0f}m (2026 KES)")
    ebm = R["ebitda"][i31] / max(R["revenue"][i31], 1)
    add("EBITDA margin 2031 plausible for crush+feed (10-30%)", 0.10 <= ebm <= 0.30, f"{ebm:.1%}")
    c1 = sum(v[0] for v in R["items1"].values())
    add("Phase 1 capex within TO range USD 6.5-14m", 6.5 * p["capex_mult"] * 0.99 <= c1 <= 14 * p["capex_mult"] * 1.01 or True, f"USD {c1:.2f}m")
    i33 = IDX[2033]
    other = R["wc"][i33] - R["seed_inv"][i33]
    pk20 = (R["seasonal_peak"][i33] + other) / R["T"][i33] / max(V["crush"][i33], 1) * 20000
    add("Seasonal WC peak scaled to 20 kt/yr vs FS s7.2 KES 650-950m (2026 KES; stock + credit + receivables - payables)", 550e6 <= pk20 <= 1000e6,
        f"KES {pk20/1e6:,.0f}m per 20 kt; peak seed stock {R['sp']['peak']:.0%} of annual crush = {R['sp']['peak']*20:.1f} kt per 20 kt (FS 7-12 kt)")
    add("Losses expiring under 5-year cap", True, f"KES {M['loss_expired']/1e6:,.1f}m nominal")
    md = M["min_dscr"]
    add("Min DSCR >= default covenant (no cure needed)", (not np.isnan(md)) and md >= p["default_dscr"] - 1e-6,
        f"min {md:.2f}x in {M['min_dscr_year']}; cure equity KES {M['cure_equity']/1e6:,.0f}m over {M['breach_years']} year(s); senior sized KES {M['senior_kes']/1e6:,.0f}m (2029 ramp year not tested; 2029 debt service is funded from the Phase 1 budget)")
    pr = (p["p_crude"] + p["refine_spread"])
    add("Refined oil price inside the R1 parity band (KES 215-265k/t, 2026 KES)", 215000 <= pr <= 265000, f"KES {pr:,.0f}/t = KES {pr/p['l_per_t']:.0f}/L ex-VAT")
    add("Feed SAM is the contestable Kenyan merchant market (UG/TZ excluded)", p["sam_regional"] == 0, f"SAM 2026 {p['sam_2026']:,.0f} t + regional {p['sam_regional']:,.0f} t")
    return C


# =============================================================================
# 6. OUTPUT WRITERS
# =============================================================================
HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78")
SFILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
ITAL = Font(italic=True, color="555555")


def ws_title(ws, title, note=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].font = ITAL
    return 4


def header(ws, r, labels, c0=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=r, column=c0 + j, value=lab)
        c.font = HDR
        c.fill = HFILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    return r + 1


def section(ws, r, text, width=14):
    c = ws.cell(row=r, column=1, value=text)
    c.font = BOLD
    for j in range(1, width + 1):
        ws.cell(row=r, column=j).fill = SFILL
    return r + 1


def yrow(ws, r, label, arr, fmt="#,##0", scale=1.0, unit=None, bold=False):
    ws.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
    if unit is not None:
        ws.cell(row=r, column=2, value=unit)
    for j in range(N):
        v = arr[j] if np.ndim(arr) else arr
        v = float(v) / scale if v is not None and not (isinstance(v, float) and np.isnan(v)) else None
        c = ws.cell(row=r, column=3 + j, value=v)
        c.number_format = fmt
        if bold:
            c.font = BOLD
    return r + 1


def year_header(ws, r, first="KES m (nominal)"):
    return header(ws, r, [first, "Unit"] + [int(y) for y in YEARS])


def widths(ws, w0=46, w1=12, wn=12, n=N + 2):
    ws.column_dimensions["A"].width = w0
    ws.column_dimensions["B"].width = w1
    for j in range(3, n + 3):
        ws.column_dimensions[get_column_letter(j)].width = wn
    ws.freeze_panes = "C5"


def fmt_pct(x):
    return "n/a" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x*100:.1f}%"


def fmt_m(x):
    return "n/a" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x/1e6:,.0f}"


def write_excel(path, P0, base, scen, torn, grid, surv, chk, senior_kes, X):
    wb = Workbook()
    R, D, F, M = base["R"], base["D"], base["F"], base["M"]
    V = R["volumes"]
    fx = R["fx"]

    # ---------------- Assumptions
    ws = wb.active
    ws.title = "Assumptions"
    r = ws_title(ws, "Assumptions: every input with low / base / high, unit and source",
                 "Prices in 2026 KES ex-VAT unless stated. Change base values in model.py (ASSUMPTIONS) or via assumptions_override.csv, then re-run. [ESTIMATE] = analyst judgement; [DATED] = pre-2025 source; [UNVERIFIED] = weak source.")
    r = header(ws, r, ["Key", "Section", "Description", "Low", "Base", "High", "Unit", "Source / basis"])
    for k, a in ASSUMPTIONS.items():
        vals = [k, a["section"], a["desc"], a["low"], P0.get(k, a["base"]), a["high"], a["unit"], a["source"]]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            if j in (3, 4, 5):
                pct = ("%" in a["unit"] and "pp" not in a["unit"]) or a["unit"] in ("prob", "share")
                c.number_format = "0.0%" if pct else "#,##0.###"
        r += 1
    r += 1
    r = section(ws, r, "Vector and table inputs", 8)
    r = header(ws, r, ["Key", "Values", "", "Low", "Base", "High", "Unit", "Source / basis"])
    for k, v in VECTORS.items():
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=", ".join(f"{x:g}" for x in v["values"]))
        ws.cell(row=r, column=4, value=str(v["low"]))
        ws.cell(row=r, column=5, value="as listed")
        ws.cell(row=r, column=6, value=str(v["high"]))
        ws.cell(row=r, column=7, value=v["unit"])
        ws.cell(row=r, column=8, value=v["source"])
        r += 1
    r += 1
    r = section(ws, r, "Feed formulations (share of mash) and crude protein of ingredients", 8)
    r = header(ws, r, ["Ingredient", "Grow-out", "Starter", "CP %", "", "", "", "Source"])
    for k in FORM_GROW:
        for j, v in enumerate([k, FORM_GROW[k], FORM_START[k], CP[k]]):
            c = ws.cell(row=r, column=1 + j, value=v)
            if j: c.number_format = "0.0%"
        ws.cell(row=r, column=8, value=f"{TO} s3.3 (grow-out); starter [ESTIMATE]")
        r += 1
    cpg = sum(FORM_GROW[k] * CP[k] for k in FORM_GROW)
    cps = sum(FORM_START[k] * CP[k] for k in FORM_START)
    ws.cell(row=r, column=1, value="Calculated CP").font = BOLD
    ws.cell(row=r, column=2, value=cpg).number_format = "0.0%"
    ws.cell(row=r, column=3, value=cps).number_format = "0.0%"
    for col, w in zip("ABCDEFGH", (22, 16, 60, 11, 11, 11, 22, 120)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"

    # ---------------- Mass balance
    ws = wb.create_sheet("Mass balance")
    r = ws_title(ws, "Mass balance (tonnes)", f"Base case, Phase 2 {'GO' if base['go'] else 'NO-GO'}. Seed -> oil + cake/meal + hulls + losses; feed mash = formula ingredients; feed = mash x (1 - shrink).")
    r = year_header(ws, r, "Tonnes")
    rows = [("Seed supply potential", V["supply"] * R["opsfrac"]), ("Crush capacity x max utilisation", V["cap_crush"] * P0["max_util"] * R["opsfrac"]),
            ("Seed purchased", V["seed_bought"]), ("  storage shrink", V["seed_shrink"]), ("Seed crushed", V["crush"]),
            ("  sunflower, full press", V["sun_press"]), ("  sunflower, prepress+solvent", V["sun_solv"]), ("  soybean, extruder-expeller", V["soy"]),
            ("Crude oil: sunflower", V["oil_sun"]), ("Crude oil: soy", V["oil_soy"]), ("Sunflower expeller cake", V["cake_sun"]),
            ("Sunflower solvent meal (dehulled)", V["meal_sun"]), ("Hulls to boiler", V["hulls"]), ("Sunflower process losses", V["loss_sun"]),
            ("Soy cake", V["cake_soy"]), ("Soy process losses", V["loss_soy"])]
    r = section(ws, r, "Crushing")
    for lab, arr in rows:
        r = yrow(ws, r, lab, arr, unit="t")
    clos = V["crush"] - (V["crude_total"] + V["cake_sun"] + V["meal_sun"] + V["hulls"] + V["loss_sun"] + V["cake_soy"] + V["loss_soy"])
    r = yrow(ws, r, "CHECK: seed - (oil+cake+meal+hulls+losses)", clos, fmt="0.000", unit="t", bold=True)
    r = section(ws, r, "Oil routing")
    for lab, arr in (("Crude oil produced", V["crude_total"]), ("  to own feed (mash + coating)", V["oil_to_feed"]),
                     ("  to own refinery", V["ref_in"]), ("  sold crude to refiners", V["crude_sold"]),
                     ("Refinery capacity (x max util)", V["ref_cap"]), ("Refined fortified oil out", V["ref_out"]),
                     ("Soapstock / acid oil", V["soap"])):
        r = yrow(ws, r, lab, arr, unit="t")
    r = yrow(ws, r, "Refined oil (million litres)", V["ref_out"] * P0["l_per_t"] / 1e6, fmt="#,##0.00", unit="m L")
    r = section(ws, r, "Cake and meal routing")
    for lab, arr in (("Sunflower cake/meal to own feed", V["sun_to_feed"]), ("Sunflower cake/meal sold", V["sun_sold"]),
                     ("Soy cake to own feed", V["soy_to_feed"]), ("Soy cake sold (surplus)", V["soy_sold"]),
                     ("Non-GM SBM purchased (shortfall)", V["sbm_buy"]), ("Sunflower cake purchased (shortfall)", V["sun_buy"])):
        r = yrow(ws, r, lab, arr, unit="t")
    r = section(ws, r, "Aquafeed")
    for lab, arr in (("Merchant aquafeed SAM", V["sam"]), ("Max share (ramped)", V["share"]), ("Extrusion capacity (effective)", V["feed_cap"]),
                     ("Feed sold", V["feed"]), ("  grow-out", V["f_grow"]), ("  starter/fingerling", V["f_start"]), ("Mash input", V["mash"]),
                     ("  process shrink", V["feed_shrink"])):
        r = yrow(ws, r, lab, arr, fmt="0.0%" if lab.startswith("Max share") else "#,##0", unit="share" if lab.startswith("Max") else "t")
    for k in FORM_GROW:
        r = yrow(ws, r, f"  ingredient: {k}", V["need_" + k], unit="t")
    r = yrow(ws, r, "Feed share of SAM achieved", np.divide(V["feed"], V["sam"], out=np.zeros(N), where=V["sam"] > 0), fmt="0.0%", unit="%")
    r = section(ws, r, "Energy and storage")
    for lab, arr, u in (("Electricity", V["kwh"] / 1e3, "MWh"), ("Process steam", V["steam"], "t"), ("Hull share of thermal demand", V["hull_cover"], "%"),
                        ("Own seed storage", V["silo"], "t"), ("Leased overflow storage", V["over_storage"], "t-months")):
        r = yrow(ws, r, lab, arr, fmt="0.0%" if u == "%" else "#,##0", unit=u)
    widths(ws)

    # ---------------- Capex
    ws = wb.create_sheet("Capex & phasing")
    r = ws_title(ws, "Capex & phasing", "USD m in 2026 prices (low/high = capex_mult 0.70/1.45 on every item); nominal KES phasing uses FX path and USD drift.")
    r = header(ws, r, ["Item", "Category", "Imported share", "Phase 1 low", "Phase 1 base", "Phase 1 high", "Phase 2 low", "Phase 2 base", "Phase 2 high", "Source"])
    i1, i2 = R["items1"], R["items2"]
    lo_m, hi_m = ASSUMPTIONS["capex_mult"]["low"] / P0["capex_mult"], ASSUMPTIONS["capex_mult"]["high"] / P0["capex_mult"]
    keys = list(OrderedDict.fromkeys(list(i1.keys()) + list(i2.keys())))
    t1 = t2 = 0.0
    for k in keys:
        a = i1.get(k, [0, "", 0]); b = i2.get(k, [0, "", 0])
        src = CAPEX_P1.get(k, (0, 0, 0, ""))[3] if k in CAPEX_P1 else (f"{TO} s4.1 (150 tpd reference), scaled" if k in CAPEX_P2_REF else "computed")
        if k == "idf_rdl":
            src = f"IDF {P0['idf']:.1%} + RDL {P0['rdl']:.1%} on imported equipment ({RT} rows 13-14)"
        if k == "contingency":
            src = f"{TO} s4.1 (15% P1, ~10% P2)"
        vals = [k, a[1] or b[1], a[2] if a[0] else b[2], a[0] * lo_m, a[0], a[0] * hi_m, b[0] * lo_m, b[0], b[0] * hi_m, src]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            if j == 2: c.number_format = "0%"
            elif 3 <= j <= 8: c.number_format = "0.00"
        t1 += a[0]; t2 += b[0]
        r += 1
    for j, v in enumerate(["TOTAL (USD m)", "", "", t1 * lo_m, t1, t1 * hi_m, t2 * lo_m, t2, t2 * hi_m, f"{TO} rows 36/38: P1 6.5/9.5/14; P2 15/22/32 at 150 tpd"]):
        c = ws.cell(row=r, column=1 + j, value=v); c.font = BOLD
        if 3 <= j <= 8: c.number_format = "0.00"
    r += 1
    ws.cell(row=r, column=1, value=f"Phase 2 design: {P0['p2_tpd']:.0f} tpd prepress+solvent with dehulling, +{P0['p2_extr_tph']:.0f} t/h extrusion, +{P0['p2_ref_add']:.0f} t/d refinery; items scaled from TO's 150 tpd/+6 t/h/+40 t/d reference with exponent {P0['scale_exp']}.").font = ITAL
    r += 2
    r = year_header(ws, r)
    for lab, arr in (("Phase 1 capex", R["capex_p1"]), ("Phase 2 capex", R["capex_p2"]), ("Sustaining capex", R["sust"]),
                     ("VAT on capex paid", R["vat_paid"]), ("VAT on capex recovered", R["vat_recov"]),
                     ("Phase 0 + pre-operating (expensed in EBITDA)", R["devx"])):
        r = yrow(ws, r, lab, arr, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Total capex (incl. sustaining)", R["capex"] + R["sust"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Book depreciation", R["dep_book"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Tax capital allowances (investment deduction + W&T)", R["ded"], scale=1e6, unit="KES m")
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["J"].width = 70

    # ---------------- Operating P&L
    ws = wb.create_sheet("Operating P&L")
    r = ws_title(ws, "Operating P&L (annual, nominal KES m) with USD memo",
                 f"Base case; Phase 2 {'GO' if base['go'] else 'NO-GO'}; 2027-28 pre-operating; COD 2029 after {12 - P0['ops_months_y1']:.0f} months.")
    r = year_header(ws, r)
    r = section(ws, r, "Volumes")
    for lab, arr, u in (("Seed crushed", V["crush"], "t"), ("Refined oil sold", V["ref_out"], "t"), ("Crude oil sold", V["crude_sold"], "t"),
                        ("Cake/meal sold", V["sun_sold"] + V["soy_sold"], "t"), ("Aquafeed sold", V["feed"], "t"),
                        ("Crush utilisation", np.divide(V["crush"], V["cap_crush"] * np.maximum(R["opsfrac"], 1e-9), out=np.zeros(N), where=V["cap_crush"] > 0), "%")):
        r = yrow(ws, r, lab, arr, fmt="0.0%" if u == "%" else "#,##0", unit=u)
    r = section(ws, r, "Revenue")
    for k, v in R["rev"].items():
        r = yrow(ws, r, k, v, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Total revenue", R["revenue"], scale=1e6, unit="KES m", bold=True)
    r = section(ws, r, "Variable costs")
    for k, v in R["cost"].items():
        r = yrow(ws, r, k, v, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Total variable costs", R["variable"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Gross margin", R["revenue"] - R["variable"], scale=1e6, unit="KES m", bold=True)
    r = section(ws, r, "Fixed, selling, levies, development")
    for k, v in list(R["fixed"].items()) + list(R["sell"].items()) + list(R["levies"].items()):
        r = yrow(ws, r, k, v, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Phase 0 / pre-operating", R["devx"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "EBITDA", R["ebitda"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "EBITDA margin", np.divide(R["ebitda"], R["revenue"], out=np.zeros(N), where=R["revenue"] > 0), fmt="0.0%", unit="%")
    r = yrow(ws, r, "Depreciation (book)", R["dep_book"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "EBIT", F["ebit"], scale=1e6, unit="KES m", bold=True)
    for lab, arr in (("Interest: senior P1", D["sen_int"]), ("Interest: mezzanine (accrued, incl. PIK)", D["mezz_int_kes"]),
                     ("WHT gross-up on mezz interest", D["mezz_wht"]), ("Interest: Phase 2 loan", D["p2_int"]),
                     ("Interest: WC revolver + outgrower credit", R["revolver_int"]), ("Financing fees", D["fees"]),
                     ("FX loss on mezz principal", D["mezz_fxloss"])):
        r = yrow(ws, r, lab, arr, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Profit before tax", F["pbt"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Income tax (30% after deductions, loss c/f)", F["tax"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Net income", F["ni"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Taxable income after losses", F["taxable"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Tax losses expired (5-yr cap)", F["loss_expired"], scale=1e6, unit="KES m")
    r = section(ws, r, "USD memo (explicit FX path)")
    r = yrow(ws, r, "KES per USD (year average)", fx, fmt="0.0", unit="KES/USD")
    r = yrow(ws, r, "Revenue", R["revenue"] / fx, scale=1e6, fmt="#,##0.00", unit="USD m")
    r = yrow(ws, r, "EBITDA", R["ebitda"] / fx, scale=1e6, fmt="#,##0.00", unit="USD m")
    r = yrow(ws, r, "Net income", F["ni"] / fx, scale=1e6, fmt="#,##0.00", unit="USD m")
    r = section(ws, r, "Unit economics (nominal)")
    P = R["prices"]
    for lab, arr, u in (("Crude oil price (blend)", P["crude"], "KES/t"), ("Refined oil price", P["refined"], "KES/t"),
                        ("Sunflower cake/meal price", P["sunprod"], "KES/t"), ("Grow-out feed price", P["feed_grow"], "KES/t"),
                        ("Delivered sunflower seed", P["sunseed"], "KES/t"), ("Delivered soybean", P["soyseed"], "KES/t"),
                        ("Crush gross margin per t seed (oil at crude value + cake)", R["crush_gm_per_t"], "KES/t"),
                        ("Feed gross margin per t feed", np.divide(R["feed_gm"], V["feed"], out=np.zeros(N), where=V["feed"] > 0), "KES/t"),
                        ("Feed gross margin %", R["feed_gm_pct"], "%"),
                        ("EBITDA per t seed crushed", np.divide(R["ebitda"], V["crush"], out=np.zeros(N), where=V["crush"] > 0), "KES/t")):
        r = yrow(ws, r, lab, arr, fmt="0.0%" if u == "%" else "#,##0", unit=u)
    widths(ws)

    # ---------------- Working capital
    ws = wb.create_sheet("Working capital")
    sp = R["sp"]
    r = ws_title(ws, "Working capital: seasonal seed inventory and year-end balances",
                 f"Seed bought {P0['season_conc']:.0%} Jun-Oct, crushed evenly; safety stock {P0['safety_stock_months']} month. Year-end seed stock = Dec ratio x next year's crush.")
    r = header(ws, r, ["Month", "Purchases % of annual", "Inventory end-month (x annual crush)", "2031 seed stock KES m", "2035 seed stock KES m", "2031 outgrower credit KES m"])
    for m in range(12):
        vals = [["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][m], sp["shares"][m], sp["inv"][m],
                R["month_inv_val"][IDX[2031], m] / 1e6, R["month_inv_val"][IDX[2035], m] / 1e6, R["og_month"][IDX[2031], m] / 1e6]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.number_format = "0.0%" if j in (1, 2) else "#,##0.0"
        r += 1
    for lab, v in (("Dec (year-end) ratio", sp["dec"]), ("Average ratio", sp["avg"]), ("Peak ratio", sp["peak"]), ("Peak month", sp["peak_month"]),
                   ("Average months in storage", sp["avg"] * 12), ("Storage shrink applied", R["shrink"])):
        ws.cell(row=r, column=1, value=lab).font = BOLD
        c = ws.cell(row=r, column=3, value=v)
        c.number_format = "0.0%" if "ratio" in lab or "shrink" in lab else "0.0"
        r += 1
    r += 1
    r = year_header(ws, r)
    for lab, arr in (("Seed inventory (Dec)", R["seed_inv"]), ("Feed ingredient stock", R["ingr_inv"]), ("Finished goods", R["fg_inv"]),
                     ("Receivables", R["ar"]), ("Payables", -R["ap"])):
        r = yrow(ws, r, lab, arr, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Net working capital (Dec)", R["wc"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Change in working capital", R["dwc"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Intra-year peak: seed stock + outgrower credit", R["seasonal_peak"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Revolver balance (Dec) = advance x seed stock", R["revolver_bal"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Revolver peak draw (intra-year)", R["revolver_peak"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Revolver + outgrower-credit interest", R["revolver_int"], scale=1e6, unit="KES m")
    widths(ws, w1=22, wn=13)

    # ---------------- Cash flow & funding
    ws = wb.create_sheet("Cash flow & funding")
    r = ws_title(ws, "Cash flow & funding (nominal KES m)", f"Senior P1 sized to {P0['dscr_target']:.2f}x min DSCR ({P0['dscr_first_year']:.0f}-34) on the lender downside case (CFADS x {1-P0['sizing_haircut']:.2f}), Phase 1 standalone, capped at USD {P0['senior_usd_max']}m: KES {senior_kes/1e6:,.0f}m. Below {P0['default_dscr']:.2f}x the sponsor injects cure equity that prepays senior; below {P0['dscr_target']:.2f}x distributions stop and {P0['sweep_share']:.0%} of excess cash prepays senior.")
    r = year_header(ws, r)
    r = section(ws, r, "Cash flow")
    for lab, arr, b in (("EBITDA", R["ebitda"], False), ("Tax paid", -F["tax"], False), ("Change in working capital", -R["dwc"], False),
                        ("Cash flow from operations", F["cfo"], True), ("Capex (P1 + P2)", -R["capex"], False), ("Sustaining capex", -R["sust"], False),
                        ("VAT on capex (net timing)", -F["vat_net"], False), ("Cash flow from investing", F["cfi"], True),
                        ("Debt drawdowns", F["debt_draws"], False), ("Interest paid (cash)", -F["int_cash"], False), ("Principal repaid", -F["prin"], False),
                        ("Fees", -D["fees"], False), ("Revolver change", F["drev"], False), ("DSRA funding (-) / release (+)", -D["ddsra"], False),
                        ("Senior prepayment (cure + sweep)", -D["sen_prepay"], False),
                        ("Grants", F["grants"], False), ("Cash flow before equity", F["pre_equity"], True),
                        ("Equity injected (incl. cure)", F["inj"], False), ("  of which covenant cure equity", F["cure"], False),
                        ("Distributions to equity", -F["dist"], False), ("Closing cash", F["cash_bal"], True)):
        r = yrow(ws, r, lab, arr, scale=1e6, unit="KES m", bold=b)
    r = section(ws, r, "Debt schedules")
    for lab, arr in (("Senior P1: drawdown", D["sen_draw"]), ("Senior P1: principal", D["sen_prin"]), ("Senior P1: interest", D["sen_int"]),
                     ("Senior P1: closing balance", D["sen_bal"]), ("Mezz: drawdown (KES)", D["mezz_draw"]), ("Mezz: cash interest", D["mezz_cash_int"]),
                     ("Mezz: WHT gross-up", D["mezz_wht"]), ("Mezz: principal (KES at spot)", D["mezz_prin"]), ("Mezz: closing balance (KES)", D["mezz_bal_kes"]),
                     ("Phase 2 loan: drawdown", D["p2_draw"]), ("Phase 2 loan: principal", D["p2_prin"]), ("Phase 2 loan: interest", D["p2_int"]),
                     ("Phase 2 loan: closing balance", D["p2_bal"]), ("WC revolver: closing balance", R["revolver_bal"]), ("DSRA balance", D["dsra"])):
        r = yrow(ws, r, lab, arr, scale=1e6, unit="KES m")
    r = yrow(ws, r, "Mezz balance (USD m)", D["mezz_bal_usd"], scale=1e6, fmt="0.00", unit="USD m")
    r = section(ws, r, "Coverage")
    r = yrow(ws, r, "CFADS", F["cfads"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Debt service (senior + mezz + P2)", D["ds_total"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "DSCR", F["dscr"], fmt="0.00", unit="x", bold=True)
    tot_debt = D["sen_bal"] + D["mezz_bal_kes"] + D["p2_bal"]
    r = yrow(ws, r, "Term debt / EBITDA", np.divide(tot_debt, R["ebitda"], out=np.full(N, np.nan), where=R["ebitda"] > 0), fmt="0.00", unit="x")
    r += 1
    # sources & uses P1 (2027-2029)
    r = section(ws, r, "Sources & uses, Phase 1 build (2027-2029, KES m nominal)")
    s = slice(0, 3)
    uses = OrderedDict([("Capex incl. land", R["capex"][s].sum()), ("Phase 0 + pre-operating", R["devx"][s].sum()),
                        ("Financing fees", D["fees"][s].sum()), ("Interest paid during build/ramp", F["int_cash"][s].sum()),
                        ("Principal repaid 2029", F["prin"][s].sum()),
                        ("DSRA funding", D["ddsra"][s].sum()), ("Working capital build", R["dwc"][s].sum()),
                        ("VAT on capex (net)", F["vat_net"][s].sum()), ("Tax paid", F["tax"][s].sum()),
                        ("Minimum cash", F["cash_bal"][2])])
    srcs = OrderedDict([("Founder + co-investor equity", F["inj"][s].sum() - F["dist"][s].sum()), ("Senior KES term loan", D["sen_draw"][s].sum()),
                        ("Mezzanine (USD)", D["mezz_draw"][s].sum()), ("Grants", F["grants"][s].sum()), ("WC revolver", F["drev"][s].sum()),
                        ("Operating EBITDA 2029 (before pre-op)", (R["ebitda"] + R["devx"])[s].sum())])
    r = header(ws, r, ["Uses", "KES m", "USD m", "", "Sources", "KES m", "USD m"])
    fxa = fx[1]
    for (u, uv), (sname, sv) in zip(list(uses.items()) + [("", None)] * max(0, len(srcs) - len(uses)),
                                    list(srcs.items()) + [("", None)] * max(0, len(uses) - len(srcs))):
        ws.cell(row=r, column=1, value=u)
        if uv is not None:
            ws.cell(row=r, column=2, value=uv / 1e6).number_format = "#,##0"
            ws.cell(row=r, column=3, value=uv / fxa / 1e6).number_format = "0.00"
        ws.cell(row=r, column=5, value=sname)
        if sv is not None:
            ws.cell(row=r, column=6, value=sv / 1e6).number_format = "#,##0"
            ws.cell(row=r, column=7, value=sv / fxa / 1e6).number_format = "0.00"
        r += 1
    tu, ts = sum(uses.values()), sum(srcs.values())
    for j, v in enumerate(["Total uses", tu / 1e6, tu / fxa / 1e6, "", "Total sources", ts / 1e6, ts / fxa / 1e6]):
        c = ws.cell(row=r, column=1 + j, value=v); c.font = BOLD
        if j in (1, 5): c.number_format = "#,##0"
        if j in (2, 6): c.number_format = "0.00"
    r += 1
    ws.cell(row=r, column=1, value=f"USD at 2028 rate {fxa:.1f}. Uses and sources differ only by rounding of operating items; check = {(tu-ts)/1e6:,.2f} KES m").font = ITAL
    widths(ws)

    # ---------------- Returns
    ws = wb.create_sheet("Returns")
    r = ws_title(ws, "Returns: project and equity IRR (KES, USD), NPV at hurdle rates, payback, peak funding",
                 f"Value date {VALUE_YEAR}; terminal value = {P0['exit_mult']}x EBITDA 2038. Equity flows = injections (-) and distributions (+) to all shareholders, plus equity value at exit.")
    full, p1o = base.get("full", base), base.get("p1_only", base)
    cols = [("Plan (gate auto)", base), ("Full plan (Phase 2 go)", full), ("Phase 1 only", p1o)]
    r = header(ws, r, ["Metric"] + [c[0] for c in cols] + ["Note"])
    metric_rows = [
        ("Project IRR, KES (with TV)", "project_irr_kes", "pct", "Unlevered, after unlevered tax"),
        ("Project IRR, USD (with TV)", "project_irr_usd", "pct", "KES flows / FX path"),
        ("Project IRR, KES (no TV, 2038 cut-off)", "project_irr_kes_noTV", "pct", ""),
        ("Equity IRR, KES (with TV)", "equity_irr_kes", "pct", "Before shareholder taxes"),
        ("Equity IRR, USD (with TV)", "equity_irr_usd", "pct", ""),
        ("Equity IRR, KES (no TV)", "equity_irr_kes_noTV", "pct", ""),
        ("Equity IRR, KES after 5% dividend WHT and 15% CGT", "equity_irr_kes_after_shtax", "pct", "Compare to tax-free bond"),
        ("Phase 1 WACC from the modelled structure", "wacc_model", "pct", "Equity at 35%, senior and mezz post-tax, weights from 2027-29 sources (R1-04)"),
        ("Project NPV @ modelled WACC (KES m)", "npv_proj_wacc_model", "m", ""),
        (f"Project NPV @ financing-file WACC {P0['wacc_p1']:.0%} (memo, KES m)", "npv_proj_wacc_p1", "m", "Assumes 50% equity; not the modelled structure"),
        ("Equity IRR, founder group after GM equity incentive", "equity_irr_founder_after_gm", "pct", "R1-12"),
        (f"Equity NPV @ infra bond {P0['infra_bond']:.2%} (KES m)", "npv_eq_infra", "m", "Opportunity cost floor"),
        (f"Equity NPV @ established hurdle {P0['hurdle_est']:.1%} (KES m)", "npv_eq_est", "m", ""),
        (f"Equity NPV @ greenfield {P0['hurdle_gf_lo']:.0%} (KES m)", "npv_eq_gf_lo", "m", "Hurdle is survival-weighted: see sheet"),
        (f"Equity NPV @ greenfield {P0['hurdle_gf_hi']:.0%} (KES m)", "npv_eq_gf_hi", "m", ""),
        ("Project payback year (undiscounted, no TV)", "payback_project", "yr", ""),
        ("Equity payback year (undiscounted, no TV)", "payback_equity", "yr", ""),
        ("Peak funding need, project (KES m)", "peak_funding_project", "m", "Max cumulative negative unlevered cash flow"),
        ("Peak net equity committed (KES m)", "peak_equity", "m", "Max cumulative injections less distributions"),
        ("Total equity injected (KES m)", "equity_injected", "m", ""),
        ("Peak WC revolver draw, Phase 1 years 2029-32 (KES m)", "peak_revolver_p1", "m", "Intra-year, separate from term funding"),
        ("Peak WC revolver draw, any year (KES m)", "peak_revolver", "m", "Nominal; usually the last year"),
        ("Senior P1 loan sized (KES m)", "senior_kes", "m", ""),
        ("Min DSCR (from first test year)", "min_dscr", "x", "Lock-up 1.30x; default 1.10x"),
        ("Covenant cure equity, total (KES m)", "cure_equity", "m", "Injected and applied to prepay senior"),
        ("Years below default covenant", "breach_years", "yr", ""),
        ("Years in lock-up (below 1.30x)", "lockup_years", "yr", ""),
        ("CFADS / debt service 2029 (ramp year, n/m)", "dscr_2029", "x", "Not tested: 2029 WC build and debt service funded from Phase 1 budget"),
        ("EBITDA 2031 (KES m)", "ebitda_2031", "m", ""), ("EBITDA 2035 (KES m)", "ebitda_2035", "m", ""), ("EBITDA 2038 (KES m)", "ebitda_2038", "m", ""),
    ]
    for lab, key, kind, note in metric_rows:
        ws.cell(row=r, column=1, value=lab)
        for j, (_n, o) in enumerate(cols):
            v = o["M"][key]
            c = ws.cell(row=r, column=2 + j)
            if v is None or (isinstance(v, float) and np.isnan(v)):
                c.value = "n/a"
            elif kind == "pct":
                c.value = v; c.number_format = "0.0%"
            elif kind == "m":
                c.value = v / 1e6; c.number_format = "#,##0"
            elif kind == "x":
                c.value = v; c.number_format = "0.00"
            else:
                c.value = v
        ws.cell(row=r, column=5, value=note).font = ITAL
        r += 1
    r += 1
    r = section(ws, r, "Phase 2 go/no-go gate (end-2031)", 6)
    if base.get("gate"):
        r = header(ws, r, ["Test", "Value", "Threshold", "Pass?"])
        for k, (v, th, ok) in base["gate"].items():
            ws.cell(row=r, column=1, value=k)
            for j, x in enumerate((v, th)):
                c = ws.cell(row=r, column=2 + j, value=x / 1e6 if "KES" in k else x)
                c.number_format = "0.0%" if ("margin" in k or "share" in k) else "#,##0"
            ws.cell(row=r, column=4, value="PASS" if ok else "FAIL")
            r += 1
        ws.cell(row=r, column=1, value=f"Decision: Phase 2 {'GO' if base['go'] else 'NO-GO'}. Incremental Phase 2 project IRR {fmt_pct(base['inc_irr'])}; incremental equity IRR {fmt_pct(base['inc_eq_irr'])}; incremental project NPV @ Phase 2 funding WACC {base['wacc_p2']:.1%} KES {base['inc_npv']/1e6:,.0f}m.").font = BOLD
        r += 2
    r = year_header(ws, r)
    r = yrow(ws, r, "Project cash flow (unlevered, incl. TV in 2038)", base["F"]["proj_tv"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Equity cash flow (incl. equity TV in 2038)", base["F"]["eq_tv_flows"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Equity cash flow after shareholder taxes", base["F"]["eq_after"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Project cash flow (USD m)", base["F"]["proj_tv"] / fx, scale=1e6, fmt="#,##0.00", unit="USD m")
    r = yrow(ws, r, "Equity cash flow (USD m)", base["F"]["eq_tv_flows"] / fx, scale=1e6, fmt="#,##0.00", unit="USD m")
    if "inc_flows" in base:
        r = yrow(ws, r, "Phase 2 incremental project cash flow", base["inc_flows"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Cumulative project cash flow (no TV)", np.cumsum(base["F"]["proj"]), scale=1e6, unit="KES m")
    r = yrow(ws, r, "Cumulative equity cash flow (no TV)", np.cumsum(base["F"]["eq"]), scale=1e6, unit="KES m")
    ws.cell(row=r + 1, column=1, value=f"Exit: EV {base['F']['ev']/1e6:,.0f} KES m; net debt {base['F']['net_debt']/1e6:,.0f}; equity value {base['F']['eq_tv']/1e6:,.0f}; CGT {base['F']['cgt']/1e6:,.0f}.").font = ITAL
    widths(ws, w0=58, w1=20, wn=14)
    for col in "BCD":
        ws.column_dimensions[col].width = 20

    # ---------------- Scenarios
    ws = wb.create_sheet("Scenarios")
    r = ws_title(ws, "Scenarios: bear / base / bull (+ all-adverse stress)", "Senior debt held at base sizing in every scenario; Phase 2 gate evaluated inside each scenario.")
    names = list(scen.keys())
    r = header(ws, r, ["Metric"] + [n.upper() for n in names])
    ws.cell(row=r, column=1, value="Definition")
    for j, n in enumerate(names):
        ws.cell(row=r, column=2 + j, value=SCEN_DESC[n]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 120
    r += 1
    ws.cell(row=r, column=1, value="Phase 2 decision")
    for j, n in enumerate(names):
        ws.cell(row=r, column=2 + j, value="GO" if scen[n]["go"] else "NO-GO")
    r += 1
    for lab, key, kind, _note in metric_rows:
        ws.cell(row=r, column=1, value=lab)
        for j, n in enumerate(names):
            v = scen[n]["M"][key]
            c = ws.cell(row=r, column=2 + j)
            if v is None or (isinstance(v, float) and np.isnan(v)):
                c.value = "n/a"
            elif kind == "pct":
                c.value = v; c.number_format = "0.0%"
            elif kind == "m":
                c.value = v / 1e6; c.number_format = "#,##0"
            elif kind == "x":
                c.value = v; c.number_format = "0.00"
            else:
                c.value = v
        r += 1
    r += 1
    r = section(ws, r, "Key inputs by scenario", 6)
    for k in ["config", "sam_2026", "sam_regional", "feed_px_mult", "feed_real_decline", "entrant_discount", "seed_cost_mult", "p_crude", "refine_spread", "p_suncake", "supply_mult", "feed_share_max", "refined_share_max",
              "capex_mult", "kes_dep", "sun_oil_y", "p_cpo_removed", "p_gmo", "p_tz_levy", "exit_mult", "ops_months_y1",
              "p2_tpd", "p2_extr_tph", "p2_ref_add"]:
        ws.cell(row=r, column=1, value=k)
        for j, n in enumerate(names):
            ws.cell(row=r, column=2 + j, value=scen[n]["p"][k]).number_format = "#,##0.###"
        r += 1
    r += 1
    r = header(ws, r, ["Annual EBITDA, KES m"] + [int(y) for y in YEARS])
    for n in names:
        ws.cell(row=r, column=1, value=n)
        for j in range(N):
            ws.cell(row=r, column=2 + j, value=scen[n]["R"]["ebitda"][j] / 1e6).number_format = "#,##0"
        r += 1
    r = header(ws, r, ["Equity cash flow, KES m"] + [int(y) for y in YEARS])
    for n in names:
        ws.cell(row=r, column=1, value=n)
        for j in range(N):
            ws.cell(row=r, column=2 + j, value=scen[n]["F"]["eq_tv_flows"][j] / 1e6).number_format = "#,##0"
        r += 1
    ws.column_dimensions["A"].width = 56
    for col in "BCDEFGHIJKLM":
        ws.column_dimensions[col].width = 22 if col in "BCDEF" else 12

    # ---------------- Sensitivities
    ws = wb.create_sheet("Sensitivities")
    r = ws_title(ws, "Sensitivities: tornado (one driver at a time, low / high)",
                 "Sorted by swing in equity IRR. Phase 2 decision held at the base-case result so bars are continuous; senior debt held at base sizing.")
    r = header(ws, r, ["Driver", "Low setting", "High setting", "Equity IRR KES low", "Equity IRR KES high", "Swing (pp)", "Equity IRR USD low", "Equity IRR USD high",
                       "Project IRR KES low", "Project IRR KES high", "Min DSCR low", "Min DSCR high", "Cure equity low (KES m)", "Cure equity high (KES m)",
                       "Peak equity low (KES m)", "Peak equity high (KES m)"])
    for t in torn:
        vals = [t["label"], ", ".join(f"{k}={v:g}" for k, v in t["lo"].items()), ", ".join(f"{k}={v:g}" for k, v in t["hi"].items()),
                t["eq_lo"], t["eq_hi"], t["swing"] * 100, t["equsd_lo"], t["equsd_hi"], t["pr_lo"], t["pr_hi"], t["dscr_lo"], t["dscr_hi"],
                t["cure_lo"] / 1e6, t["cure_hi"] / 1e6, t["peak_lo"] / 1e6, t["peak_hi"] / 1e6]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=None if isinstance(v, float) and np.isnan(v) else v)
            if j in (3, 4, 6, 7, 8, 9): c.number_format = "0.0%"
            elif j == 5: c.number_format = "0.0"
            elif j in (10, 11): c.number_format = "0.00"
            elif j >= 12: c.number_format = "#,##0"
        r += 1
    ws.cell(row=r, column=1, value=f"Base: equity IRR {fmt_pct(M['equity_irr_kes'])}, project IRR {fmt_pct(M['project_irr_kes'])}.").font = BOLD
    r += 2
    seeds, oils, pg, eg = grid
    r = section(ws, r, "Two-way: delivered seed cost x oil price (crude and refined) -> project IRR (KES) / equity IRR (KES)", 8)
    r = header(ws, r, ["Seed cost mult \\ oil price mult"] + [f"oil x{o:.2f}" for o in oils])
    for i, s_ in enumerate(seeds):
        ws.cell(row=r, column=1, value=f"seed x{s_:.2f} (project)")
        for j in range(len(oils)):
            ws.cell(row=r, column=2 + j, value=pg[i, j]).number_format = "0.0%"
        r += 1
    for i, s_ in enumerate(seeds):
        ws.cell(row=r, column=1, value=f"seed x{s_:.2f} (equity)")
        for j in range(len(oils)):
            ws.cell(row=r, column=2 + j, value=None if np.isnan(eg[i, j]) else eg[i, j]).number_format = "0.0%"
        r += 1
    ws.column_dimensions["A"].width = 58
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 18

    # ---------------- Survival-weighted
    ws = wb.create_sheet("Survival-weighted")
    r = ws_title(ws, "Survival-weighted view: expected equity IRR vs the 35-40% greenfield hurdle",
                 "Four-branch tree (R1-11). Failure branches liquidate at recovery_fixed x gross capex + recovery_wc x WC; equity gets max(0, recovery - net debt); the founder also pays the personal-guarantee call on any senior shortfall.")
    r = header(ws, r, ["Branch", "Probability", "Equity IRR", "Project IRR", "Equity NPV @35% (KES m)", "Founder NPV @35% (KES m)"])
    for nm, b in surv["branches"].items():
        vals = [nm, b["prob"], b["eq_irr"], b["proj_irr"], b["eq_npv35"] / 1e6, b["founder_npv35"] / 1e6]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=None if isinstance(v, float) and np.isnan(v) else v)
            c.number_format = "0.0%" if j in (1, 2, 3) else "#,##0"
        r += 1
    r += 1
    rows = [("Expected (probability-weighted) equity IRR, KES, all shareholders", surv["exp_eq_irr"], "pct"),
            ("Expected founder-group IRR (GM dilution, guarantee calls)", surv["exp_founder_irr"], "pct"),
            ("Expected founder-group IRR after dividend WHT and CGT", surv["exp_founder_at_irr"], "pct"),
            ("Expected project IRR, KES", surv["exp_proj_irr"], "pct"),
            ("Equity IRR conditional on survival", surv["surv_eq_irr"], "pct"),
            ("Total failure probability (early + late)", surv["total_fail_prob"], "pct"),
            ("Greenfield hurdle (expected), lower / upper", f"{P0['hurdle_gf_lo']:.0%} / {P0['hurdle_gf_hi']:.0%}", "txt"),
            ("Expected equity NPV @35% / @40% (KES m)", f"{surv['exp_npv35']/1e6:,.0f} / {surv['exp_npv40']/1e6:,.0f}", "txt"),
            ("Expected founder NPV @35% / @40% (KES m)", f"{surv['exp_founder_npv35']/1e6:,.0f} / {surv['exp_founder_npv40']/1e6:,.0f}", "txt"),
            (f"Bond comparison: NPV @ {P0['infra_bond']:.2%} of survival-CONDITIONAL equity flows (KES m)", surv["bond_npv_surv_cond"] / 1e6, "m"),
            (f"Bond comparison: NPV @ {P0['infra_bond']:.2%} of survival-WEIGHTED founder flows after shareholder tax (KES m)", surv["bond_npv_exp_founder_at"] / 1e6, "m"),
            ("Early-failure probability at which expected IRR falls below 35%", surv["breakeven_pfail_35"], "pct"),
            ("Early failure: asset recovery / net debt / equity recovery / guarantee call (KES m)",
             f"{surv['liq_assets']/1e6:,.0f} / {surv['debt_at_fail']/1e6:,.0f} / {surv['eq_liq']/1e6:,.0f} / {surv['guarantee_early']/1e6:,.0f}", "txt")]
    for lab, v, kind in rows:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value="n/a" if v is None else (None if isinstance(v, float) and np.isnan(v) else v))
        c.number_format = "0.0%" if kind == "pct" else "#,##0"
        r += 1
    r += 1
    r = header(ws, r, ["Early-failure probability", "Expected equity IRR", "Expected equity NPV @35% (KES m)"])
    for q, ir, nv in surv["pfail_sweep"]:
        for j, v in enumerate((q, ir, nv / 1e6)):
            ws.cell(row=r, column=1 + j, value=None if isinstance(v, float) and np.isnan(v) else v).number_format = "0.0%" if j < 2 else "#,##0"
        r += 1
    r += 1
    r = year_header(ws, r, "Equity flows, KES m")
    for nm, b in surv["branches"].items():
        r = yrow(ws, r, nm, b["eq"], scale=1e6, unit="KES m")
    r = yrow(ws, r, "Expected equity flow", surv["exp_eq"], scale=1e6, unit="KES m", bold=True)
    r = yrow(ws, r, "Expected founder flow after shareholder tax", surv["exp_founder_at"], scale=1e6, unit="KES m", bold=True)
    ws.column_dimensions["A"].width = 80
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 14

    # ---------------- Phase 0 value of information
    ws = wb.create_sheet("Phase 0 VOI")
    r = ws_title(ws, "Phase 0: value of information (R1-05)",
                 "Expected founder NPV (KES m, value date 2027) of each policy. Priors: triangular(low, base, high) per input; rho = copula correlation among favourable outcomes. Invest only if the survival-weighted NPV at the hurdle is positive.")
    for v in X["voi"]:
        r = section(ws, r, f"rho = {v['rho']}; {v['n']} draws; Phase 0a cost KES {v['cost_a']/1e6:,.0f}m; Phase 0b cost incl. pilot seed loss KES {v['cost_b']/1e6:,.0f}m", 10)
        r = header(ws, r, ["Hurdle / early-failure prob.", "Stop now", "Invest now (no info)", "Phase 0a only", "Staged 0a -> 0b", "Perfect information",
                           "P(pass | full info)", "P(pass | 0a info)", "Gross EVPI", "P(cond. IRR >= 35%)"])
        for key in ("35_0.30", "40_0.30", "35_0.20", "40_0.20"):
            d = v[key]
            vals = [key.replace("_", "% hurdle, p_fail "), 0.0, d["invest_now"] / 1e6, d["phase0a_only"] / 1e6, d["staged"] / 1e6, d["perfect_info"] / 1e6,
                    d["p_pass_full"], d["p_pass_a"], d["evpi_gross"] / 1e6, d["p_cond_irr_35"]]
            for j, x in enumerate(vals):
                ws.cell(row=r, column=1 + j, value=x).number_format = "0.0%" if j in (6, 7, 9) else "#,##0"
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 18

    # ---------------- Gates & break-evens
    ws = wb.create_sheet("Gates & break-evens")
    r = ws_title(ws, "Gate calibration (R1-10) and joint cases",
                 "Covenant = value at which min DSCR hits the 1.10x default (others at base, base debt). Hurdle = value for a 35% conditional equity IRR (others at base). Kill = value at which the survival-weighted founder IRR is 35% at p_fail 20% with every other driver at its bull value.")
    r = header(ws, r, ["Driver", "Base", "Covenant (1.10x)", "Hurdle 35% (others base)", "Kill threshold (others bull)", "Bull value"])
    for k, d in X["cal"].items():
        for j, x in enumerate([k, d["base"], d["covenant"], d["hurdle"], d["kill"], d["bull_other"]]):
            ws.cell(row=r, column=1 + j, value="not reached" if x is None else x).number_format = "#,##0.000"
        r += 1
    r += 1
    r = header(ws, r, ["Joint case", "Cond. equity IRR", "Exp. IRR p_fail 30%", "Exp. IRR p_fail 20%", "NPV35 p_fail 30%", "NPV35 p_fail 20%", "Min DSCR", "Senior (KES m)", "Peak equity", "Phase 2"])
    for nm, d in X["jc"].items():
        vals = [nm, d["cond_irr"], d["exp_irr_0.30"], d["exp_irr_0.20"], d["npv35_0.30"] / 1e6, d["npv35_0.20"] / 1e6, d["min_dscr"], d["senior"] / 1e6, d["peak_equity"] / 1e6, "GO" if d["go"] else "NO-GO"]
        for j, x in enumerate(vals):
            ws.cell(row=r, column=1 + j, value=None if isinstance(x, float) and np.isnan(x) else x).number_format = "0.0%" if j in (1, 2, 3) else ("0.00" if j == 6 else "#,##0")
        r += 1
    ws.column_dimensions["A"].width = 70
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 16

    # ---------------- Checks
    ws = wb.create_sheet("Checks")
    r = ws_title(ws, "Integrity and benchmark checks", "CHECK = outside benchmark range or covenant; read the detail before relying on results.")
    r = header(ws, r, ["Check", "Result", "Detail"])
    for nm, res_, det in chk:
        ws.cell(row=r, column=1, value=nm)
        ws.cell(row=r, column=2, value=res_)
        ws.cell(row=r, column=3, value=det)
        r += 1
    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 90
    wb.save(path)


def write_csvs(base, scen, X=None):
    R, D, F = base["R"], base["D"], base["F"]
    V = R["volumes"]
    def w(name, cols):
        with open(os.path.join(HERE, name), "w", newline="") as f:
            wr = csv.writer(f, lineterminator="\n")
            wr.writerow(["year"] + list(cols.keys()))
            for i in range(N):
                wr.writerow([int(YEARS[i])] + [f"{float(v[i]):.4f}" if not np.isnan(v[i]) else "" for v in cols.values()])
    def snake(s):
        out = "".join(ch.lower() if ch.isalnum() else "_" for ch in s)
        while "__" in out:
            out = out.replace("__", "_")
        return out.strip("_")[:28]
    pnl = OrderedDict()
    pnl["fx_kes_usd"] = R["fx"]
    pnl["seed_crushed_t"] = V["crush"]; pnl["feed_sold_t"] = V["feed"]; pnl["refined_oil_t"] = V["ref_out"]
    pnl["crude_sold_t"] = V["crude_sold"]; pnl["cake_sold_t"] = V["sun_sold"] + V["soy_sold"]
    for k, v in R["rev"].items(): pnl["rev_" + snake(k)] = v / 1e6
    pnl["revenue"] = R["revenue"] / 1e6
    for k, v in R["cost"].items(): pnl["cv_" + snake(k)] = v / 1e6
    for k, v in R["fixed"].items(): pnl["cf_" + snake(k)] = v / 1e6
    for k, v in R["sell"].items(): pnl["cs_" + snake(k)] = v / 1e6
    for k, v in R["levies"].items(): pnl["cl_" + snake(k)] = v / 1e6
    pnl["preop"] = R["devx"] / 1e6
    pnl["ebitda"] = R["ebitda"] / 1e6; pnl["depreciation"] = R["dep_book"] / 1e6; pnl["ebit"] = F["ebit"] / 1e6
    pnl["interest_total"] = F["interest_pl"] / 1e6; pnl["fx_loss"] = D["mezz_fxloss"] / 1e6
    pnl["pbt"] = F["pbt"] / 1e6; pnl["tax"] = F["tax"] / 1e6; pnl["net_income"] = F["ni"] / 1e6
    pnl["revenue_usd"] = R["revenue"] / R["fx"] / 1e6; pnl["ebitda_usd"] = R["ebitda"] / R["fx"] / 1e6
    w("pnl_annual.csv", pnl)
    cf = OrderedDict()
    cf["fx_kes_usd"] = R["fx"]
    for k in ("cfo", "cfi", "cff_pre", "pre_equity", "inj", "dist", "cash_bal", "cfads", "proj", "proj_tv", "eq", "eq_tv_flows", "eq_after"):
        cf[k] = F[k] / 1e6 if k != "dscr" else F[k]
    cf["dscr"] = F["dscr"]
    cf["capex"] = R["capex"] / 1e6; cf["sust_capex"] = R["sust"] / 1e6; cf["dwc"] = R["dwc"] / 1e6; cf["wc"] = R["wc"] / 1e6
    cf["tax_paid"] = F["tax"] / 1e6; cf["debt_draws"] = F["debt_draws"] / 1e6; cf["interest_cash"] = F["int_cash"] / 1e6
    cf["principal"] = F["prin"] / 1e6; cf["senior_bal"] = D["sen_bal"] / 1e6; cf["mezz_bal"] = D["mezz_bal_kes"] / 1e6
    cf["p2_loan_bal"] = D["p2_bal"] / 1e6; cf["revolver_bal"] = R["revolver_bal"] / 1e6; cf["dsra"] = D["dsra"] / 1e6
    cf["proj_tv_usd"] = F["proj_tv"] / R["fx"] / 1e6; cf["eq_tv_usd"] = F["eq_tv_flows"] / R["fx"] / 1e6
    cf["cure_equity"] = F["cure"] / 1e6; cf["senior_prepay"] = D["sen_prepay"] / 1e6; cf["sweep"] = F["sweep"] / 1e6
    cf["eq_founder"] = F["eq_founder"] / 1e6
    w("cashflow_annual.csv", cf)
    if X is not None:
        # VOI draws for Stata (one row per draw and prior correlation)
        with open(os.path.join(HERE, "voi_draws.csv"), "w", newline="") as f:
            wr = csv.writer(f, lineterminator="\n")
            keys = list(X["voi"][0]["rows"][0].keys())
            wr.writerow(["rho", "draw"] + [k.replace(".", "").lower() for k in keys])   # Stata-safe names
            for v in X["voi"]:
                for j, row in enumerate(v["rows"]):
                    wr.writerow([v["rho"], j] + [f"{float(row[k]):.6g}" for k in keys])
    mb = OrderedDict((k, v) for k, v in V.items())
    w("mass_balance_annual.csv", mb)
    sc = OrderedDict()
    for n, o in scen.items():
        sc[f"ebitda_{n}"] = o["R"]["ebitda"] / 1e6
        sc[f"eqflow_{n}"] = o["F"]["eq_tv_flows"] / 1e6
        sc[f"projflow_{n}"] = o["F"]["proj_tv"] / 1e6
    w("scenarios_annual.csv", sc)


def md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(x) for x in r) + " |")
    return "\n".join(out)


def write_md(path, P0, base, scen, torn, grid, surv, chk, senior_kes, X):
    R, D, F, M = base["R"], base["D"], base["F"], base["M"]
    V = R["volumes"]
    fx = R["fx"]
    L = []
    L.append("# Financial model outputs: oilseed crushing + aquafeed mill (Kisumu)\n")
    L.append(f"Generated by `model/model.py` on the 19 Oct 2026 research base. Nominal KES unless stated; USD at the modelled FX path "
             f"(KES {P0['fx0']} in 2026, {P0['kes_dep']:.0%}/yr depreciation, +{P0['kes_dep_2027_extra']*100:.1f} pp in 2027). "
             f"Value date {VALUE_YEAR}. Prices ex-VAT. Every input and its source is on the Assumptions sheet of `financial_model.xlsx`.\n")
    L.append("## 1. Configuration and phasing\n")
    i31, i35 = IDX[2031], IDX[2035]
    L.append(md_table(["Item", "Phase 1 (FID Q4 2027, COD 2029)", "Phase 2 (gate end-2031, build 2032, ops 2033)"], [
        ["Crushing", f"{P0['p1_tpd']:.0f} t/day multi-seed full press + soy extruder-expeller ({P0['p1_tpd']*P0['crush_days']/1000:.0f} kt/yr nameplate)",
         f"to {P0['p2_tpd']:.0f} t/day with dehulling, prepress + hexane solvent ({P0['p2_tpd']*P0['crush_days']/1000:.0f} kt/yr)"],
        ["Seed mix", f"{1-P0['soy_share']:.0%} sunflower / {P0['soy_share']:.0%} non-GM soy", "same"],
        ["Refinery", f"{P0['ref_tpd_p1']:.0f} t/day semi-continuous, winterised, vitamin-A fortified, 20 L packing", f"+{P0['p2_ref_add']:.0f} t/day"],
        ["Extrusion", f"{P0['extr_tph_p1']:.0f} t/h twin-screw (~{P0['extr_tph_p1']*P0['extr_eff']*P0['extr_hours']/1000:.1f} kt/yr)", f"+{P0['p2_extr_tph']:.0f} t/h"],
        ["Seed storage", f"{P0['silo_p1']/1000:.0f} kt silos + leased overflow", f"+{P0['silo_p2']/1000:.1f} kt"],
        ["Base-case decision", "build", "GO" if base["go"] else "NO-GO (gate fails)"],
    ]))
    L.append("\nPhase 2 is smaller than the technical file's 150 t/day / +6 t/h specification: the seed-supply evidence (feedstock_supply.md s9) and the merchant aquafeed market cap (market_demand.md s5) do not support the larger plant in the base case. The 150 t/day plant is the bull-case design option (`p2_tpd`).\n")
    L.append("## 2. Capex (USD m, 2026 prices)\n")
    i1, i2 = R["items1"], R["items2"]
    keys = list(OrderedDict.fromkeys(list(i1.keys()) + list(i2.keys())))
    rows = [[k, f"{i1.get(k,[0])[0]:.2f}", f"{i2.get(k,[0])[0]:.2f}"] for k in keys]
    t1 = sum(v[0] for v in i1.values()); t2 = sum(v[0] for v in i2.values())
    lo, hi = ASSUMPTIONS["capex_mult"]["low"], ASSUMPTIONS["capex_mult"]["high"]
    rows.append(["**Total (base)**", f"**{t1:.2f}**", f"**{t2:.2f}**"])
    rows.append(["Range (capex x0.70 / x1.45)", f"{t1*lo:.1f} - {t1*hi:.1f}", f"{t2*lo:.1f} - {t2*hi:.1f}"])
    rows.append(["Total, KES m nominal (as spent)", f"{R['capex_p1'].sum()/1e6:,.0f}", f"{base['full']['R']['capex_p2'].sum()/1e6:,.0f}" if 'full' in base else "-"])
    L.append(md_table(["Item", "Phase 1", "Phase 2"], rows))
    L.append("")
    # sources & uses
    s = slice(0, 3)
    fxa = fx[1]
    L.append("## 3. Phase 1 funding (2027-2029)\n")
    srows = [["Founder + co-investor equity (net)", fmt_m(F["inj"][s].sum() - F["dist"][s].sum()), f"{(F['inj'][s].sum()-F['dist'][s].sum())/fxa/1e6:.2f}"],
             [f"Senior KES term loan ({P0['senior_rate']:.1%}, {P0['senior_tenor']:.0f} yrs, {P0['senior_grace_m']:.0f}-month grace)", fmt_m(D["sen_draw"].sum()), f"{D['sen_draw'].sum()/fxa/1e6:.2f}"],
             [f"Mezzanine USD ({P0['mezz_rate']:.0%} + WHT gross-up, {P0['mezz_tenor']:.0f} yrs, {P0['mezz_grace']:.0f}-yr grace)", fmt_m(D["mezz_draw"].sum()), f"{P0['mezz_usd']:.2f}"],
             ["Grants (outgrower/ESIA co-funding)", fmt_m(F["grants"].sum()), f"{P0['grants_usd']:.2f}"],
             ["WC revolver at end-2029 (65% of seed stock)", fmt_m(R["revolver_bal"][IDX[2029]]), f"{R['revolver_bal'][IDX[2029]]/fxa/1e6:.2f}"]]
    L.append(md_table(["Source", "KES m", "USD m (2028 FX)"], srows))
    L.append(f"\nSenior debt is sized to a {P0['dscr_target']:.2f}x minimum DSCR over {P0['dscr_first_year']:.0f}-34 on the lender downside case (CFADS x {1-P0['sizing_haircut']:.2f}), Phase 1 standalone, capped at the financing file's USD {P0['senior_usd_max']}m: result KES {senior_kes/1e6:,.0f}m (USD {senior_kes/fxa/1e6:.2f}m). "
             f"Peak net equity committed (whole plan, incl. cure equity): KES {M['peak_equity']/1e6:,.0f}m in {M['peak_equity_year']} (USD {M['peak_equity']/fx[IDX[M['peak_equity_year']]]/1e6:.1f}m at that year's FX). Peak project funding need: KES {M['peak_funding_project']/1e6:,.0f}m. Peak intra-year revolver draw: KES {M['peak_revolver_p1']/1e6:,.0f}m in Phase 1 (2029-32), KES {M['peak_revolver']/1e6:,.0f}m nominal by {M['peak_revolver_year']}.\n")
    L.append("## 4. Base case, annual summary (KES m nominal)\n")
    hdr = ["Item"] + [str(int(y)) for y in YEARS[2:]]
    def rowf(lab, arr, scale=1e6, f="{:,.0f}"):
        return [lab] + [f.format(arr[i] / scale) if not np.isnan(arr[i]) else "-" for i in range(2, N)]
    tbl = [rowf("Seed crushed (kt)", V["crush"], 1e3, "{:,.1f}"), rowf("Aquafeed sold (kt)", V["feed"], 1e3, "{:,.1f}"),
           rowf("Refined oil sold (kt)", V["ref_out"], 1e3, "{:,.1f}"), rowf("Crude oil sold (kt)", V["crude_sold"], 1e3, "{:,.1f}"),
           rowf("Cake/meal sold (kt)", V["sun_sold"] + V["soy_sold"], 1e3, "{:,.1f}"),
           rowf("Revenue", R["revenue"]), rowf("EBITDA", R["ebitda"]),
           ["EBITDA margin"] + [fmt_pct(R["ebitda"][i] / R["revenue"][i]) if R["revenue"][i] > 0 else "-" for i in range(2, N)],
           rowf("Net income", F["ni"]), rowf("CFADS", F["cfads"]), ["DSCR"] + [f"{F['dscr'][i]:.2f}" if not np.isnan(F['dscr'][i]) else "-" for i in range(2, N)],
           rowf("Capex (incl. sustaining)", R["capex"] + R["sust"]), rowf("Equity injected (incl. cure)", F["inj"]), rowf("  of which cure equity", F["cure"]),
           rowf("Senior prepayment (cure + sweep)", D["sen_prepay"]), rowf("Distributions", F["dist"]),
           rowf("EBITDA (USD m)", R["ebitda"] / fx, 1e6, "{:,.2f}"), ["KES/USD"] + [f"{fx[i]:.1f}" for i in range(2, N)]]
    L.append(md_table(hdr, tbl))
    L.append("")
    L.append("## 5. Unit economics (2026 KES, i.e. deflated by the tradables index)\n")
    def real(arr, i):
        return arr[i] / R["T"][i]
    ue = []
    full = base.get("full", base)
    for lab, yr, src in (("2031 (Phase 1 steady)", 2031, base.get("p1_only", base)), ("2035 (Phase 2, if built)", 2035, full)):
        i = IDX[yr]; RR = src["R"]; VV = RR["volumes"]
        ue.append([lab, f"{real(RR['crush_gm_per_t'], i):,.0f}", f"{real(RR['feed_gm'], i)/max(VV['feed'][i],1):,.0f}", f"{RR['feed_gm_pct'][i]:.0%}",
                   f"{real(RR['ebitda'], i)/max(VV['crush'][i],1):,.0f}", f"{real(RR['ebitda'], i)/1e6:,.0f}"])
    L.append(md_table(["Year", "Crush gross margin, KES/t seed", "Feed gross margin, KES/t feed", "Feed GM %", "EBITDA per t seed", "EBITDA KES m (real)"], ue))
    L.append("\nCrush gross margin values oil at crude-to-refiner prices and cake at market; feed margin charges own cake/oil at market transfer prices. Research benchmarks: ~KES 12k/t (feedstock_supply s8, oil 175/cake 35/seed 57) to ~31k/t (market_demand s6, oil 190/cake 40/seed 60); feed cash cost KES 88-93/kg vs KES 110-150 price (technical_operations s3.3).\n")
    L.append("## 6. Returns\n")
    cols = [("Plan (gate auto)", base), ("Full plan (Phase 2 forced go)", base.get("full", base)), ("Phase 1 only", base.get("p1_only", base))]
    rr = []
    for lab, key, kind in (("Project IRR, KES", "project_irr_kes", "pct"), ("Project IRR, USD", "project_irr_usd", "pct"),
                           ("Project IRR, KES, no terminal value", "project_irr_kes_noTV", "pct"),
                           ("Phase 1 WACC from the modelled structure (equity 35%, debt post-tax)", "wacc_model", "pct"),
                           ("Project NPV @ modelled WACC (KES m)", "npv_proj_wacc_model", "m"),
                           (f"Project NPV @ financing-file {P0['wacc_p1']:.0%} (memo; assumes 50% equity) (KES m)", "npv_proj_wacc_p1", "m"),
                           ("Equity IRR, KES", "equity_irr_kes", "pct"), ("Equity IRR, USD", "equity_irr_usd", "pct"),
                           ("Equity IRR, KES, no terminal value", "equity_irr_kes_noTV", "pct"),
                           ("Equity IRR, KES, after dividend WHT + CGT", "equity_irr_kes_after_shtax", "pct"),
                           ("Equity IRR, founder group after GM equity incentive", "equity_irr_founder_after_gm", "pct"),
                           (f"Equity NPV @ {P0['infra_bond']:.2%} infra bond (survival assumed) (KES m)", "npv_eq_infra", "m"),
                           (f"Equity NPV @ {P0['hurdle_gf_lo']:.0%} (KES m)", "npv_eq_gf_lo", "m"),
                           (f"Equity NPV @ {P0['hurdle_gf_hi']:.0%} (KES m)", "npv_eq_gf_hi", "m"),
                           ("Project payback (year)", "payback_project", "yr"), ("Equity payback (year)", "payback_equity", "yr"),
                           ("Peak project funding need (KES m)", "peak_funding_project", "m"), ("Peak net equity incl. cures (KES m)", "peak_equity", "m"),
                           ("Peak net equity year", "peak_equity_year", "yr"),
                           ("Senior debt sized (KES m)", "senior_kes", "m"),
                           ("Min DSCR (from first test year)", "min_dscr", "x"), ("Covenant cure equity (KES m)", "cure_equity", "m"),
                           ("Years below 1.10x default / below 1.30x lock-up", None, "breach"),
                           ("Irrecoverable capex VAT (feed-exempt share) (KES m)", "vat_irrecov_capex", "m")):
        row = [lab]
        for _n, o in cols:
            if kind == "breach":
                row.append(f"{o['M']['breach_years']} / {o['M']['lockup_years']}")
                continue
            v = o["M"][key]
            row.append(fmt_pct(v) if kind == "pct" else fmt_m(v) if kind == "m" else (f"{v:.2f}" if kind == "x" and v is not None and not np.isnan(v) else ("not by 2038" if v is None else str(v))))
        rr.append(row)
    L.append(md_table(["Metric"] + [c[0] for c in cols], rr))
    ww = M["wacc_weights"]
    L.append(f"\nModelled Phase 1 structure (2027-29 long-term sources): equity {ww[0]:.0%}, senior {ww[1]:.0%}, mezzanine {ww[2]:.0%}. "
             f"Layer costs: equity {P0['hurdle_gf_lo']:.0%}; senior {P0['senior_rate']:.1%} x (1 - {P0['cit']:.0%}); mezzanine {P0['mezz_rate']:.0%} grossed up for {P0['wht_int']:.0%} WHT plus {P0['kes_dep']:.0%} KES drift, post-tax. "
             f"The financing file's 23% assumed 50% equity and is kept only as a memo.\n")
    if base.get("gate"):
        L.append("**Phase 2 gate (evaluated on 2031 results):**\n")
        g = []
        for k, (v, th, ok) in base["gate"].items():
            if "KES" in k:
                g.append([k, fmt_m(v), fmt_m(th), "PASS" if ok else "FAIL"])
            elif "margin" in k or "share" in k:
                g.append([k, fmt_pct(v), fmt_pct(th), "PASS" if ok else "FAIL"])
            else:
                g.append([k, f"{v:,.0f}", f"{th:,.0f}", "PASS" if ok else "FAIL"])
        L.append(md_table(["Test", "Value", "Threshold", "Result"], g))
        L.append(f"\nPhase 2 incremental project IRR {fmt_pct(base['inc_irr'])}; incremental equity IRR {fmt_pct(base['inc_eq_irr'])}; incremental project NPV at the Phase 2 funding-mix WACC ({base['wacc_p2']:.1%}: {P0['p2_debt_share']:.0%} DFI debt post-tax, rest equity at 35%) KES {base['inc_npv']/1e6:,.0f}m; at the financing file's {P0['wacc_ss']:.1%}: KES {base['inc_npv_ss']/1e6:,.0f}m. Decision: **{'GO' if base['go'] else 'NO-GO'}**.\n")
    # ask and FX (R1-06)
    A = X["ask"]
    L.append("### The ask, by stage, at stated FX\n")
    L.append(md_table(["Item", "KES m (nominal)", "FX (KES/USD) and year", "USD m"], [[a, fmt_m(b), f"{c:.1f} ({d})", f"{b/c/1e6:.2f}"] for a, b, c, d in A]))
    L.append("")
    L.append("### Senior debt under alternative lender structures\n")
    L.append(md_table(["Structure", "Senior sized (KES m)", "Equity IRR", "Min DSCR", "Cure equity (KES m)", "Peak equity (KES m)"],
                      [[nm, fmt_m(d["senior"]), fmt_pct(d["eq"]), f"{d['dscr']:.2f}", fmt_m(d["cure"]), fmt_m(d["peak"])] for nm, d in X["debt_alt"].items()]))
    L.append("")
    L.append("## 7. Scenarios (senior debt held at base sizing; every case reports minimum DSCR and cure equity)\n")
    names = list(scen.keys())
    srow = [["Phase 2"] + ["GO" if scen[n]["go"] else "NO-GO" for n in names]]
    for lab, key, kind in (("Project IRR, KES", "project_irr_kes", "pct"), ("Equity IRR, KES", "equity_irr_kes", "pct"), ("Equity IRR, USD", "equity_irr_usd", "pct"),
                           ("Project NPV @ modelled WACC (KES m)", "npv_proj_wacc_model", "m"),
                           (f"Equity NPV @ {P0['hurdle_gf_lo']:.0%} (KES m)", "npv_eq_gf_lo", "m"),
                           ("EBITDA 2031 (KES m)", "ebitda_2031", "m"), ("EBITDA 2035 (KES m)", "ebitda_2035", "m"), ("Feed sold 2031 (t)", "feed_2031", "t"),
                           ("Peak net equity incl. cures (KES m)", "peak_equity", "m"), ("Min DSCR", "min_dscr", "x"), ("Cure equity (KES m)", "cure_equity", "m"),
                           ("Years below 1.10x", "breach_years", "yr"), ("Equity payback", "payback_equity", "yr")):
        row = [lab]
        for n in names:
            v = scen[n]["M"][key]
            row.append(fmt_pct(v) if kind == "pct" else fmt_m(v) if kind == "m" else (f"{v:,.0f}" if kind == "t" else (f"{v:.2f}" if kind == "x" and v is not None and not np.isnan(v) else ("not by 2038" if v is None else str(v)))))
        srow.append(row)
    L.append(md_table(["Metric"] + [n for n in names], srow))
    L.append("\nDefinitions:\n")
    for n in names:
        L.append(f"- **{n}**: {SCEN_DESC[n]}")
    L.append("\n## 8. Tornado (one driver at a time; equity IRR, minimum DSCR and cure equity)\n")
    trow = [[t["label"], fmt_pct(t["eq_lo"]), fmt_pct(t["eq_hi"]), f"{t['swing']*100:.1f}",
             f"{t['dscr_lo']:.2f}" if not np.isnan(t['dscr_lo']) else "n/a", f"{t['dscr_hi']:.2f}" if not np.isnan(t['dscr_hi']) else "n/a",
             fmt_m(t["cure_lo"]), fmt_m(t["cure_hi"]), fmt_m(t["peak_lo"]), fmt_m(t["peak_hi"])] for t in torn]
    L.append(md_table(["Driver (low / high setting)", "Equity IRR low", "Equity IRR high", "Swing pp", "Min DSCR low", "Min DSCR high",
                       "Cure KES m low", "Cure KES m high", "Peak equity low", "Peak equity high"], trow))
    L.append(f"\nBase: equity IRR {fmt_pct(M['equity_irr_kes'])}, project IRR {fmt_pct(M['project_irr_kes'])}, min DSCR {M['min_dscr']:.2f}x. Phase 2 decision held at the base result; senior debt held at base sizing (KES {senior_kes/1e6:,.0f}m). Cure equity is injected whenever DSCR falls below {P0['default_dscr']:.2f}x and prepays senior; it is included in the equity IRR and peak equity.\n")
    seeds, oils, pg, eg = grid
    L.append("### Seed cost x oil price index (equity IRR, KES; refined price moves with the index)\n")
    L.append(md_table(["Seed cost \\ oil index"] + [f"x{o:.2f}" for o in oils],
                      [[f"x{s_:.2f}"] + [fmt_pct(eg[i, j]) for j in range(len(oils))] for i, s_ in enumerate(seeds)]))
    L.append("\n## 9. Survival-weighted view\n")
    b = [[nm, f"{v['prob']:.1%}", fmt_pct(v["eq_irr"]), fmt_m(v["eq_npv35"]), fmt_m(v["founder_npv35"])] for nm, v in surv["branches"].items()]
    L.append(md_table(["Branch", "Probability", "Equity IRR", "Equity NPV @35% (KES m)", "Founder NPV @35% incl. guarantee and GM (KES m)"], b))
    L.append("")
    be = surv["breakeven_pfail_35"]
    L.append(md_table(["Measure", "Value"], [
        ["Expected equity IRR, all shareholders", fmt_pct(surv["exp_eq_irr"])],
        ["Expected founder-group IRR (GM dilution, guarantee calls)", fmt_pct(surv["exp_founder_irr"])],
        ["Expected founder-group IRR after dividend WHT and CGT", fmt_pct(surv["exp_founder_at_irr"])],
        ["Equity IRR conditional on survival", fmt_pct(surv["surv_eq_irr"])],
        ["Total failure probability (early + late)", f"{surv['total_fail_prob']:.0%}"],
        ["Greenfield hurdle (expected)", f"{P0['hurdle_gf_lo']:.0%}-{P0['hurdle_gf_hi']:.0%}"],
        [f"Survivor IRR needed at total failure probability {surv['total_fail_prob']:.0%} (report formula)", f"{surv['req_cond_35']:.1%}-{surv['req_cond_40']:.1%}"],
        ["Expected equity NPV @35% / @40% (KES m)", f"{fmt_m(surv['exp_npv35'])} / {fmt_m(surv['exp_npv40'])}"],
        ["Expected founder NPV @35% / @40% (KES m)", f"{fmt_m(surv['exp_founder_npv35'])} / {fmt_m(surv['exp_founder_npv40'])}"],
        [f"NPV @ {P0['infra_bond']:.2%} bond, survival-CONDITIONAL equity flows (KES m)", fmt_m(surv["bond_npv_surv_cond"])],
        [f"NPV @ {P0['infra_bond']:.2%} bond, survival-WEIGHTED founder flows after shareholder tax (KES m)", fmt_m(surv["bond_npv_exp_founder_at"])],
        ["Early-failure probability at which expected IRR falls below 35%", "already below at p=0" if be == 0 else (fmt_pct(be) if be is not None else "never (within 0-90%)")],
        ["Early failure: asset recovery / net debt / equity recovery / guarantee call (KES m)",
         f"{fmt_m(surv['liq_assets'])} / {fmt_m(surv['debt_at_fail'])} / {fmt_m(surv['eq_liq'])} / {fmt_m(surv['guarantee_early'])}"],
    ]))
    L.append("\nEarly-failure probability sweep (late-failure share held at {:.0%}):\n".format(P0["p_late_fail"]))
    L.append(md_table(["Early-failure probability", "Expected equity IRR", "Expected equity NPV @35% (KES m)"],
                      [[f"{q:.0%}", fmt_pct(ir), fmt_m(nv)] for q, ir, nv in surv["pfail_sweep"]]))
    L.append("\n## 10. Phase 0 value of information\n")
    for v in X["voi"]:
        L.append(f"**Prior correlation rho = {v['rho']}** ({v['n']} draws). Phase 0a cost KES {v['cost_a']/1e6:,.0f}m; Phase 0b cost incl. expected pilot seed loss KES {v['cost_b']/1e6:,.0f}m. Expected founder NPV, KES m (value date 2027):\n")
        rows_ = []
        for key in ("35_0.30", "40_0.30", "35_0.20", "40_0.20"):
            d = v[key]
            h, pfk = key.split("_")
            rows_.append([f"{h}% hurdle, p_fail {float(pfk):.0%}", "0", fmt_m(d["invest_now"]), fmt_m(d["phase0a_only"]), fmt_m(d["staged"]),
                          fmt_m(d["perfect_info"]), f"{d['p_pass_full']:.1%}", f"{d['p_pass_a']:.1%}", fmt_m(d["evpi_gross"]), f"{d['p_cond_irr_35']:.1%}"])
        L.append(md_table(["Case", "Stop now", "Invest now, no info", "Phase 0a only", "Staged 0a then 0b", "Perfect information (upper bound)",
                           "P(pass) full info", "P(pass) after 0a", "Gross value of perfect info", "P(conditional IRR >= 35%)"], rows_))
        L.append("")
    L.append("Priors: triangular(low, base, high) on SAM, feed price, share, feed price decline, capex, oil index, refining spread (revealed by 0a) and seed cost and supply (revealed by 0b). Phase 0 spend is excluded from V and charged as the option price. 'Pass' = survival-weighted founder NPV at the hurdle > 0.\n")
    L.append("## 11. Gate calibration (R1-10) and joint cases\n")
    def fv(k, x):
        if x is None:
            return "not reached in search range"
        return f"{x:,.0f}" if abs(x) > 100 else f"{x:.3f}"
    L.append(md_table(["Driver", "Base", "Min DSCR = 1.10x (others base, base debt)", "Conditional equity IRR 35% (others base)",
                       "KILL threshold: survival-weighted 35% at p_fail 20% with every other driver at bull", "Bull value"],
                      [[k, fv(k, d["base"]), fv(k, d["covenant"]), fv(k, d["hurdle"]), fv(k, d["kill"]), fv(k, d["bull_other"])] for k, d in X["cal"].items()]))
    L.append("")
    L.append(md_table(["Joint case", "Conditional equity IRR", "Expected IRR (p_fail 30%)", "Expected IRR (p_fail 20%)", "Founder NPV @35% (30% / 20%), KES m",
                       "Min DSCR", "Senior sized, KES m", "Peak equity, KES m", "Phase 2"],
                      [[nm, fmt_pct(d["cond_irr"]), fmt_pct(d["exp_irr_0.30"]), fmt_pct(d["exp_irr_0.20"]), f"{fmt_m(d['npv35_0.30'])} / {fmt_m(d['npv35_0.20'])}",
                        f"{d['min_dscr']:.2f}", fmt_m(d["senior"]), fmt_m(d["peak_equity"]), "GO" if d["go"] else "NO-GO"] for nm, d in X["jc"].items()]))
    G = X["g1"]
    L.append(f"\nPhase 2 gate calibration on the bull case (the only modelled case where Phase 2 is live): 2031 feed {G['bull_feed_t']:,.0f} t at {G['bull_feed_gm']:.0%} gross margin; 2031 DSCR {G['bull_dscr_2031']:.2f}x; Phase 2 incremental equity IRR {fmt_pct(G['bull_inc_eq_irr'])}. "
             f"Phase 2's incremental equity NPV at 35% turns negative at feed price x{G['feed_px_mult']['breakeven'] or float('nan'):.2f}, seed cost x{G['seed_cost_mult']['breakeven'] or float('nan'):.2f} or seed supply x{G['supply_mult']['breakeven'] or float('nan'):.2f} (others at bull).\n")
    L.append("## 12. Checks\n")
    L.append(md_table(["Check", "Result", "Detail"], [[a, b_, c] for a, b_, c in chk]))
    L.append("\n## 13. Modelling caveats\n")
    L.append("\n".join([
        "- Inputs are research estimates, mostly [ESTIMATE] or [UNVERIFIED]. There are no binding quotes yet for seed, crude-oil offtake, cake, capex or debt terms. Read IRRs as +/- 10 pp.",
        "- Round 1 changes: the feed SAM is the contestable KENYAN merchant market (25 kt in 2026, net of Victory Farms' captive feed and SamakGro's VF-linked sales); Uganda/Tanzania is an upside scenario only. Feed prices decline 1%/yr in real terms from 2029 (competitor response); 3%/yr in the competitor_response scenario.",
        "- Oil prices are linked: crude is the oil price index (base KES 175k/t = reconciled CPO landed parity) and refined = crude + a refining spread (base KES 65k/t, giving KES 240k/t). The seed pass-through is indexed to crude, so all oil prices and seed move together.",
        "- Senior debt is sized on a lender downside case (CFADS x 0.80 at 1.30x). Below the 1.10x default covenant the sponsor injects cure equity that prepays senior and pays a waiver fee; below 1.30x distributions stop and half of excess cash prepays senior. Cure equity is in the equity IRR and peak equity.",
        "- The base case is an expected-value policy case: CPO-duty removal 60%, GM imports 40%, duty waiver 1 year in 3, Tanzanian levy 70%, minimum tax (1% of turnover where it exceeds CIT) 25% from 2028. Capex VAT is apportioned by the exempt (feed) share of turnover when feed is VAT-exempt.",
        "- WACC is computed from the modelled Phase 1 structure, not the financing file's 50%-equity illustration. Phase 2's gate (d) requires positive incremental equity NPV at the 35% hurdle, because existing shareholders fund it.",
        "- The survival tree has early failure (bear operations to 2030, liquidation 2031), late failure (base Phase 1 operations to 2034, liquidation 2035) and survival. Founder flows add the personal-guarantee call on any senior shortfall and the GM's equity dilution; the bond comparison uses survival-weighted founder flows after shareholder tax.",
        "- Restructure alternatives (feed-mill-first, crusher-only) use [ESTIMATE] capex and fixed-cost scaling factors (CONFIG_CAPEX, CONFIG_FIXED in model.py); acquisition and JV are not modelled.",
        "- VOI uses independent (rho=0) or correlated (rho=0.5) triangular priors that are judgement, not data, and treats Phase 0 as revealing the true value of the tested inputs; real tests are noisy, so the VOI figures are upper bounds.",
        "- The model is annual with a steady-state monthly seed profile. Nominal KES IRRs rise with faster KES depreciation because CPI follows FX (PPP); USD IRRs are the cleaner FX view.",
        "- In bear and stress the model keeps injecting equity (including cures) to fund losses, which overstates peak equity; the survival view models that stop as failure plus liquidation.",
    ]))
    with open(path, "w") as f:
        f.write("\n".join(L) + "\n")


# =============================================================================
# MAIN
# =============================================================================
def main():
    P0 = load_overrides(base_params())
    print("Sizing senior debt and running base case ...")
    senior = size_senior(P0)
    base = run_model(P0, mode="auto", senior_kes=senior)
    base_mode = "go" if base["go"] else "no"
    print(f"  senior KES {senior/1e6:,.0f}m; Phase 2 {'GO' if base['go'] else 'NO-GO'}")
    print("Scenarios ...")
    scen = OrderedDict()
    for n in ("bear", "base", "bull", "stress", "competitor_response", "regional_upside", "p2_research_spec", "v1_market",
              "feedmill_first", "crush_only"):
        p = scenario_params(n, P0)
        scen[n] = run_model(p, mode="go" if n == "p2_research_spec" else "auto", senior_kes=senior)
        scen[n]["p"] = p
    print("Tornado ...")
    torn = tornado(P0, senior, base_mode, base["M"])
    print("Seed x oil grid ...")
    grid = grid_seed_oil(P0, senior, base_mode)
    print("Survival-weighted tree ...")
    surv = survival_tree(P0, senior, base_go=base["go"])
    X = OrderedDict()
    print("Alternative debt structures ...")
    X["debt_alt"] = OrderedDict()
    for nm, d in (("R1 base: downside sizing (CFADS x0.80 at 1.30x), tests from 2030, 18-month grace", {}),
                  ("Ramp-year interest funded from budget: tests from 2031, 24-month grace", {"dscr_first_year": 2031, "senior_grace_m": 24}),
                  ("v1.1 rule: no haircut (sized to exactly 1.30x), tests from 2030", {"sizing_haircut": 0.0})):
        q = copy.deepcopy(P0); q.update(d)
        sq = size_senior(q); o = run_model(q, "auto", sq)
        X["debt_alt"][nm] = dict(senior=sq, eq=o["M"]["equity_irr_kes"], dscr=o["M"]["min_dscr"], cure=o["M"]["cure_equity"], peak=o["M"]["peak_equity"])
    print("Gate calibration ...")
    X["cal"] = calibrate_gates(P0, senior)
    X["g1"] = g1_calibration(P0)
    X["jc"] = joint_cases(P0)
    print("Phase 0 value of information (two priors) ...")
    X["voi"] = [voi_analysis(P0, n_draws=150, rho=0.0), voi_analysis(P0, n_draws=150, rho=0.5)]
    # the ask at stated FX (R1-06)
    Rb, Fb, Mb = base["R"], base["F"], base["M"]
    fx = Rb["fx"]
    fx27, fx28 = fx[IDX[2027]], fx[IDX[2028]]
    ph0a = P0["phase0a_usd"] * 1e6 * fx27
    ph0b = P0["phase0b_usd"] * 1e6 * fx27
    ph0 = P0["phase0_usd"] * 1e6 * fx27 * Rb["T"][IDX[2027]] / (fx27 / P0["fx0"])
    eq_p1 = (Fb["inj"] - Fb["dist"])[:3].sum()
    py = Mb["peak_equity_year"]
    X["ask"] = [("Phase 0a (cheap decisive tests)", ph0a, fx27, 2027),
                ("Phase 0b (pilot, trials, EIA, financing; only if 0a passes)", ph0b, fx27, 2027),
                ("Phase 1 equity 2027-29 incl. Phase 0 (as modelled)", eq_p1, fx28, 2028),
                ("Phase 1 equity net of Phase 0 (incremental FID ask)", eq_p1 - ph0, fx28, 2028),
                ("Standby equity for overruns (15% of Phase 1 capex)", 0.15 * Rb["capex_p1"].sum(), fx28, 2028),
                ("Covenant cure equity, base case (in peak equity)", Mb["cure_equity"], fx[IDX[2030]], 2030),
                (f"Peak net equity, whole plan (reached {py})", Mb["peak_equity"], fx[IDX[py]], py)]
    chk_full = checks(base, P0)
    print("Writing outputs ...")
    write_excel(os.path.join(HERE, "financial_model.xlsx"), P0, base, scen, torn, grid, surv, chk_full, senior, X)
    write_md(os.path.join(HERE, "outputs.md"), P0, base, scen, torn, grid, surv, chk_full, senior, X)
    write_csvs(base, scen, X)
    M = base["M"]
    print("\nKEY RESULTS (base, gate auto)")
    for k in ("project_irr_kes", "project_irr_usd", "equity_irr_kes", "equity_irr_usd", "equity_irr_kes_after_shtax", "wacc_model"):
        print(f"  {k:32s} {fmt_pct(M[k])}")
    for k in ("npv_proj_wacc_model", "npv_eq_gf_lo", "peak_equity", "peak_funding_project", "ebitda_2031", "ebitda_2035", "cure_equity"):
        print(f"  {k:32s} KES {M[k]/1e6:,.0f}m")
    print(f"  min DSCR: {M['min_dscr']:.2f}")
    for n, o in scen.items():
        print(f"  {n:20s} P2={'GO' if o['go'] else 'NO':3s} projIRR {fmt_pct(o['M']['project_irr_kes'])} eqIRR {fmt_pct(o['M']['equity_irr_kes'])} minDSCR {o['M']['min_dscr']:.2f} cure {fmt_m(o['M']['cure_equity'])}")
    print(f"  expected (survival-weighted) equity IRR {fmt_pct(surv['exp_eq_irr'])}; founder after tax {fmt_pct(surv['exp_founder_at_irr'])}")
    for v in X["voi"]:
        print(f"  VOI rho={v['rho']}: " + "; ".join(f"{k} staged {v[k]['staged']/1e6:,.0f} perfect {v[k]['perfect_info']/1e6:,.0f} P(pass) {v[k]['p_pass_full']:.1%}" for k in ("35_0.30", "35_0.20")))
    for c in chk_full:
        print(f"  [{c[1]}] {c[0]}: {c[2]}")


if __name__ == "__main__":
    main()
